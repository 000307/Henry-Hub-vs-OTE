# Python 3.10+; pip install numpy pandas openpyxl plotly statsmodels scipy pillow "xlrd>=2.0.1"
# Při tvorbě reportu se doplní původní OTE Excel a načtou data EIA, ČNB, Fed a CME.
# Místní vstupy: původní OTE Excel a Forwardy OTE.xlsx (EUR/MWh). Henry Hub: list Data 1 ze staženého XLS EIA. Logo je vloženo přímo do skriptu.
from __future__ import annotations

import argparse
import base64
from io import BytesIO
from bisect import bisect_left
from html.parser import HTMLParser
import json
import math
import sys
from numbers import Real
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
import plotly.graph_objects as go
from plotly.utils import PlotlyJSONEncoder
from plotly.offline import get_plotlyjs
from statsmodels.nonparametric.smoothers_lowess import lowess
from scipy.optimize import least_squares


import re
from datetime import date, datetime, timedelta, timezone
from urllib.error import HTTPError, URLError
from urllib.parse import urlencode
from urllib.request import Request, urlopen

import unicodedata
from datetime import date, datetime
from pathlib import PurePosixPath
from urllib.parse import quote, unquote, urljoin, urlsplit, urlunsplit
from zipfile import BadZipFile, ZipFile
from copy import copy
from datetime import date, datetime, timedelta
import os
import posixpath
import shutil
import tempfile
import xml.etree.ElementTree as ET
from zipfile import ZipFile
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from openpyxl.styles.numbers import is_date_format

CME_FORWARD_SOURCE = "https://www.cmegroup.com/markets/energy/natural-gas/natural-gas.settlements.html"
# Endpoint used by the CME web table; CME may change this public web interface.
CME_FORWARD_API = "https://www.cmegroup.com/CmeWS/mvc/Settlements/Futures"
_FORWARD_MONTHS = {name: index for index, name in enumerate(
    ("JAN", "FEB", "MAR", "APR", "MAY", "JUN", "JUL", "AUG", "SEP", "OCT", "NOV", "DEC"), 1)}
_FORWARD_MONTH_NAMES = ("JANUARY", "FEBRUARY", "MARCH", "APRIL", "MAY", "JUNE",
                        "JULY", "AUGUST", "SEPTEMBER", "OCTOBER", "NOVEMBER", "DECEMBER")


def _forward_timestamp() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def _forward_date(value) -> date:
    """Parse CME US/ISO dates without depending on the process locale."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%m/%d/%Y", "%Y-%m-%d", "%m/%d/%y", "%Y%m%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    raise ValueError("Neplatné datum vypořádání CME: " + text[:80])


def _forward_contract(value) -> tuple[str, str]:
    text = " ".join(str(value).replace("\u00a0", " ").upper().strip().split())
    match = re.fullmatch(r"([A-Z]+)[\s-]+(\d{2}|\d{4})", text)
    if not match:
        raise ValueError("Neplatný měsíc kontraktu CME: " + text[:80])
    name, raw_year = match.groups()
    if name not in _FORWARD_MONTHS and name not in _FORWARD_MONTH_NAMES:
        raise ValueError("Neznámý měsíc kontraktu CME: " + name[:30])
    month = _FORWARD_MONTHS.get(name) or (_FORWARD_MONTH_NAMES.index(name) + 1)
    year = int(raw_year) + (2000 if len(raw_year) == 2 else 0)
    contract_date = date(year, month, 1)
    label = str(value).strip()  # Preserve the Month column exactly as CME supplied it.
    return label, contract_date.isoformat()


def _forward_settle(value) -> float | None:
    if value is None:
        return None
    text = str(value).strip().replace("\u00a0", "")
    if text.upper() in ("", "-", "--", "—", "N/A", "NA", "NAN", "NONE", "NULL"):
        return None
    if not re.fullmatch(r"[+-]?(?:\d+(?:\.\d+)?|\.\d+|\d{1,3}(?:,\d{3})+(?:\.\d+)?)[abAB]?", text):
        raise ValueError("Neplatná cena Settle CME: " + text[:80])
    result = float(re.sub(r"[abAB]$", "", text).replace(",", ""))
    if not math.isfinite(result):
        raise ValueError("Cena Settle CME není konečné číslo.")
    return result


def parse_cme_forward_rows(raw_rows: list) -> list[dict]:
    by_month = {}
    for raw in raw_rows:
        if not isinstance(raw, dict):
            raise ValueError("CME vrátilo neplatný řádek vypořádání.")
        raw_month = str(raw.get("month", "")).strip()
        if raw_month.upper() in ("TOTAL", "TOTALS", ""):
            continue
        month, contract_date = _forward_contract(raw_month)
        settle = _forward_settle(raw.get("settle"))
        if settle is None:
            continue
        row = {"month": month, "contract_date": contract_date, "settle": settle}
        if contract_date in by_month and by_month[contract_date]["settle"] != settle:
            raise ValueError("CME vrátilo rozdílné ceny pro stejný měsíc: " + month)
        by_month[contract_date] = row
    if not by_month:
        raise ValueError("CME nevrátilo žádné platné ceny Month / Settle.")
    return [by_month[key] for key in sorted(by_month)]


class _CMEAccessError(RuntimeError):
    """A transport/access failure must not trigger requests for additional dates."""


def _forward_json(url: str):
    request = Request(url, headers={
        "User-Agent": "Mozilla/5.0 (compatible; HenryHubForwardReport/1.0)",
        "Accept": "application/json", "Referer": CME_FORWARD_SOURCE,
    })
    try:
        with urlopen(request, timeout=10) as response:
            body = response.read(5_000_001)
    except HTTPError as exc:
        if exc.code in (401, 403, 429):
            raise _CMEAccessError(f"CME odmítlo přístup (HTTP {exc.code}).") from exc
        raise
    except (URLError, TimeoutError) as exc:
        raise _CMEAccessError("Síťové připojení k CME selhalo: " + str(exc)[:150]) from exc
    if len(body) > 5_000_000:
        raise ValueError("Odpověď CME je příliš velká.")
    try:
        return json.loads(body.decode("utf-8-sig"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise _CMEAccessError("CME nevrátilo JSON; přístup mohl být odmítnut.") from exc


def _forward_trade_dates(payload) -> list[date]:
    if isinstance(payload, dict):
        payload = payload.get("tradeDates", payload.get("dates"))
    if not isinstance(payload, list):
        raise ValueError("Neznámý formát seznamu dat vypořádání CME.")
    result = set()
    for item in payload:
        if isinstance(item, dict):
            item = item.get("tradeDate", item.get("formattedDate", item.get("date")))
        try:
            result.add(_forward_date(item))
        except ValueError:
            continue
    if not result:
        raise ValueError("CME nevrátilo datum vypořádání.")
    return sorted(result, reverse=True)


def _forward_settlement_pages(trade_date: date, request_json=_forward_json) -> dict:
    page_size = 500
    raw_rows = []
    fingerprints = set()
    total_expected = None
    pages_expected = None
    update_time = ""
    for page in range(1, 21):
        query = urlencode({"tradeDate": trade_date.strftime("%m/%d/%Y"), "pageSize": page_size,
                           "pageNumber": page, "isProtected": "false"})
        payload = request_json(CME_FORWARD_API + "/Settlements/444/FUT?" + query)
        if not isinstance(payload, dict):
            raise ValueError("Neznámý formát vypořádání CME.")
        outer = payload
        payload = payload.get("dailySettlementList", payload)
        if isinstance(payload, list) and len(payload) == 1:
            payload = payload[0]
        if not isinstance(payload, dict) or not isinstance(payload.get("settlements"), list):
            raise ValueError("V odpovědi CME chybí seznam settlements.")
        returned_date = payload.get("tradeDate", outer.get("tradeDate"))
        if returned_date and _forward_date(returned_date) != trade_date:
            raise ValueError("CME vrátilo jiné datum vypořádání než požadované.")
        batch = payload["settlements"]
        total = payload.get("total", outer.get("total"))
        if total is not None:
            if not re.fullmatch(r"\d+", str(total)):
                raise ValueError("CME vrátilo neplatný počet kontraktů.")
            total = int(total)
            if total_expected is not None and total_expected != total:
                raise ValueError("Počet kontraktů CME se během načítání změnil.")
            total_expected = total
        total_pages = payload.get("totalPages", outer.get("totalPages"))
        if total_pages is not None:
            if not re.fullmatch(r"\d+", str(total_pages)):
                raise ValueError("CME vrátilo neplatný počet stránek.")
            total_pages = int(total_pages)
            if total_pages > 20:
                raise ValueError("CME překročilo limit stránek forwardové křivky.")
            if total_pages == 0 and (batch or raw_rows):
                raise ValueError("CME vrátilo řádky při nulovém počtu stránek.")
            if pages_expected is not None and pages_expected != total_pages:
                raise ValueError("Počet stránek CME se během načítání změnil.")
            pages_expected = total_pages
        if not batch:
            if pages_expected is not None and page < pages_expected:
                raise ValueError("CME vrátilo prázdnou stránku před koncem forwardové křivky.")
            if total_expected is not None and len(raw_rows) < total_expected:
                raise ValueError("CME vrátilo neúplnou stránku forwardové křivky.")
            break
        fingerprint = json.dumps(batch, sort_keys=True)
        if fingerprint in fingerprints:
            raise ValueError("CME opakuje stejnou stránku forwardové křivky.")
        fingerprints.add(fingerprint)
        raw_rows.extend(batch)
        update_time = str(payload.get("updateTime", outer.get("updateTime", update_time)))
        if pages_expected is not None:
            if page >= pages_expected:
                if total_expected is not None and len(raw_rows) < total_expected:
                    raise ValueError("CME vrátilo neúplnou forwardovou křivku.")
                break
        elif total_expected is not None:
            if len(raw_rows) >= total_expected:
                break
        elif len(batch) < page_size:
            break
    else:
        raise ValueError("CME překročilo limit stránek forwardové křivky.")
    return {"rows": parse_cme_forward_rows(raw_rows), "trade_date": trade_date.isoformat(),
            "update_time": update_time, "source_url": CME_FORWARD_SOURCE}


def fetch_cme_forward(request_json=_forward_json) -> dict:
    try:
        candidates = _forward_trade_dates(request_json(CME_FORWARD_API + "/TradeDates/444"))[:3]
    except _CMEAccessError:
        raise
    except Exception:
        try:
            from zoneinfo import ZoneInfo
            today = datetime.now(ZoneInfo("America/Chicago")).date()
        except (ImportError, KeyError):
            today = datetime.now(timezone.utc).date()
        candidates = [today - timedelta(days=offset) for offset in range(8)]
    errors = []
    for trade_date in candidates:
        try:
            return _forward_settlement_pages(trade_date, request_json)
        except _CMEAccessError:
            raise
        except Exception as exc:
            errors.append(str(exc))
    raise RuntimeError("Data CME se nepodařilo načíst. " + errors[0][:250])


def load_forward_for_report() -> dict[str, Any]:
    """Načte aktuální CME data jednou pro tento report, bez diskové cache."""
    try:
        data = fetch_cme_forward()
        return {
            "rows": parse_cme_forward_rows(data["rows"]),
            "trade_date": _forward_date(data["trade_date"]).isoformat(),
            "fetched_at": _forward_timestamp(),
            "error": None,
        }
    except Exception as exc:
        return {
            "rows": [], "trade_date": None, "fetched_at": None,
            "error": str(exc)[:500] or type(exc).__name__,
        }


_OTE_FORWARD_MONTHS = {name: number for number, name in enumerate(
    ("JAN", "FEB", "MAR", "APR", "MAY", "JUN", "JUL", "AUG", "SEP", "OCT", "NOV", "DEC"), 1)}
_OTE_FORWARD_MISSING = {"", "-", "--", "—", "N/A", "NA", "ND", "NAN", "NONE", "NULL"}


def _ote_forward_day(value, epoch=None) -> date:
    if value is None or isinstance(value, bool):
        raise ValueError("Chybí datum řádku forwardů OTE.")
    if isinstance(value, datetime):
        if pd.isna(value):
            raise ValueError("Neplatné datum řádku forwardů OTE.")
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, Real):
        if epoch is not None and math.isfinite(float(value)):
            converted = from_excel(value, epoch)
            if isinstance(converted, datetime):
                return converted.date()
        raise ValueError("Neplatné číselné datum řádku forwardů OTE.")
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d. %m. %Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    raise ValueError("Neplatné datum řádku forwardů OTE: " + text[:80])


def _ote_forward_price(value) -> float | None:
    if value is None:
        return None
    if isinstance(value, bool):
        raise ValueError("Cena forwardu OTE nesmí být logická hodnota.")
    if isinstance(value, Real):
        result = float(value)
    else:
        text = str(value).strip().replace("\u00a0", "").replace(" ", "")
        if text.upper() in _OTE_FORWARD_MISSING:
            return None
        # A comma or period denotes the decimal separator; Excel numeric cells
        # require no locale conversion. Mixed separators are ambiguous here.
        if not re.fullmatch(r"[+-]?(?:\d+(?:[.,]\d+)?|[.,]\d+)", text):
            raise ValueError("Neplatná cena forwardu OTE: " + text[:80])
        result = float(text.replace(",", "."))
    if not math.isfinite(result):
        raise ValueError("Cena forwardu OTE musí být konečné číslo.")
    return result


def _ote_forward_product(value) -> dict:
    product = " ".join(str(value).replace("\u00a0", " ").strip().split())
    text = product.upper()
    month = re.fullmatch(r"([A-Z]{3})\s*-\s*(\d{4})", text)
    quarter = re.fullmatch(r"Q([1-4])\s*-\s*(\d{4})", text)
    annual = re.fullmatch(r"CAL\s*-\s*(\d{4})", text)
    if month and month.group(1) in _OTE_FORWARD_MONTHS:
        year, month_number = int(month.group(2)), _OTE_FORWARD_MONTHS[month.group(1)]
        start = date(year, month_number, 1)
        end = date(year + (month_number == 12), month_number % 12 + 1, 1)
        kind = "month"
    elif quarter:
        year, quarter_number = int(quarter.group(2)), int(quarter.group(1))
        month_number = 3 * quarter_number - 2
        start = date(year, month_number, 1)
        end = date(year + (quarter_number == 4), (month_number + 2) % 12 + 1, 1)
        kind = "quarter"
    elif annual:
        year = int(annual.group(1))
        start, end, kind = date(year, 1, 1), date(year + 1, 1, 1), "calendar"
    else:
        raise ValueError("Neznámý produkt ve forwardech OTE: " + product[:80])
    return {"product": product, "kind": kind, "delivery_start": start.isoformat(), "delivery_end": end.isoformat()}


def read_ote_forward(path: Path, *, today: date | None = None) -> dict:
    """Choose one nearest calendar date; ties prefer the earlier observation.

    A missing price is skipped, never borrowed from another observation. Period
    ends are exclusive. The renderer plots monthly prices as separate markers,
    quarters and calendar years as separate horizontal segments.
    """
    path = Path(path)
    today = _ote_forward_day(today) if today is not None else pd.Timestamp.now(tz="Europe/Prague").date()
    with path.open("rb") as source:
        workbook = load_workbook(source, read_only=True, data_only=True)
        try:
            sheet = workbook.worksheets[0]
            records = sheet.iter_rows(values_only=True)
            header = next(records, None)
            if not header:
                raise ValueError("Soubor forwardů OTE je prázdný.")
            active_columns = [(index, value) for index, value in enumerate(header) if value is not None and str(value).strip()]
            dates = [index for index, value in active_columns if str(value).strip().casefold() in {"datum", "date"}]
            if len(dates) != 1:
                raise ValueError("Forwardy OTE musí obsahovat právě jeden sloupec Datum nebo Date.")
            date_column = dates[0]
            products = [(index, _ote_forward_product(value)) for index, value in active_columns if index != date_column]
            if not products:
                raise ValueError("Forwardy OTE neobsahují žádné cenové produkty.")
            keys = [(product["kind"], product["delivery_start"]) for _, product in products]
            if len(set(keys)) != len(keys):
                raise ValueError("Forwardy OTE obsahují duplicitní sloupce stejného produktu.")
            candidates = []
            for number, row in enumerate(records, start=2):
                if all(value is None or not str(value).strip() for value in row):
                    continue
                try:
                    observed = _ote_forward_day(row[date_column], workbook.epoch)
                except (ValueError, OverflowError) as exc:
                    raise ValueError(f"Forwardy OTE: neplatné datum na řádku {number}.") from exc
                candidates.append((observed, number, row))
            if not candidates:
                raise ValueError("Forwardy OTE neobsahují žádné datované řádky.")
            closest = min((observed for observed, _, _ in candidates), key=lambda day: (abs((day - today).days), day))
            selected = []
            for observed, number, row in candidates:
                if observed != closest:
                    continue
                try:
                    prices = tuple(_ote_forward_price(row[index] if index < len(row) else None) for index, _ in products)
                except ValueError as exc:
                    raise ValueError(f"Forwardy OTE: vadná cena na vybraném řádku {number}. {exc}") from exc
                selected.append(prices)
            if any(prices != selected[0] for prices in selected[1:]):
                raise ValueError(f"Forwardy OTE obsahují rozdílné ceny pro vybrané datum {closest.isoformat()}.")
            rows, skipped = [], []
            for (_, product), price in zip(products, selected[0]):
                if price is None:
                    skipped.append(product["product"])
                else:
                    rows.append({**product, "price": price})
            return {"rows": rows, "quote_date": closest.isoformat(),
                    "fetched_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
                    "source_file": path.name, "skipped_products": skipped,
                    "source_currency": "EUR", "source_unit": "mwh"}
        finally:
            workbook.close()


# Vstupní a výstupní soubory
INPUT_OTE_DEFAULT = "spotové ceny OTE - plyn.xlsx"
INPUT_OTE_FORWARD_DEFAULT = "Forwardy OTE.xlsx"
LOGO_DATA_URI = "data:image/webp;base64,UklGRu4OAABXRUJQVlA4TOEOAAAvTUEoEP8nFkzmD51C7/TP/wQkdJyHNWz/1zetvknTrIUa7pzgLt1wdzrrcOeszH3IBHeX4RPc3d1hwtiguLu7tJSupWmaz4t8/7/f7//Lv6d5dRbR/wmgf5Ef0Xtg2UASUacXO6oHkqjMeSzIFUCiyF/x4I2AUbSTsu2Dd4gtQPTm7kiKPAHMCQoMFcbfkeR6CiyyBYSCgH1B1B7A2IAQJQEDiLYCaBcQSgOS81J5L/CscAAoAgB6E20BsCQA9JrPLqK2AFAt8PO+z1OibB4AqwI/q32SiSgegDt3oCc0yecIEa0AgM8DPR/A9w0imuazKdCzFkBSHBHRFJ8HgR57ky875CDfpT7IEeARPcJUCBRFpDH1AkXdwTYMENnPcRUDRO+DzxcYyvOUe0EBYdtm8AcDQz1hOCEg1DDdqHEgqOhDGCYEB4DCTsD4Z8ocRzt0KxyjSWg37YZmt6KgLRCsmTkUSe+qW+x2Td5fpFvO5C+saAYED5POY6awJf1hIo7rNhRl9Dh5VrcBuBZkPb0h2l6rRLCN/CAyCWim2SZM16IxPKF6vfIAaG85bTJErgRlDt8C2KHZHSRH6rAeqKXXBwAOW03NFIh2p0zBeQcAKmmVB0AvDVwZwKda2c4DQH1rcT2E6Dl75tAdvgu1igFwzW7eJACztXoLvpssJeIMhGMpU7CdYtyFdOoLAO+Ylj0BQLxWvzLeMhYStA3COylzaAF+vE6rffaa9hkApAVrVB38LAsZB2FP+Uxit0FiuEZXfFDJJNt5H1TVaKVBal7L6AzxKZQ5VIZxb30iwc4yqSXY9/RxeQwwwioqJIs9zpFJLBa46dCmEZeay5xt3Ax9psL4STZrCDsH8Y+JcmUGhd0C6KJNbw59TSnt5Q5qkyNZAJ9bw0KIx9spOiP+q5yWNwGiR7VZbHDbYcZ08P8E6dIfolftVtAe4t6aZDsEIGVOeWuLeC6EprqcNUAHEyJeGKCsJs77QmhnAYWeSswmigPrXV/Byr4B6+W2aZItw+gPE3rBuKsm74P1cn9ZwDqIP8tNkQ85IGNWbssKvs1952FQUY9aEHxNmf2qwEQ9bGe5Qf8wqOd3sZD8nGgyRJ90tqpuYO86V3Dz9fhcZKGytyG4T483wSaG/cRt9LdstyTOO6iiRwhYkt2aTnIDqBbnLqjFbBF3PlV7RRL12M9NorKct4yfDYDk60R7IXuysBU1B5uah+hvBmO1OCaCIYoqQLi4DtXAZriIdjCY5V+5EyV2Er0N+dslLWgXN5eIunCJYRoEu4UeONXMEmurwwpuHRHFcKl5/WoCxD0VyHFOAe4UsZzK4CsRUfBdBr00eBXi3ZXkTBEbpYHLwzUgItsFBsP9KSJJYj7Rx1B6IrvVLOb2km9/7qbDvPe5ZdxhJd+DXZXObNdgCtjj5Psp9yTUj3pC3F2Uwh6owWyLKeTmYpncqQy6mDeTK3edQR0Fjltc9dPMA/Oikrl3mewJDD7zH9s1iZ+J+kJ1c2sZD/aqnaHZ3FHzDjEpQX24FQragT1IyxkUMq0v2IdOhsZzl+1+UwvinuIU9kTZGbuVhCdyXxNfiUNjs4L+YY5SZDKTXljud64D9efeNMt5jxtKfFEPg7Z+M0liGdHXUN/eSvqATQo3oD3cVrPKg11ANJ3BaKlosLcc9DY32KweYN35DGg196ff3JSoTo6bJvxuIY5b3BQyfptDBZO6cr2JSnuZx6EyC7nviFzcepNsZ7iFZFyPQ10/KQ7xw0TtYGZR6+gK1ltcwH6Zm2vSD1xzItrC4H2JvGnMPzmJbEnMLZPeAB8tQPHcej/pKtGDaJ8pn1jHcW4jiX7FuQuYs48rQETNuJMSg8D+RET0J4Nc5uzlfiPR7py3tH/MFEuNoOIwdbFlNAPfWCj8OYPRptieM0+IiGxnGDQWct7nyvnM4pqb8hr4NkLO+wx+9o/9YiuIBptz0jJ2cKdIfDKXEGZGcbAHfOhDbr1QF7A7yPdL7ntTlnE3goRoEJeaxy+ui7UjituTakaSVVQC/55EMS+Dr81oz01jQp8yGcVEDnMxTCNuhRn/8XB9SDzvSwbD/CJFyBNFRBTSar1bGZwWsZB7HCJB67kbDhNGcx8yNJrBJIFaYC/YmJzcJTMmg02OkqD53JNQf4Dwb2RYYEqqqihrKOTmRpBsIw6dTNjB1eIKpTMJYUbLuU+Jv8d4w9VFJXMzSLYKh8/8bYARUcENVjYerLugFJ3g4k14xIVztIzBZwaF0pmE7AY7GTRQ1xest4wU7ecu2/2suQhRxxdKbJYQlsgtIfkeHBorKwT2OhnW5M7buJFgx5PhJO4rZc673FaSb8WhjR8kCOUUo4p3FLwkS+gFvrqCVx5xW5XFcpuM6BCDGCbkMeMpahTHLVL2LvgWCuzXuIN+cF7kJsmWvCd32RIcN7mDpHI4hwqqBnOjBDpx25geYFeT8avcKVW209xZmwLqxaGOfltEdkjRay+l9lpCF/BpiSqTDeao2sh1EnDcYbxlfE5ydQVCMhhPqKIY8O5ElS8M1uo3TmSuHPWR+sESjhqYm5Zf0S2uogD1ZTCdiBqCjSfRiwxqKNpjYK63lHbtRIYrsB+W6WoFTaHnKDV5wLqDRXKlMC8iidZy3YTWcJ+oiYaeP2qXR+RjBdRUpoAVbNPkWZiS5txJEv6FQS9yZTD3nUJDuFlqlmqSmls3OizQRQUdEjtFFlARun6ppC+3VKw8d9U+HuwgEm7DHVFS1KMJhmjXT6CVkk/FBlnBfG2uO1Ss5PqK0S4GnROYl3nESnFpDhWToOvjUN2Kes3JlSFU2gIKurnzG9Qf4NBBxWXuTYk3uGSw80jcnsKgsoKoF9yVDep3c/hEN9pu1E4JnRTZTRYwFnxlUu+8yx1REAH+PxK2i4xhZQmK53oo+A58I1Jvv8BdsuvW3OhDNT+JvG0BYYncbjLzew4N5Rpwz0n2c6H9JDufmybnvMvFk5kfcWitm+24wXdqvhY4brOAnuBbmhKVzG2R68X9IRWWIPKOVC/uoNy74DubEvqIO6gbvWEwVU1bgVbkf44b3GmbKTSFQzmpRdzPUjRB4Kpdqjn3wi5jO83ddJhCQzjU0Y0OcOvUNDL63WYBncDHkbkuDzdH6iz3uZzLY9STpPNzKCvTAnxvMjd3KrdWu4oe5qyahgbeaLKAeO6e0yRayaXllwj1cA3kaI1BUoQcPeY6y+zmEsNNop85b0ndaAKTHqzkVYMpZAGNwfcjs6tzGCVRC3wOBfUNppLCfdwEiSrgJ5DZpb0MZmoXcsEHrylpwV3OpsuCodKtRLZyyTlMo9+5Z9nFPuHukMqjjLeEiqncXoklnLuwabSRS8mtG1V3+3yqpAvjqUu6KFwgUAH8NDI/lsMXYrO57Uq6MxtJ5QfcM7Ei6dwSMr8eh0HaUR+fZUr6MN+RFczlPMU0sF/krgUJxXMTlDgf+DRRUotDMaGJ4KtoQIe5R6Ha2dYCeGhTMcdnk80KCri51aTjJxzaiwSncf9VQoMBnCKlYQZtRCKTuD2kYwcOH2lHYacARKv4C8CZCLKC0eBraBH6mDssUhV8VTV5XwLvq6Fr3EiRb8G/roXjOnfRrh257gPDFDiSgYcusoKwBO4P0nMYh/oC73GeEDU0H49DFG3gtgo473BnbFrQ1xze0Y+ik3BaQS3gRQ2yhK/Bv6NJnlRus8AM7jwproIRpHgk90CgO/j3SM+wBO4PP6AmKagi1xcvm5CeRV2qcxvkdfF2Taigiy0qkNfF5ldFOwuqinDxArldvEMTyufi/YHedk+U2+eOpf+N89L/orF3nDK5UmMpkBvTWqZ7LAV2I2Vy0L/+//9HIzr9tGXN5BinVPm4WMrXZ8W2WW2dKnL2mLVt1YQmdqkacTFRgzat+/wVsZxxv2xZP7Otg0rGtRMIejeunlB0XFxo1sDWMxHszRiZnjjRIRG+FytIBQ1OAXu2jswUnLsC4C+7SMi4VLAXa8fhhgBtx6NQkXicpCxh0BIgdfWQMQcB70dSzzNwavzQzV7cjZQI2QU8Xzpw/BHAHSsFHP1+Q2USzHkIeDp/yA8ngLTVYk2AjwVqAh9mDcYA2wsQETW6B09tGeB7IqLm6fhGYi6wPCcR0duJSCkrtTaIhIMOABOzERE1vQ6I0RFcshstREK2LEGZDOx0EFsuGYek1hO7AjvFagNLbMTWTsdGqSIk/gHQm/hcZ2XaA20Mcr3EZMoSzkZKfjL8BqglU4P7HpfFtuBRBBmOh7eExA0St93EHpsBlfZI2K/gT4Nv4S2ZJbA/xhwyDkvFcIk0O/cx7gtlS8MoMi7sxRcSf0m8CjQmwVUS9ClQj7Ffw3bKEpYGDi0QfIKdEjeJj0OiUH1g/wLBf7BY4oDE53geJPKZTOgjbGTeBN7KGjSA7GmJ84o6QHa3KaNwkkTbyNBAeMv6bMU1e9agBbBpgfBEiROKugArFwgPN2UMjpqSMxmziai4F99Q1rAy0JiUq2sBlCflcl/jiVBPKZqCl/mJJiA1ZxYh9CX6iQRrUgB4VyTYpLpABZF1ckU9GEUhTzGPsoq7cCXYyHF0WqQWdAJ/24zCLozMZkrQY8wUKOyWoyVICIsDorMMMcBYoxFAIz3igJ5GPyO9kinUF556Bs7dUFAZ+PoIDlHWcScw1unjGAT8RHo4jiGjr93nlcnAUDIn5Byet2SK7YUK2o4EoGsWIvcZ4MbQVq2GXgF2OzUh1x3g/IC32oy5DSy1m0TF7wC7e7X9dEUqMvaraArgoTMLQTnXgPdOcZIuVHAPePewIDKLCuwAf7tlDxUUD4ykrGX1H47cu/fX+HIkW2voJwZVh/aVIWo889iDOwdHuEg2ZmicHFHt6YeunVjVPZR6KhkIT+EsRiY6AifkbGexjgKnC29MEVqH3XJNgcYBlJk4JRL8BJPlNuKsLYDSDd4yAp2BZlLFvPicAqjZH+E3p0Hu27hkl5qEpPBACnUFfi/NRJ9GRgOSzZ6A6RRYHQ54fx//7ZhfvfB0I+lPgHIBFmp9A/zx2iR/Fnsp4OpoMnDeqrn9atpIvuatW28EXv79NQA="
PRINT_LOGO_DATA_URI = "data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAU4AAACiCAYAAAAuoiQkAAAex0lEQVR4nO2dd7gU1fnHP1eKDQEFNAIq2BMMXgRsP02ssddYo8ResCKWmGJNokls1xqVWGPFRI3Yo+jPWGNMsGKMJkhMNIoGFYWoQP54Z7PDMOfMnKln730/z8Ozszvt3GX3u+9522k7s30YSnZOnfJy3UNQFKViFql7AIqiKK2GCqeiKIojKpyKoiiOqHAqiqI4osKpKIriiAqnoiiKIyqciqIojqhwKoqiOKLCqSiK4ogKp6IoiiMqnIqiKI50r3sAIXoDBwNLAL8CXq13OIqiKPH4ZHF+BJwPvA48C9wPjK51RIqiKDH4JJwNbgFGASsBvweuBfrVOSBFUZQwPgonwJ+B9YHHgP2Al4Ftax2RoihKgI/COQLoCXwIbAM8AiwH3A2cBrTVNzRFURQ/hXMgcA/QB/gU2AV4HhHM04EJQLe6BqcoiuKjcD4PbAE8iIjnh8DOwAfB/oOAa1DLU1GUmvBRON8OHtcF7kCsy2nA4aFjxgBnVzssRVEUwUfhnAt8HGxvCpwcbN8G3Bs67jvAbhWOS1EUBfBTOAEWDW1/F1g22D4RmB/adyUwuKpBKYqigJ/C2RuJqjdYEtg32H6FBa3OpYGfVDQuRVEUwE/hXC3mta1D29dE9u2DJMwriqJUgo/C2R7z2sjQ9n2IHzTMSaWNRlEUJYKPwrl1zGthn+enwJTI/p2B/iWNR1EUZQF8E87Fga1iXp8aef565HkPYM9SRqQoihLBN+HcB1gq5vXTIs/fjTkmTnAVRVEKxzfh3Cby/GPgABaMpAPMizl33VJGpCiKEsGnRsYAuwObAMOAfwG/Bf4dc9yAmNeWQ9KT4o5XFEUpDN+Ecx4wOfhnY3XD6wNR4VQUpWR8m6qnoTcw3LBv6SoHoihK16QVhXMnFqwsCqPt5hRFKZ1WE85FkNp1E+9XNRBFUbourSacBwBftux/r6qBKIrSdWkl4RwA/MyyfxYSiVcURSmVVhHONuBqYBnLMS9UNBZFUbo4rSKcxwLbJxzzZBUDURRFaQXh/DpwTorjotVFiqIopeC7cK4ITCQ5UX8m8Hjpo1EURcFv4ewF3EVz2QwbtwKflzscJQNn4V91mguDiW9z6AOL01wZoVU5HVnhoeXwVTi7AbcAa6c8PtoVXqmfFZA1olq53d86iH/dR74FfKPuQeRgGeAEJMWw5fDVGrgQ2C7lsc8Cz5Q4lqI5G7EW4riEhXuNtirHIp+vE4Abax5LVtZBLM41gD/XPJYox+Dv9zcNYxFrczzwcxZe1cFrfHzjxwNHOhyfJnDkE2OBPoZ9d9I5hLMPcGiw3Q5sATxU22iys07weBRwdJ0DibAp0q9hLvIjPLve4TizKM33cyiwK7L8d8vg21R9V9yE8A3g9pLGomTnUBZsSH1CXQPJSUM4D8D8Y1cHxwSP3UjvzvKJMUgbyAYn1jWQrPgknOsBN+A2pjNoMRO/C9CThf2CWwFfrWEseRgADAq2lwQOrHEsYYYAO4aej6hpHFlpA46PvDYa2LiGsWTGF+EcgkTQTb6/OKbSur6zzsyeNAUnTKtZnSMjz4/Gj+9LdBzrmA70lO2ANWNebymr04cPQm/gbtKlHYU5mfglNJT6aMP8BdibeEH1lfbI86HADjWMI0yc5dtqFqfp87E9EoRrCeoWzkba0TDH8x5ELFTFL7bEPCXvgb+pPXFELU5o+hbrYj+gb+S1ryLvbSswGviaYV/cFN5b6hbOs1h4gbYk5gLHlTAWJT9J061o0Mhn4iy5zajPV9tGfGS/J+6GR10kuWvG4D7zrIU6hXNv4KQM510CvFLwWJT8DEfSjmz0AQ6pYCx56QOsYthXV1rSN4j3DUJr+DmHAN9MOGYx/Er7MlKXcA4DJmQ4bwYSSVf8I61zv5EY7zM2v+EYoF9VAwlhc3O0gnAeR7qlbcYCS5Q8ltzUIZy9kGTXLDWqP6C5imUdH14lnsGkL61cEdijxLEUgU2IFqN6q3l17DXzvgeIlgYOSnlsP1qgDLMO4bwU+/IXJp6jaaWOAN4F/oA47G0NjpXyORa3AIXvQYAkITqSaq3moxEfp4l2/F6osFFemZbx1B9/sVL14HYHvp3hvPnIh3Ue8gG6DBn7SKSu/S3gF8BXihmm4kBvmuWVaVkH2LyEsRRF0tR3MFLlVgW9gf0TjlkCWK38oWSiJ+5+y5Wp7v3NRJXCOQi4POO5V9Fs5LEfsH5k/+LIVOAl4A5aJ8rYGTgU+XLHMd9ynq9W5xKYgzBhqkpNOghxbyURlz7lA2OALxn22T4fXhdMVCmcF5NtSv1vmksC9wF+ajm2DdgZWX/oSqB/hvsp6UnKzTwZc0nsNsBahY8oP8NJ9734P8oXq0WQBiNp8NHP2YZMu02cBnxq2LcesFHhIyqIqoRzR2CXjOeegkTTQRqfpsnzWgRx4L+KpD0p5bAnMm2N459AB/Ary/k+Wp0uYli21bkdMm1Ng4+R9W0xu88+BC4Arrec720ZZhXCuQSSe5mFV4Ergu21cPeV9ANuQpqHtGSnac+xfbAvAz5DxNPEPsDAIgdUAO0Ox+7Ngl1+imacw7E+Cqdtun0VsqT3hZZjdsDTMswqhPM4pBt4FsYDXwTbF5I9crgPsgqmyTpS3NkSmdbGMQdxlQA8DfzecFwP6i9jjOJicfYADi9pHMOQSqW09CG9dVoFo4BNDPvmIa47EOPoAcNx3pZhli2c/clubj8I3Bds74DbhyiO4cBTwKo5r6MItv/Xm4D3Qs9tVsXhpAt+VEEP3P2uY5HIcdFkqev3yc9pszZ/A0wLPe+wHOtlGWbZwvkdsjWAnUvTqdwde0DIhcHAo2S3gBVhOGJxmuiIPL8N8XnG4VMZZpaGGctR/LpKy5BtITZfputDgN0s+zsizx/AvDTJYqQPkFVGmcLZGzgs47m/BF4Otg8iW8K8iUHAJNTnmQebtTkZeDHy2udI4YOJcfhRhmmz2G627Cu6vvoQzL1pb6Ppvorii3COw+xWmwI8FnltPvZZyRG49eotnTKF80CydcL5HElTAJnClVGbvjb26YFiZhB2C8v0BbgS8X3G4UsZpk14fsiC08swo4ENCxpDd+xrbp2D+AXj8GGq3hc42LK/w/D69cBMwz7vyjDLEs42sjv9rwamB9tHUV7U8mDs000lHlt55V+RptRxzECyG0z4EAQwCeds4DXs2SFFBbl2xuxKehJZ1fUlw/7lqL9Z9OGYZ3PvYrbcP8He+Oc4PCrDLGsg6yEds12ZC/ws2O5F+dUDF+DRf0YLsBT28sqLsHflv8iybx1k9ca66IY5S2Aq8tn8BfIFj+ObFJO1YQsKNaz5qCskTHsBY8hKT+w/ID9HUtRMXIK5YGJVsueCF05ZorF7xvMmIlYLiG+z7A5Iw0juEag0ORRzsO9jZLZg40XgYcv+Osvs1sDczqwhVB8C1xiO6Y744vIwAnO1zN9pruhqE846Sy/3AZY37PscEU4b05GSaRM+zEqA8oTTFlGzcUHw2B17qVaR+JZH6CvdsVtDVyPimYQtCLAt9fUZsPkHXwhtX4y5xtoW1EnDOMu+S2gGhWzCWZefMynn8mbgXymuY/t8bICUutZOGcK5MuLsd+XZ4B+InyfLNbKwUYX3amX2xOx7m499Gh7mHuB1y/6qfjCj2Cy1sFC9BtxrOK4/8K2M918W2Muw71PETdDgTcw/UnVZnNtg/9HrSHmdx5EWkia8aP5RhnBukPG8sBmfd8rjyrYV368VsX1gJ9F0sSQRrhqJYwzm6V6ZtFv2RYMxHZZjs85gDsOcSH898EHo+fyYMTVYgXqafNuszd8Bf3K4ls3q3Alp7FwrZQhnlrSMOcCvg+2VqT5I4IX57zFbYBeWDsfrXQN8ZNjXg+rXnWnDbKm9D7wdee1hmnnGUYbj/vntid1YiLPmTcIJ1edzjsRe2dfheL1bgXcM+5I6LlVCGcKZpZnwb2h+kbJUTOTFFE1VBJu1+SLwiOP1kgJJY6m2DHMo5p6icQKVlLDtanXuhrln5QNIVD+Kzc9ZtXDarM03ke+3C58hTWJM7AcMcLxmoZQhnEMynBNuPTYd+UU3JUuXQZbUqa7CV4GtLPs7Ml73IsxBlr6kX6OmCGx+wRcMr9/AgtPnMDvi9pkaZ9nXYXjdlwDRStiLFy7GnGJk4wrgP4Z9tZdhliGcrgnrc4GHQs+vRaaGSyN5W3ciqQxlshTlNGroDNiszRlIQ48s/A27JXIc1ZVh2oTGJFCzabY8jOLSgHh9pPIojj9j7hxkEnSo1uIch7m88hOkfVwWbMnyINVVtZVhliGcrn/Mk8SXWs1BRHMX5FftQsq1Qr1fkrQGBmFvBH05+f5PbNPdlcie1uaKTWhslt2lmOvGDySduyEp4d1klX/Awr7XBquSrdzZlb7YG7Rci7mMMg0dln21lmH6UDVj+kUN8zbyy7YK7v4SJTu28so0Cc1JPAo8b9lfVeqJTThNQSCAfyBNN+LoS/LChIMw/zjMxN4dHcyi3kY103VbeeV87NkTaXge+YyYqK0M0wfhNDW5jeOfSI7nXkj36CL5sODrtTq9sJdXTsTcKs4Fm9U5kvIzLAZhDjRMIzmp3zb+pGV9x2J2R0zAXN7ZwBZZb084Ny9Jq1feh7lVnAu293dVRA8qpwzhnOl4vC3Z1cStSL7oPzKcG8d/sK+41xU5BHsvVdsH2oVo0+MoZVudaRPfTTyDdLmPY03MgbXFMLddnEu65WZs4xuV4vw87I192ZOiPh93If5wE7WUYZYhnKb8qzimY45MJvES0prf5Odx4e8FXKMz0R17pLfRpacI/oN92eiyyzDbLfvSCCdkS03aG/MqrHfQ7BBmwxYgak9xflbasP+gvQL8tqB7JRVMbEhxLf1SU0bU8g3SrUsN8flpLryOLKvxBLBojuuk+ZB2JfbAXoY6inxO/yhJn8PjsPd4zENeixMkne5c4lu6bY00EIlOW8dZrteR8r6vIMISZwB9BQnUzk55LRe2wr7EyGrIst5FkbTW2PHIj3lllGFx2pzpUYqYaj8HfD/nNWwBiq5I0vSnJzKNL+pfUjf+MZgTxPNiC6LYfIhhvsA8tW5j4dSkr2MuungOMQTSMAdz3X83JAe3DJLcJz0o9vORlJ2wCyLWlVGGcLpM4YryUV7geN8oWfysnZXN8WcJhgZJgYisDMDcuORz3IIbEzBbd/uzoL/YVlnU4XBPqL6CaATyGfGJNmRWUhllCOf/OxxblHDOA76b4/zJBY2jM+BNz8MIZZRhtlv2TcWt8OJ9ZK2sOHoheZ0glXU7G457B8lWcKFq4fSiO1EMB2D2GRdOGcL5HumtvzT9G9PyMObopo0XKSbA1BlYC2kP5iNLU3zCs82/mXaaHsYWJDoK+b4dgfl7dxn2DulxVFl6uSLFr+hZFIthX6upUMrK47R1cQ7zacH3TUoYjsOUwNwV8dXabDCeYgOaWUotbdiiySsjomOqtEnKLjBhG+dwin2/jiU5UFMnR1FRGWZZtcA3Aj/GnvxbBhMRJ73LD8KtJY2l1RiILH1g4lWKSWhOoi8SPIljCLAr7tNZE1lLLW10YF4EcALmQFhSPquJNxDfapxg9ESi67a0pbT0xV4Q8QbZrHRXlkR6WcTRH+mclOUHyImyhHM6Ukq5dcX3fx/5z0vbJu4hpKO3IgELU3kliLVUxBcwiZ5IwrMpufpEihHO3kjliYmsInAf8pmKa7Zryx7oyHi/eYila3I7jKSY/7dDsfuYD8ZeHlkUiyB/7xqG/eORpahtiwYWMoiyOD/FMUuXcF+XfK60yz10dnohdccmHqIa0QTx8dn+X0ZhtkhdaLfs+wjpI5kFl2VEGjxKvve3bD9nT+zNSJ6jGtEEEUSbtqyGdIkvlTKF8yFgSsIxZQjnqymPm4J5DfCuxsHYyyvPrWogAVdgr9M+sYB7FJH4buI63AoEOnLezya6RaxBlFReeV4B93Dhl9RbpluqcM4nOTE9rtIiL2+lPO4MtD4dkssrXwIerGYo/2MmCy5OFmU74Ms571FE4ruJWdjHH+avyJpNebCNdzj5vudJ5ZXTqT7AOhtp6Wei9DLMsrsj3Ys9r9OUfJyH91Mc8zjanq7BbkjvSxPnUc8PTAf2zuF5MwDKtDhBvthpOp9fTH5/nG28vci3uNmW2MsrL8Tck7RMLsXeC7bUDJEq2sodhfkDZHLw5iHpSz4fCYSotSnYrIm3yd7hPS/TaC7gF0eeMszFsX/2ihDOaUgjbhtJay+l5R2kG7+JPH5Om1vkQyRToA5mIC4RE7tgD/7logrhfAkpiYxjdeyR3CwkJdVfhNtSpZ2ZTbFbXhfjnpBdJDbfap4yzLWx5yMWIZyQ3FrNttqnK2VUEK2NOfUHRDSLLGJx5XzMBlCpZZhVrelyCtLFKPor3x3xwRRZK24ru3od+F6B9yqa/ZFWeXmYQrKl08BmbX5CBflwCTyLuFU2MuwfC5xFcsPfKO2Wff+guM4+vwP+SLxwZYm+23gRc9PnrMJpszY/p7iem1l5DfEP72jYfwBwGnZrPBNVCeccZGr1BAtbmOtSjXDORYSp6GqlItmvgGtcRzrhHIb0ujRxNcW2BsvKOZiFs1GGmabpb5iy/ZthOoivaJuEJI0XRdEpSStgL6+cSPpAbJmcg1k4F0d+XH9Y9E2rXDrjWeIbcWxc8H2WN7z+PdK36+oKjLfsm4vZvVI1d2MvUhiPexlgmRH1KLcC/4p5vWhrzTbupXFfAvtY7IbVOY7XK4vHsffGOJoSyjCrXnPofOD2yGubU2xpZlyayiT8+Y/2geWRGYCJO7AvV1Al87CL+FCkDDMtPbD3qSza4vwMad4RvUfRHbmSBN/F6uyDvbzyYfzqYWvzhQ/A/lnPRNXCOR+ZjoY/nMtSbJv/aOrEy8C+aBQ9zNHYg3K+/chch91P5ZIQPwwJLJkoo0LqcqSJR4MyfIOzsP/Yufg5D8W+vHDVCe9J3I5kMZg4noK1ro5VLmchPonw2kQ7F3Tt7iwonO8C21Nc5LIz0Avx+5h4AreVR6tgNgtbbWFGk97lYxOQuaSvPHPhXeDmYHsG0gSnDGyin9biTCqvfBm4P/WIquEL7NVXqyPB6cKoa3ngaUhgopHK4DLVsjGKZiOFWcibNa2ga3cWDkI63ZjwzdpscAn2hOeTUl7HJpx/SbhHHjqCx8tLvIdtup629HJP7BV9thSgOrkKe5lrEWW6/6PtzPYyFxBMZDPE+b84Ml3P6zc5GTgbmRZtSwWd3U+d4rLEEiDNYMv8wZqFfVq7LLCEZf90Su4sk4OBmKfZ80i36J7t75+D2yqtrjyAZAEUsR59HL2BZSz7p6W4Rn/sXZDeop5KoTQshz0QNK2oG1WVjmRiMvIL92vEgZtXOL+B5Jftgb/LYdS9oua7Nd8/D0UITp1//5iS7/8R+d1SMygh77Ei4rIXSqGuqXqYSUi99F7YnfZJ9AM2CK51VwHjUpSiaeUfLSVE3RZng7uQtI3tSL/sRpRtEetVRVNRlFLxRThBInW2npBJ3I0flS6KonRyfJiqh/kwx7kqmoqiVIJvwqkoiuI9KpyKoiiOqHAqiqI4osKpKIriiAqnoiiKIyqciqIojqhwKoqiOKLCqSiK4ogKp6IoiiMqnIqiKI6ocCqKojiiwqkoiuKICqeiKIojKpyKoiiOqHAqiqI4osKpKIriiAqnoiiKIz4tnaH4S29kTaevAysgy+j+HVnudjKyXlRWvgKsC3yArBe1HLAvMDq471vBfSblvE+YZYCdkMX9BiNLKr8J3Ac8Sr7lkddF/qZ3gKeBo5G/5QvgYeAXyPLVruPdMRjvIGQl138i7/1vgmuvCmwU/C2/SnHNbsj73Aa8DjzuMJ4RwNrB9q3AbIdzOwV1r6ve8mRYV72VaAOOBU7HvB7UdOAwZM2oLIwDLkCWhj4LuNJwr9eAXYE8b3g34PvAyZjX334FOAR4MuM9OpD3bCqyausqoX3PABuSXpgXA84Ajgm243gN2B9YA7gG+QEYkvL69wFbA+8BK5FeAP8AjAReoCmgXQq1OBUT3YDrgH2C53OQBfFeQr7EX0NEYEXgXuBwRPSyMhS4GXEfvYgI8afIF3Q7YHXgQcSay7I21WKIRbtl8Pwj5O+Zili2mwKjgus/Sv5lpr8cPP4RmIhYi6eSXjSXAe4B1g+eNyzyacgPy2aIaK0ejDfLWM9BhHMA8G3gihTnrIf8nwBckuGenQIVTsXEj2iK5v3AgcDbkWM2AW4ClgcuQwTvqYz36x08ngz8NLJvS0ScByLW4LkZrn8pTdG8BTgSEaMwOwC/RITpFmAd4NUM92pwO7AHMNfxvG7BuQ3RPA8R3U8jx22OTP2HIELvymTEehwFHA9MIFnYjwgeZwI3Zrhnp0CDQ0ocawAnBdsPIoISFU0QS2dz4BPky35+zvveycKiCfBb4NfB9pYx+5PYABF+kC/7t1hYNEH8qNsgPsPFDWNxYRzuoglwAOJPBhG0E1hYNEF8pqMQ90JWfhY8rgbsnHBsP2DPYPtaw5i6BCqcShwnIJ+N2cB+iJCYmIr4QEEspPXNhyZytmXfn4LHlTNc9/vB43uIxTTfcuxTiM8V5AdjFcuxNt5EAmiutAGnBNsPh8Zi4n1gF7IJNMgP0hvB9gkJxx4ILIq8f5dmvF+nQIVTibIITcvjJiQ6nMTliA8UJPqehc+QaaOJmcFjL8frLkHTSp2A+DaTuAgRhzbEAs1CmvctjhGI3xjgx9hFvsFryNQ+C/MQVwCIZb6R4bhFgLHB9gNIJL7Loj5OJcqqQP9gexgyJUvDJ0gAZr2M930Hu3+tkcJjioabGIlEt0GCWdemPG82IrrrkS0IMif5kFg2CB4/Ah5zOO9RYPeM97wWid4PAE4kPjVpGySAB13c2gQVTmVhvhTazjL1HpjxvmX5y5YPbW+S4fwvJR9SKI33bxpu0++sFi7Ij8RFwA8R98SaLBwUawSF/oYE6ro0KpxKlHC+4CTigyg2ZmS8r2tSeFrCn/GJuCdrTy9wLGlouM9sfuUyuAzJaFgSCUgdEtq3Mk2XxWXkKxDoFKhwKlHClssFwCN1DaQgwkJ+Bvki0FXQeP9XcjxvxeRDrHyApDYdC4xBAlSNsYxF/L1zgKtz3qdToMEhJcpfaFp/LtP0HiWMpQheDG2Pdjivrr+nESDrh/iY0/K1Au59PuIeWBQpFQWZgTRSuW7GfQbSKVHhVKLMphmUOIh0AtIdqcu+CHNpZl28jZRzgvjp2lKc0wsR3B8hAaIqeZqmlXyE7cAQg4HtC7j3dCTxv3HvXsBeSBUTwMUF3KNToMKpxNFIZF8FCRgkcRpSZXM0kk7jGx3B47rIVDSJc5EigO+QPY8zK3Nppgcdhjk9qEFPpEa9KAu5kRDfF/nhPDJ4/jTNXNoujwqnEsf9SMUQiHj8hGZKT5juiC/sB8Hzy5G0GN+4geaX/jwkCBL32V8U8eseFjz/MQtO9auiAyks6IbUq29lOG4o0qhjiwLv/QLNhi2nI5VJ0IXr0uPQ4JBiYh9EBIch4rkXku83JdjfjrQla1hkD5HOmquDL5DOSo8jbdnORiqibkCm8T0R/+e+yLQXJPn/zMpHKsxBpt6PIeO9H3l/70Wm08siJZk7IT7IecGxmxR0/3OR5h99g+fvArcVdO1OgQqnYmIGEnC4EvgmEuU9Lea4+Yhv8ySK65dZBtOQqfr1SH39mogPM8rniLCeSb1pN39FxPwaxOLcgnjL8i3gYCT/c5OC7v0w8BzNLkgT8Pv/tnJUOBUbHyBdd0YDewMbIxYQiOXzGPLFnprjHk8jaUJJCdxTguOyVuSANP/dAmkhtxtSpTMQsUjfRLoFXYWIbFbuR8pD81yjwduI5bcB0phkFNLo+SOkzHIS0rR4NtJQpEjuRIRzLuKCUUKocCppeDb4VwZPB/+SmELTTZCXRygvP/V+sjd1NvEUye36+gWPMwu4XxvyQwnS5/OtAq7ZqdDgkKL4xXWItdrheF4j5/O9AsawGdLQGTQFKRYVTkXxi1mIP3kzh3N60EyAf6mAMRwTPL6Cn1kStaPCqSh+0ZiSr4XkkqZhN5pT9Xty3n8o0ugDpC49TVu7LocKp6L4xR3IdLsNiWbH5c+G6U8zaf0vSIArD0cF9/4YyUBQYlDhVBS/+AQ4LtjeGEkNWt1w7AgkyDUYSZ06mHwpVEvSrEu/DhFPJQaNqiuKf9yITNNPQUouXwWeQKbx7yG14xsiwtqGpAztj1vj4zjG0Ex6vyzntTo1KpyK4ienImWiFyDBoo2Ir1ufgjTkyLq6aJhGR6TJ5MvN7fSocCqKv9yBJLk31rBfGelY9DGy5s9k4PcUE8BZD1gKWWDuvIRjuzwqnIriN18gApk36JPEM+Rvhtxl0OCQoiiKIyqciqIojqhwKoqiOKLCqSiK4ogKp6IoiiMqnIqiKI6ocCqKojiiwqkoiuKICqeiKIojKpyKoiiOqHAqiqI48l/W9yz/8fmpRAAAAABJRU5ErkJggg=="
OUTPUT_DEFAULT = "henry_hub_ote_interaktivni_analyza.html"

# 1 MWh = 3.4121416 MMBtu. Proto při převodu ceny platí:
# USD/MWh = USD/MMBtu * 3.4121416; USD/MMBtu = USD/MWh / 3.4121416.
MMBTU_PER_MWH = 3.4121416
UNITS = {
    "mmbtu": {"label": "USD/MMBtu", "short": "USD/MMBtu"},
    "mwh": {"label": "USD/MWh", "short": "USD/MWh"},
}
BASES = {
    "nominal": {"label": "nominální", "description": "nominální ceny"},
}

# Barevné oblasti jsou analytické intervaly, nikoli automaticky nalezené režimy.
MARKET_REGIMES = [
    ("1997-01-01", "2005-12-31", "Před břidlicovou<br>revolucí", "rgba(99,110,250,0.07)"),
    ("2006-01-01", "2014-12-31", "Nástup břidlicového<br>plynu", "rgba(0,204,150,0.07)"),
    ("2015-01-01", "2019-12-31", "Nízké ceny a růst LNG", "rgba(171,99,250,0.07)"),
    ("2020-01-01", "2021-12-31", "Covid a<br>zotavení", "rgba(239,85,59,0.07)"),
    ("2022-01-01", "2023-12-31", "Evropská energetická<br>krize", "rgba(255,161,90,0.09)"),
    ("2024-01-01", "2026-12-31", "Vysoká produkce a<br>nové exportní<br>kapacity", "rgba(25,211,243,0.07)"),
]

# Události: datum, název, kategorie, možný přenosový kanál, odkaz (odkaz se v tabulce nevypisuje).
EVENTS = [
    ("2000-03-10", "Dot-com bublina: vrchol Nasdaq", "Makroekonomie/finance", "Nepřímý vliv přes hospodářský cyklus a poptávku", "https://www.federalreservehistory.org/essays/stock-market-crash-of-2000"),
    ("2001-09-11", "Teroristické útoky v USA", "Geopolitika", "Nejistota, hospodářská aktivita a energetické trhy", "https://www.911memorial.org/911-faqs"),
    ("2001-10-07", "Zahájení amerických operací v Afghánistánu", "Válka", "Geopolitická riziková prémie; nepřímá vazba na americký plyn", "https://history.state.gov/milestones/2001-2009/afghanistan"),
    ("2003-03-20", "Zahájení války v Iráku", "Válka", "Ropa, geopolitická nejistota a energetický komplex", "https://history.state.gov/milestones/2001-2009/iraq"),
    ("2005-08-29", "Hurricane Katrina", "Počasí/infrastruktura", "Produkce a infrastruktura v Mexickém zálivu", "https://www.nhc.noaa.gov/data/tcr/AL122005_Katrina.pdf"),
    ("2005-09-21", "Hurricane Rita", "Počasí/infrastruktura", "Produkce a infrastruktura v Mexickém zálivu", "https://www.nhc.noaa.gov/data/tcr/AL182005_Rita.pdf"),
    ("2008-09-15", "Lehman Brothers vyhlásil bankrot", "Makroekonomie/finance", "Recese a pokles průmyslové energetické poptávky", "https://www.federalreservehistory.org/essays/lehman-brothers-files-for-bankruptcy"),
    ("2011-02-15", "Občanská válka v Libyi", "Válka", "Ropa a globální energetická riziková prémie", "https://www.britannica.com/event/Libya-Revolt-of-2011"),
    ("2014-02-23", "Ruská okupace Krymu", "Geopolitika", "Evropský plyn a geopolitická nejistota", "https://www.britannica.com/place/Crimea/History"),
    ("2014-11-27", "OPEC ponechal těžební cíl", "Ropa/těžba", "Cena ropy a ekonomika doprovodného plynu", "https://www.eia.gov/todayinenergy/detail.php?id=19451"),
    ("2016-02-24", "První zásilka LNG ze Sabine Pass", "LNG/infrastruktura", "Strukturální růst exportní poptávky po americkém plynu", "https://www.eia.gov/todayinenergy/detail.php?id=25072"),
    ("2017-08-25", "Hurricane Harvey", "Počasí/infrastruktura", "Produkce, poptávka a exportní infrastruktura", "https://www.nhc.noaa.gov/data/tcr/AL092017_Harvey.pdf"),
    ("2020-03-11", "WHO označila COVID-19 za pandemii", "Pandemie", "Propad ekonomické aktivity a energetické poptávky", "https://www.who.int/europe/emergencies/situations/covid-19"),
    ("2020-04-20", "WTI uzavřela se zápornou cenou", "Ropa/finance", "Ropa, vrty a budoucí produkce doprovodného plynu", "https://www.eia.gov/todayinenergy/detail.php?id=43495"),
    ("2021-02-13", "Winter Storm Uri", "Počasí/infrastruktura", "Propad produkce a skoková poptávka po vytápění", "https://www.eia.gov/todayinenergy/detail.php?id=46856"),
    ("2021-08-29", "Hurricane Ida", "Počasí/infrastruktura", "Omezení produkce a infrastruktury v Mexickém zálivu", "https://www.nhc.noaa.gov/data/tcr/AL092021_Ida.pdf"),
    ("2022-02-24", "Ruská invaze na Ukrajinu", "Válka", "Globální poptávka po LNG a evropská energetická krize", "https://www.un.org/press/en/2022/ga12407.doc.htm"),
    ("2022-06-08", "Požár a odstávka Freeport LNG", "LNG/infrastruktura", "Pokles domácí poptávky po plynu pro zkapalňování", "https://www.eia.gov/todayinenergy/detail.php?id=52859"),
    ("2023-10-07", "Útok Hamásu na Izrael", "Geopolitika", "Regionální geopolitické riziko a energetický komplex", "https://www.un.org/unispal/document/security-council-press-statement-07oct2023/"),
    ("2026-02-28", "Zahájení války USA a Izraele proti Íránu", "Válka", "Hormuzský průliv, globální energie a riziková prémie", "https://www.britannica.com/event/2026-Iran-war"),
]
ELECTIONS = [
    ("2000-11-07", "Prezidentské volby USA 2000"),
    ("2004-11-02", "Prezidentské volby USA 2004"),
    ("2008-11-04", "Prezidentské volby USA 2008"),
    ("2012-11-06", "Prezidentské volby USA 2012"),
    ("2016-11-08", "Prezidentské volby USA 2016"),
    ("2020-11-03", "Prezidentské volby USA 2020"),
    ("2024-11-05", "Prezidentské volby USA 2024"),
]
for event_date, event_name in ELECTIONS:
    EVENTS.append((event_date, event_name, "Volby", "Regulace energetiky, LNG a těžby; krátkodobý efekt nemusí být významný", "https://www.archives.gov/electoral-college/election-results"))
EVENTS.sort(key=lambda row: pd.Timestamp(row[0]))

OVERLAY_LABELS = {
    "raw": "Cena",
    "ma": "12M klouzavý průměr",
    "median": "12M klouzavý medián",
    "loess": "LOESS",
    "exp": "Exponenciální trend",
    "break": "Strukturální režimy",
}
OVERLAY_COLORS = {
    "raw": "#6c757d",
    "ma": "#1f77b4",
    "median": "#ff7f0e",
    "loess": "#2ca02c",
    "exp": "#9467bd",
    "break": "#d62728",
}
VOL_SPECS = [
    ("vol30", "30denní anualizovaná volatilita", "#d62728"),
    ("vol90", "90denní anualizovaná volatilita", "#ff7f0e"),
    ("vol365", "365denní anualizovaná volatilita", "#1f77b4"),
    ("robust_vol90", "90denní robustní volatilita MAD", "#2ca02c"),
    ("downside_vol90", "90denní volatilita poklesů", "#9467bd"),
    ("abs_change_pct", "Absolutní denní změna (%)", "#8c564b"),
    ("price_vol90", "90denní volatilita změny ceny", "#17becf"),
]


def _to_numeric(values: pd.Series) -> pd.Series:
    """Převede číselnou řadu včetně případných desetinných čárek na čísla."""
    if pd.api.types.is_numeric_dtype(values):
        return pd.to_numeric(values, errors="coerce")
    text = values.astype(str).str.strip().str.replace("\u00a0", "", regex=False).str.replace(" ", "", regex=False)
    # Pokud se v textu vyskytuje čárka bez tečky, bereme ji jako desetinný oddělovač.
    comma_decimal = text.str.contains(",", regex=False) & ~text.str.contains(".", regex=False)
    text.loc[comma_decimal] = text.loc[comma_decimal].str.replace(",", ".", regex=False)
    # Běžné tisícové čárky odstraníme, pokud je zároveň přítomna tečka.
    comma_thousands = text.str.contains(",", regex=False) & text.str.contains(".", regex=False)
    text.loc[comma_thousands] = text.loc[comma_thousands].str.replace(",", "", regex=False)
    return pd.to_numeric(text, errors="coerce")


EIA_PRICE_SOURCE_URL = "https://www.eia.gov/dnav/ng/hist/rngwhhdd.htm"
EIA_PRICE_XLS_URL = "https://www.eia.gov/dnav/ng/hist_xls/RNGWHHDd.xls"
_EIA_PRICE_COLUMN = "Henry Hub Natural Gas Spot Price (Dollars per Million Btu)"
_EIA_XLS_SIGNATURE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def _eia_price(value) -> float:
    if isinstance(value, bool):
        return math.nan
    if isinstance(value, Real):
        return float(value)
    text = "".join(str(value).split())
    if "," in text and "." not in text:
        text = text.replace(",", ".")
    elif "," in text and "." in text:
        text = text.replace(",", "")
    if not re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?", text):
        return math.nan
    return float(text)


def _eia_date(value):
    if pd.isna(value) or isinstance(value, bool):
        return pd.NaT
    try:
        if isinstance(value, (datetime, date, pd.Timestamp)):
            return pd.Timestamp(value).normalize()
        if isinstance(value, Real):
            converted = pd.to_datetime(float(value), unit="D", origin="1899-12-30", errors="coerce")
            return pd.NaT if pd.isna(converted) else converted.normalize()
        text = str(value).strip()
        for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%m/%d/%Y", "%m/%d/%y"):
            try:
                return pd.Timestamp(datetime.strptime(text, fmt)).normalize()
            except ValueError:
                pass
    except (ValueError, OverflowError, TypeError):
        pass
    return pd.NaT


def parse_henry_hub_xls(content: bytes) -> pd.DataFrame:
    """Read the named EIA data sheet and identify its two columns by their headers."""
    if not content.startswith(_EIA_XLS_SIGNATURE):
        prefix = content.lstrip()[:200].lower()
        if prefix.startswith((b"<!doctype", b"<html", b"<?xml")):
            raise ValueError("EIA vrátila HTML/XML místo souboru XLS; přístup mohl být odmítnut.")
        raise ValueError("Odpověď EIA není platný binární soubor XLS.")
    try:
        raw = pd.read_excel(BytesIO(content), sheet_name="Data 1", header=None, engine="xlrd")
    except ImportError as exc:
        raise RuntimeError("Pro načtení EIA XLS chybí knihovna xlrd. Nainstalujte: pip install 'xlrd>=2.0.1'.") from exc
    except ValueError as exc:
        if "Data 1" in str(exc):
            raise ValueError("V souboru EIA XLS chybí požadovaný list 'Data 1'.") from exc
        raise ValueError("Soubor EIA XLS se nepodařilo přečíst: " + str(exc)[:200]) from exc
    except Exception as exc:
        raise ValueError("Soubor EIA XLS je poškozený nebo jej nelze přečíst: " + str(exc)[:200]) from exc
    date_column = price_column = header_row = None
    required_price = " ".join(_EIA_PRICE_COLUMN.split()).casefold()
    for row_number, row in enumerate(raw.itertuples(index=False, name=None)):
        labels = [" ".join(str(cell).split()).casefold() for cell in row]
        if "date" in labels and required_price in labels:
            if labels.count("date") != 1 or labels.count(required_price) != 1:
                raise ValueError("List EIA 'Data 1' obsahuje nejednoznačné záhlaví.")
            header_row, date_column, price_column = row_number, labels.index("date"), labels.index(required_price)
            break
    if header_row is None:
        raise ValueError("V listu EIA 'Data 1' chybí záhlaví Date a '" + _EIA_PRICE_COLUMN + "'.")
    data = raw.iloc[header_row + 1:]
    result = pd.DataFrame({
        "date": data.iloc[:, date_column].map(_eia_date),
        "price_mmbtu": data.iloc[:, price_column].map(_eia_price),
    })
    valid = result["date"].notna() & result["price_mmbtu"].map(lambda value: math.isfinite(value) and value > 0)
    result = result.loc[valid].copy()
    result["date"] = pd.to_datetime(result["date"]).dt.normalize()
    result = result.groupby("date", as_index=False)["price_mmbtu"].mean()
    result["price_mwh"] = result["price_mmbtu"] * MMBTU_PER_MWH
    result = result.loc[result[["price_mmbtu", "price_mwh"]].apply(
        lambda column: column.map(math.isfinite)).all(axis=1)]
    if len(result) < 30:
        raise ValueError("Po vyčištění EIA Henry Hub zůstalo méně než 30 platných denních pozorování.")
    return result.sort_values("date").reset_index(drop=True)


def read_henry_hub() -> pd.DataFrame:
    """Download the current EIA workbook exactly once for this report generation."""
    request = Request(EIA_PRICE_XLS_URL, headers={
        "User-Agent": "Mozilla/5.0 (compatible; HenryHubOTEReport/1.0)",
        "Accept": "application/vnd.ms-excel, application/octet-stream", "Referer": EIA_PRICE_SOURCE_URL,
    })
    try:
        with urlopen(request, timeout=20) as response:
            content = response.read(20_000_001)
        if len(content) > 20_000_000:
            raise ValueError("Soubor překračuje limit 20 MB.")
    except (HTTPError, URLError, OSError, ValueError) as exc:
        raise RuntimeError("Denní ceny Henry Hub se nepodařilo stáhnout z EIA: " + str(exc)[:250]) from exc
    try:
        return parse_henry_hub_xls(content)
    except ValueError as exc:
        raise RuntimeError("Stažená data Henry Hub z EIA nejsou platná: " + str(exc)) from exc


OTE_ANNUAL_SOURCE_TEMPLATE = "https://www.ote-cr.cz/cs/statistika/rocni-zprava?date={year}-01-01"
_OTE_PRICE_HEADERS = (
    "Použitelná cena pro záporné vyrovnávací množství (Kč/MWh)",
    "Použitelná cena pro záporné denní vyrovnávací množství (Kč/MWh)",
)


def _ote_normalize(value) -> str:
    text = unicodedata.normalize("NFKD", " ".join(str(value).split()).casefold())
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"\s*([()/])\s*", r"\1", text)


class _OteAnnualLinks(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.links, self.headings, self.title = [], [], []
        self._anchor = None
        self._heading = self._in_title = False
        self._ignored = None

    def _finish_anchor(self):
        if self._anchor is not None:
            self.links.append(self._anchor)
            self._anchor = None

    def handle_starttag(self, tag, attrs):
        if self._ignored:
            return
        if tag in ("script", "style"):
            self._ignored = tag
        elif tag == "a":
            self._finish_anchor()
            attrs = dict(attrs)
            self._anchor = [attrs.get("href", ""), attrs.get("title", ""), []]
        elif tag in ("h1", "h2", "h3"):
            self._heading = True
        elif tag == "title":
            self._in_title = True

    def handle_data(self, data):
        if self._ignored:
            return
        if self._anchor is not None:
            self._anchor[2].append(data)
        if self._heading:
            self.headings.append(data)
        if self._in_title:
            self.title.append(data)

    def handle_endtag(self, tag):
        if self._ignored:
            if tag == self._ignored:
                self._ignored = None
            return
        if tag == "a":
            self._finish_anchor()
        elif tag in ("h1", "h2", "h3"):
            self._heading = False
        elif tag == "title":
            self._in_title = False

    def close(self):
        super().close()
        self._finish_anchor()


def _ote_version(text: str) -> int:
    values = re.findall(r"(?:^|[^a-z0-9])v(\d+)(?=$|[^a-z0-9])", text.casefold())
    return max(map(int, values)) if values else -1


def _ote_pick(candidates: dict[str, int], year: int, kind: str) -> str | None:
    if not candidates:
        return None
    # OTE updates V0 daily; V1/V2 have a reporting lag and may omit recent days.
    preferred = 0 if 0 in candidates.values() else max(candidates.values())
    if preferred != 0 and year >= pd.Timestamp.now(tz="Europe/Prague").year:
        raise ValueError(f"OTE pro aktuální rok {year} nenabízí průběžnou verzi V0 ({kind}).")
    selected = [name for name, version in candidates.items() if version == preferred]
    if len(selected) != 1:
        raise ValueError(f"OTE obsahuje více nejednoznačných souborů stejné preferované verze ({kind}).")
    return selected[0]


def _ote_url(url: str) -> str:
    parts = urlsplit(url)
    if (parts.scheme.casefold() != "https" or parts.hostname not in ("www.ote-cr.cz", "ote-cr.cz")
            or parts.username or parts.password or parts.port not in (None, 443)):
        raise ValueError("Odkaz roční zprávy musí směřovat přes HTTPS přímo na ote-cr.cz.")
    return urlunsplit(("https", parts.netloc, quote(parts.path, safe="/%:@!$&'()*+,;=-._~"),
                       quote(parts.query, safe="%=&?/:;+,$@!()*'-._~"), ""))


def find_ote_gas_zip_url(html: str, year: int, page_url: str) -> str | None:
    parser = _OteAnnualLinks()
    parser.feed(html)
    parser.close()
    candidates = {}
    for href, title, label in parser.links:
        combined = _ote_normalize(unquote(href) + " " + title + " " + " ".join(label))
        if (not urlsplit(href).path.casefold().endswith(".zip") or "plyn" not in combined
                or "elektr" in combined or not re.search(rf"(?<!\d){year}(?!\d)", combined)):
            continue
        url = _ote_url(urljoin(page_url, href))
        candidates[url] = max(candidates.get(url, -1), _ote_version(combined))
    return _ote_pick(candidates, year, "ZIP")


def _ote_check_zip(archive: ZipFile):
    members = archive.infolist()
    if len({item.filename for item in members}) != len(members):
        raise ValueError("Archiv OTE obsahuje duplicitní názvy souborů.")
    if len(members) > 200 or sum(item.file_size for item in members) > 100_000_000:
        raise ValueError("Archiv OTE překračuje limit počtu souborů nebo rozbalené velikosti 100 MB.")
    if any(item.file_size > 30_000_000 or item.flag_bits & 1 for item in members):
        raise ValueError("Archiv OTE obsahuje příliš velký nebo šifrovaný soubor.")
    return members


def _ote_date(value):
    if pd.isna(value) or isinstance(value, bool):
        return pd.NaT
    try:
        if isinstance(value, (date, datetime, pd.Timestamp)):
            return pd.Timestamp(value).normalize()
        if isinstance(value, Real):
            parsed = pd.to_datetime(float(value), unit="D", origin="1899-12-30", errors="coerce")
            return pd.NaT if pd.isna(parsed) else parsed.normalize()
        text = str(value).strip()
        for fmt in ("%d.%m.%Y", "%d. %m. %Y", "%Y-%m-%d", "%Y-%m-%d %H:%M:%S"):
            try:
                return pd.Timestamp(datetime.strptime(text, fmt)).normalize()
            except ValueError:
                pass
    except (ValueError, TypeError, OverflowError):
        pass
    return pd.NaT


def _ote_price(value) -> float:
    if isinstance(value, bool):
        return math.nan
    if isinstance(value, Real):
        return float(value)
    text = "".join(str(value).split())
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    else:
        text = text.replace(",", ".")
    if not re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?", text):
        return math.nan
    return float(text)


def _ote_header(raw: pd.DataFrame) -> tuple[int, int, int]:
    wanted_date = _ote_normalize("Plynárenský den")
    wanted_prices = {_ote_normalize(value) for value in _OTE_PRICE_HEADERS}
    for end in range(min(len(raw), 50)):
        for height in range(1, min(3, end + 1) + 1):
            block = raw.iloc[end - height + 1:end + 1]
            labels = [_ote_normalize(" ".join(str(value) for value in column if pd.notna(value)))
                      for _, column in block.items()]
            dates = [pos for pos, label in enumerate(labels) if label == wanted_date]
            prices = [pos for pos, label in enumerate(labels) if label in wanted_prices]
            if dates and prices:
                if len(dates) != 1 or len(prices) != 1:
                    raise ValueError("List OTE obsahuje nejednoznačné sloupce dne nebo záporné ceny.")
                return end + 1, dates[0], prices[0]
    raise ValueError("V listu 'Odchylky NC BAL' chybí Plynárenský den nebo přesný sloupec ceny pro záporné vyrovnávací množství (Kč/MWh).")


def parse_ote_annual_zip(content: bytes, year: int) -> pd.DataFrame:
    if len(content) > 30_000_000:
        raise ValueError("ZIP OTE překračuje limit 30 MB.")
    try:
        with ZipFile(BytesIO(content)) as archive:
            members = _ote_check_zip(archive)
            candidates = {}
            for item in members:
                name = PurePosixPath(item.filename.replace("\\", "/")).name
                normalized = _ote_normalize(name)
                if (name.casefold().endswith(".xlsx") and "plyn" in normalized and "elektr" not in normalized
                        and re.search(rf"(?<!\d){year}(?!\d)", normalized)):
                    candidates[item.filename] = _ote_version(normalized)
            selected = _ote_pick(candidates, year, "XLSX")
            if selected is None:
                raise ValueError(f"ZIP OTE neobsahuje plynový XLSX pro rok {year}.")
            with archive.open(selected) as member:
                workbook = member.read(30_000_001)
        if len(workbook) > 30_000_000:
            raise ValueError("XLSX OTE překračuje limit 30 MB.")
        with ZipFile(BytesIO(workbook)) as inner:
            _ote_check_zip(inner)
        raw = pd.read_excel(BytesIO(workbook), sheet_name="Odchylky NC BAL", header=None, engine="openpyxl")
    except (BadZipFile, OSError, RuntimeError, KeyError) as exc:
        raise ValueError("Archiv nebo sešit OTE se nepodařilo přečíst: " + str(exc)[:200]) from exc
    except ImportError as exc:
        raise RuntimeError("Pro načtení roční zprávy OTE je nutná knihovna openpyxl.") from exc
    first_data_row, date_column, price_column = _ote_header(raw)
    data = raw.iloc[first_data_row:]
    result = pd.DataFrame({"date": data.iloc[:, date_column].map(_ote_date),
                           "price_czk_mwh": data.iloc[:, price_column].map(_ote_price)})
    dated = result["date"].map(lambda value: pd.notna(value) and value.year == year)
    raw_prices = data.iloc[:, price_column]
    missing = raw_prices.map(lambda value: pd.isna(value) or str(value).strip().casefold() in {"", "nd", "n/a", "na", "nan", "none", "null", "-", "--", "—"})
    if (dated & ~missing & ~result["price_czk_mwh"].map(math.isfinite)).any():
        raise ValueError("Roční zpráva OTE obsahuje neplatnou nebo nekonečnou cenu pro plynárenský den.")
    result = result.loc[result["date"].notna() & result["price_czk_mwh"].map(math.isfinite)].copy()
    result["date"] = pd.to_datetime(result["date"]).dt.normalize()
    result = result.loc[result["date"].dt.year.eq(year)]
    if result.empty:
        # A newly published current-year template may have dates but no prices yet.
        known_dates = data.iloc[:, date_column].map(_ote_date).dropna()
        if year == pd.Timestamp.now(tz="Europe/Prague").year and any(day.year == year for day in known_dates):
            return pd.DataFrame({"date": pd.Series(dtype="datetime64[ns]"), "price_czk_mwh": pd.Series(dtype="float64")})
        raise ValueError(f"Roční zpráva OTE neobsahuje platné denní ceny pro rok {year}.")
    if (result.groupby("date")["price_czk_mwh"].nunique() > 1).any():
        raise ValueError("Roční zpráva OTE obsahuje rozdílné ceny pro stejný plynárenský den.")
    return result.drop_duplicates("date").sort_values("date").reset_index(drop=True)


def _ote_download(url: str, limit: int, accept: str) -> bytes:
    request = Request(_ote_url(url), headers={"User-Agent": "Mozilla/5.0 (compatible; HenryHubOTEReport/1.0)", "Accept": accept})
    try:
        with urlopen(request, timeout=20) as response:
            content = response.read(limit + 1)
        if len(content) > limit:
            raise ValueError("Odpověď překračuje povolenou velikost.")
        return content
    except (HTTPError, URLError, OSError, ValueError) as exc:
        raise RuntimeError("Data OTE se nepodařilo stáhnout: " + str(exc)[:250]) from exc


def fetch_ote_annual_prices(year: int) -> pd.DataFrame | None:
    page_url = OTE_ANNUAL_SOURCE_TEMPLATE.format(year=year)
    body = _ote_download(page_url, 2_000_000, "text/html")
    try:
        html = body.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError("Roční stránka OTE neobsahuje platné UTF-8 HTML.") from exc
    zip_url = find_ote_gas_zip_url(html, year, page_url)
    if zip_url is None:
        parser = _OteAnnualLinks()
        parser.feed(html)
        parser.close()
        heading = _ote_normalize(" ".join(parser.title + parser.headings))
        blocked = any(word in heading for word in ("access denied", "forbidden", "pristup odepren", "too many requests", "captcha"))
        if blocked or "rocni zprava" not in heading:
            raise ValueError("OTE nevrátilo rozpoznanou stránku roční zprávy; přístup mohl být odmítnut.")
        return None
    result = parse_ote_annual_zip(_ote_download(zip_url, 30_000_000, "application/zip"), year)
    result["source_url"] = zip_url
    return result

CNB_YEARLY_RATE_TEMPLATE = (
    "https://www.cnb.cz/cs/financni-trhy/devizovy-trh/kurzy-devizoveho-trhu/"
    "kurzy-devizoveho-trhu/rok.txt?rok={year}"
)
_CNB_MAX_TEXT_BYTES = 2_000_000
_CNB_MMBTU_PER_MWH = MMBTU_PER_MWH


def _cnb_empty_rates() -> pd.DataFrame:
    return pd.DataFrame({"date": pd.Series(dtype="datetime64[ns]"),
                         "czk_per_usd": pd.Series(dtype="float64")})


def _cnb_number(value, label: str, positive: bool = False) -> float:
    if isinstance(value, bool):
        raise ValueError(f"{label}: neplatná číselná hodnota.")
    if isinstance(value, Real):
        result = float(value)
        if not math.isfinite(result) or (positive and result <= 0):
            raise ValueError(f"{label}: hodnota musí být konečná" + (" a kladná." if positive else "."))
        return result
    text = str(value).strip().replace("\u00a0", "").replace(" ", "")
    if not re.fullmatch(r"[+-]?(?:\d+(?:[.,]\d+)?|[.,]\d+)", text):
        raise ValueError(f"{label}: neplatná číselná hodnota {text[:60]!r}.")
    result = float(text.replace(",", "."))
    if not math.isfinite(result) or (positive and result <= 0):
        raise ValueError(f"{label}: hodnota musí být konečná" + (" a kladná." if positive else "."))
    return result


def _cnb_price_date(value) -> pd.Timestamp:
    # A gas-day label denotes a calendar day; never shift it to a UTC day.
    if isinstance(value, (pd.Timestamp, datetime, date)) and not pd.isna(value):
        result = pd.Timestamp(value)
        if result.tzinfo is not None:
            result = result.tz_localize(None)
        return result.normalize()
    if isinstance(value, str):
        text = value.strip()
        for pattern, fmt in ((r"\d{4}-\d{2}-\d{2}", "%Y-%m-%d"),
                             (r"\d{2}\.\d{2}\.\d{4}", "%d.%m.%Y")):
            if re.fullmatch(pattern, text):
                try:
                    return pd.Timestamp(datetime.strptime(text, fmt))
                except ValueError:
                    break
    raise ValueError(f"Neplatné datum plynového dne / kurzu ČNB: {str(value)[:60]!r}.")


def _cnb_validate_rates(rates: pd.DataFrame) -> pd.DataFrame:
    if not isinstance(rates, pd.DataFrame) or not {"date", "czk_per_usd"}.issubset(rates.columns):
        raise ValueError("Tabulka kurzů ČNB musí mít sloupce date a czk_per_usd.")
    if rates.empty:
        return _cnb_empty_rates()
    result = pd.DataFrame({
        "date": [_cnb_price_date(value) for value in rates["date"]],
        "czk_per_usd": [_cnb_number(value, "Kurz ČNB CZK/USD", positive=True)
                        for value in rates["czk_per_usd"]],
    })
    if (result.groupby("date")["czk_per_usd"].nunique() > 1).any():
        raise ValueError("ČNB obsahuje rozdílné kurzy USD pro stejné datum.")
    return result.drop_duplicates("date").sort_values("date").reset_index(drop=True)


def parse_cnb_usd_rates(text: str) -> pd.DataFrame:
    """Parse annual CNB TXT, including a valid header-only (empty-year) table."""
    if not isinstance(text, str) or len(text) > _CNB_MAX_TEXT_BYTES:
        raise ValueError("ČNB nevrátila platný nebo přiměřeně velký textový soubor.")
    lines = [line.strip() for line in text.lstrip("\ufeff\r\n \t").splitlines() if line.strip()]
    if not lines:
        raise ValueError("Prázdná odpověď ČNB neobsahuje hlavičku roční tabulky.")
    header = [cell.strip() for cell in lines[0].split("|")]
    if not header or header[0].casefold() not in {"datum", "date"}:
        raise ValueError("ČNB nevrátila roční tabulku kurzů (očekávám hlavičku Datum|…|USD).")
    usd_columns = []
    for index, name in enumerate(header[1:], start=1):
        match = re.fullmatch(r"(\d+)\s+USD", name, flags=re.IGNORECASE)
        if match:
            amount = float(match.group(1))
            if not math.isfinite(amount) or amount <= 0:
                raise ValueError("ČNB má neplatné množství USD v hlavičce.")
            usd_columns.append((index, amount))
    if len(usd_columns) != 1:
        raise ValueError("ČNB musí mít právě jeden sloupec s množstvím a kódem USD.")
    usd_column, amount = usd_columns[0]
    records = []
    for line_number, line in enumerate(lines[1:], start=2):
        cells = [cell.strip() for cell in line.split("|")]
        if len(cells) != len(header):
            raise ValueError(f"ČNB: neúplný nebo neplatný řádek {line_number}.")
        if not re.fullmatch(r"\d{2}\.\d{2}\.\d{4}", cells[0]):
            raise ValueError(f"ČNB: neplatné datum na řádku {line_number}.")
        rate = _cnb_number(cells[usd_column], "Kurz ČNB USD", positive=True) / amount
        records.append({"date": _cnb_price_date(cells[0]), "czk_per_usd": rate})
    return _cnb_validate_rates(pd.DataFrame(records, columns=["date", "czk_per_usd"]))


def fetch_cnb_usd_rates(year: int) -> pd.DataFrame:
    """Download a single annual table; never interpret an error page as no fixings."""
    if isinstance(year, bool) or not isinstance(year, int) or not 1991 <= year <= 9999:
        raise ValueError("ČNB: rok musí být celé číslo od roku 1991.")
    url = CNB_YEARLY_RATE_TEMPLATE.format(year=year)
    request = Request(url, headers={"Accept": "text/plain",
                                   "User-Agent": "HenryHubOTEReport/1.0"})
    try:
        with urlopen(request, timeout=20) as response:
            content_type = response.headers.get("Content-Type", "").casefold()
            if "html" in content_type:
                raise ValueError("ČNB vrátila HTML stránku místo roční tabulky kurzů.")
            body = response.read(_CNB_MAX_TEXT_BYTES + 1)
    except (HTTPError, URLError, TimeoutError, OSError) as exc:
        raise RuntimeError(f"Kurzy ČNB pro rok {year} se nepodařilo stáhnout: {exc}") from exc
    if len(body) > _CNB_MAX_TEXT_BYTES:
        raise ValueError(f"Odpověď ČNB pro rok {year} překročila velikostní limit.")
    try:
        text = body.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError(f"ČNB pro rok {year} nevrátila platný text UTF-8.") from exc
    result = parse_cnb_usd_rates(text)
    if not result.empty and not result["date"].dt.year.eq(year).all():
        raise ValueError(f"ČNB pro rok {year} vrátila kurzy z jiného roku.")
    result.attrs["source_url"] = url
    return result


def convert_ote_czk_to_usd(prices: pd.DataFrame) -> pd.DataFrame:
    """Use the same-day or latest earlier CNB fixing, preserving input row order.

    Downloads each required year at most once in this call. When the first gas
    day of a year precedes that year's first fixing, also downloads the previous
    year. Missing coverage fails explicitly rather than using a future fixing.
    """
    if not isinstance(prices, pd.DataFrame) or not {"date", "price_czk_mwh"}.issubset(prices.columns):
        raise ValueError("Ceny OTE musí mít sloupce date a price_czk_mwh.")
    result = prices.copy()
    if result.empty:
        result["date"] = pd.Series(index=result.index, dtype="datetime64[ns]")
        for column in ("price_mwh", "price_mmbtu", "czk_per_usd"):
            result[column] = pd.Series(index=result.index, dtype="float64")
        result["cnb_rate_date"] = pd.Series(index=result.index, dtype="datetime64[ns]")
        return result
    dates = pd.DatetimeIndex([_cnb_price_date(value) for value in result["date"]])
    czk_prices = [_cnb_number(value, "Cena OTE v CZK/MWh") for value in result["price_czk_mwh"]]
    lookups = pd.DataFrame({"date": dates, "_position": range(len(result))})
    yearly = {}

    def load_year(year: int) -> pd.DataFrame:
        if year not in yearly:
            yearly[year] = _cnb_validate_rates(fetch_cnb_usd_rates(year))
            if not yearly[year].empty and not yearly[year]["date"].dt.year.eq(year).all():
                raise ValueError(f"ČNB pro rok {year} vrátila kurzy z jiného roku.")
        return yearly[year]

    for year in sorted(set(dates.year)):
        year = int(year)
        first_day = dates[dates.year == year].min()
        rates = load_year(year)
        if not rates["date"].le(first_day).any():
            if year <= 1991:
                raise ValueError(f"Pro {first_day:%d.%m.%Y} není k dispozici předchozí kurz ČNB USD.")
            previous = load_year(year - 1)
            if not previous["date"].le(first_day).any():
                raise ValueError(f"Pro {first_day:%d.%m.%Y} chybí stejný nebo předchozí kurz ČNB USD.")
    all_rates = _cnb_validate_rates(pd.concat(list(yearly.values()), ignore_index=True))
    aligned = pd.merge_asof(lookups.sort_values("date"),
                            all_rates.rename(columns={"date": "cnb_rate_date"}),
                            left_on="date", right_on="cnb_rate_date", direction="backward",
                            allow_exact_matches=True).sort_values("_position")
    if aligned["czk_per_usd"].isna().any():
        raise ValueError("Pro některý plynový den OTE není k dispozici stejný nebo předchozí kurz ČNB USD.")
    result["date"] = dates.to_numpy()
    result["czk_per_usd"] = aligned["czk_per_usd"].to_numpy()
    result["cnb_rate_date"] = aligned["cnb_rate_date"].to_numpy()
    result["price_mwh"] = pd.Series(czk_prices, dtype="float64").to_numpy() / result["czk_per_usd"].to_numpy()
    result["price_mmbtu"] = result["price_mwh"] / _CNB_MMBTU_PER_MWH
    return result

_OTE_USD_MWH = "Spotová cena plynu OTE (USD/MWh)"
_OTE_USD_MMBTU = "Spotová cena plynu OTE (USD/MMBtu)"
_OTE_CZK_MWH = "Spotová cena plynu OTE (Kč/MWh)"
_OTE_CNB_RATE = "Kurz ČNB (Kč/USD)"
_OTE_CNB_DATE = "Datum kurzu ČNB"
_OTE_SOURCE = "Zdroj nových dat OTE"
_OTE_MMBTU_PER_MWH = MMBTU_PER_MWH
_OTE_XML_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def _ote_workbook_number(value):
    """Finite stored price, including zero and negative prices; None if missing."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, str):
        value = value.strip().replace("\u00a0", "").replace(" ", "")
        if "," in value and "." not in value:
            value = value.replace(",", ".")
        elif "," in value:
            value = value.replace(",", "")
    try:
        result = float(value)
    except (ValueError, TypeError, OverflowError):
        return None
    return result if math.isfinite(result) else None


def _ote_workbook_date(value, epoch=None):
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, datetime):
        return None if pd.isna(value) else value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        if epoch is None or not math.isfinite(float(value)):
            return None
        try:
            parsed = from_excel(value, epoch)
            return parsed.date() if isinstance(parsed, datetime) else parsed if isinstance(parsed, date) else None
        except (ValueError, OverflowError):
            return None
    text = str(value).strip()
    if not text:
        return None
    try:
        return datetime.fromisoformat(text).date()
    except ValueError:
        pass
    for pattern in ("%d.%m.%Y", "%d. %m. %Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            pass
    return None


def _ote_sheet_xml_paths(archive):
    ns = {"m": _OTE_XML_NS}
    rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    doc_rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    rels = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    targets = {node.get("Id"): node.get("Target") for node in rels.findall(f"{{{rel_ns}}}Relationship")}
    book = ET.fromstring(archive.read("xl/workbook.xml"))
    result = {}
    for sheet in book.findall("m:sheets/m:sheet", ns):
        target = targets.get(sheet.get(f"{{{doc_rel_ns}}}id"))
        if target:
            result[sheet.get("name")] = target.lstrip("/") if target.startswith("/") else posixpath.normpath("xl/" + target)
    return result


def _ote_restore_formula_caches(original_bytes, staged_path, changed_cells):
    """openpyxl drops cached formula results; retain caches of untouched formulas."""
    replacements = {}
    ns = "{" + _OTE_XML_NS + "}"
    with ZipFile(BytesIO(original_bytes)) as original, ZipFile(staged_path) as staged:
        old_paths, new_paths = _ote_sheet_xml_paths(original), _ote_sheet_xml_paths(staged)
        for name, old_path in old_paths.items():
            new_path = new_paths.get(name)
            if not new_path:
                continue
            old_root = ET.fromstring(original.read(old_path))
            cached = {cell.get("r"): cell for cell in old_root.iter(ns + "c")
                      if cell.find(ns + "f") is not None and cell.find(ns + "v") is not None
                      and cell.find(ns + "v").text is not None}
            if not cached:
                continue
            new_root = ET.fromstring(staged.read(new_path))
            modified = False
            for cell in new_root.iter(ns + "c"):
                coord = cell.get("r")
                old_cell = cached.get(coord)
                if old_cell is None or (name, coord) in changed_cells or cell.find(ns + "f") is None:
                    continue
                # The updater never edits formula definitions. Shared formulas may
                # be expanded by openpyxl, so match their unchanged cell identity.
                for existing in list(cell):
                    if existing.tag in (ns + "v", ns + "is"):
                        cell.remove(existing)
                cell.append(copy(old_cell.find(ns + "v")))
                if "t" in old_cell.attrib:
                    cell.set("t", old_cell.get("t"))
                else:
                    cell.attrib.pop("t", None)
                modified = True
            if modified:
                replacements[new_path] = ET.tostring(new_root, encoding="utf-8")
        if replacements:
            stream = BytesIO()
            with ZipFile(stream, "w") as rebuilt:
                for info in staged.infolist():
                    rebuilt.writestr(info, replacements.get(info.filename, staged.read(info.filename)))
    if replacements:
        staged_path.write_bytes(stream.getvalue())


def _ote_workbook_new_data(annual_fetcher, cutoff, today, notes):
    pieces = []
    for year in range((cutoff + timedelta(days=1)).year, today.year + 1):
        frame = annual_fetcher(year)
        if frame is None:
            if year == today.year:
                notes.append(f"Roční soubor OTE pro rok {year} zatím nezveřejněno; doplněna jsou pouze dostupná data.")
                continue
            raise ValueError(f"Roční soubor OTE pro potřebný rok {year} není dostupný.")
        if not isinstance(frame, pd.DataFrame) or not {"date", "price_czk_mwh", "source_url"}.issubset(frame.columns):
            raise ValueError(f"Roční data OTE pro rok {year} nemají sloupce date, price_czk_mwh a source_url.")
        for record in frame.to_dict("records"):
            day = _ote_workbook_date(record["date"])
            if day is None:
                raise ValueError(f"Roční data OTE pro rok {year} obsahují neplatné datum.")
            if not cutoff < day <= today:
                continue
            if day.year != year:
                raise ValueError(f"Roční data OTE pro rok {year} obsahují datum z jiného roku: {day}.")
            price = _ote_workbook_number(record["price_czk_mwh"])
            source = record["source_url"]
            if price is None or not isinstance(source, str) or not source.strip():
                raise ValueError(f"OTE nevrátilo platnou cenu nebo zdroj pro {day}.")
            pieces.append({"date": day, "price_czk_mwh": price, "source_url": source.strip()})
    if not pieces:
        return pd.DataFrame(columns=["date", "price_czk_mwh", "source_url"])
    data = pd.DataFrame(pieces)
    if (data.groupby("date")["price_czk_mwh"].nunique() > 1).any():
        raise ValueError("OTE vrátilo rozdílné ceny pro stejný nový den.")
    return data.drop_duplicates("date").sort_values("date").reset_index(drop=True)


def _ote_workbook_converted(currency_converter, source):
    converted = currency_converter(source.copy(deep=True))
    required = {"date", "price_mwh", "price_mmbtu", "czk_per_usd", "cnb_rate_date"}
    if not isinstance(converted, pd.DataFrame) or not required.issubset(converted.columns):
        raise ValueError("Převod OTE do USD nevrátil všechny požadované cenové a kurzové sloupce.")
    by_date = {row["date"]: row for row in source.to_dict("records")}
    result = []
    for row in converted.to_dict("records"):
        day = _ote_workbook_date(row["date"])
        if day not in by_date:
            raise ValueError("Převod OTE do USD vrátil nečekané datum.")
        values = {key: _ote_workbook_number(row[key]) for key in ("price_mwh", "price_mmbtu", "czk_per_usd")}
        rate_date = _ote_workbook_date(row["cnb_rate_date"])
        if any(value is None for value in values.values()) or values["czk_per_usd"] <= 0 or rate_date is None or rate_date > day:
            raise ValueError(f"Chybí platná cena USD nebo kurz ČNB pro {day}.")
        original = by_date[day]
        if not math.isclose(values["price_mwh"], original["price_czk_mwh"] / values["czk_per_usd"], rel_tol=1e-7, abs_tol=1e-7):
            raise ValueError(f"Cena OTE v USD/MWh neodpovídá ceně Kč/MWh a kurzu ČNB pro {day}.")
        if not math.isclose(values["price_mmbtu"], values["price_mwh"] / _OTE_MMBTU_PER_MWH, rel_tol=1e-7, abs_tol=1e-7):
            raise ValueError(f"Cena OTE v USD/MMBtu neodpovídá převodu jednotek pro {day}.")
        result.append({**original, **values, "cnb_rate_date": rate_date})
    if len(result) != len(by_date) or len({row["date"] for row in result}) != len(by_date):
        raise ValueError("Převod OTE do USD nevrátil právě jeden záznam pro každý nový den.")
    return sorted(result, key=lambda row: row["date"])


def update_ote_workbook(path: Path, *, today: date | None = None, annual_fetcher=None, currency_converter=None) -> dict:
    """Extend original OTE XLSX; leave history unchanged and back it up on success.

    Injected fetcher(year) supplies date/price_czk_mwh/source_url. The converter
    preserves those rows and adds USD price_mwh/price_mmbtu plus ČNB metadata.
    """
    path = Path(path)
    today = _ote_workbook_date(today) if today is not None else pd.Timestamp.now(tz="Europe/Prague").date()
    if today is None:
        raise ValueError("Neplatné datum konce aktualizace OTE.")
    original_bytes = path.read_bytes()
    wb = load_workbook(BytesIO(original_bytes), data_only=False)
    cached = load_workbook(BytesIO(original_bytes), read_only=True, data_only=True)
    try:
        ws, values_ws = wb.worksheets[0], cached.worksheets[0]
        headers = {str(cell.value).strip().casefold(): cell.column for cell in ws[1] if cell.value is not None}
        date_col = headers.get("date")
        price_cols = [headers[name.casefold()] for name in (_OTE_USD_MWH, _OTE_USD_MMBTU) if name.casefold() in headers]
        if date_col is None or not price_cols:
            raise ValueError("První list OTE musí obsahovat Date a alespoň jeden cenový sloupec USD/MWh nebo USD/MMBtu.")
        rows_by_date, valid_rows = {}, []
        for row_index, cells in enumerate(values_ws.iter_rows(min_row=2, max_col=max(date_col, *price_cols)), 2):
            day = _ote_workbook_date(cells[date_col - 1].value, wb.epoch)
            has_price = any(_ote_workbook_number(cells[column - 1].value) is not None for column in price_cols)
            if has_price and day is None:
                raise ValueError(f"Řádek {row_index} OTE má cenu, ale neplatné datum Date.")
            if day is not None:
                rows_by_date.setdefault(day, []).append(row_index)
                if has_price:
                    valid_rows.append((day, row_index))
        last_date = max((day for day, _ in valid_rows), default=None)
        cutoff = last_date or ((min(rows_by_date) if rows_by_date else date(today.year, 1, 1)) - timedelta(days=1))
    finally:
        cached.close()
        wb.close()
    status = {"added_rows": 0, "updated_rows": 0, "last_date": last_date.isoformat() if last_date else None,
              "backup_path": None, "notes": []}
    if cutoff >= today:
        return status
    if annual_fetcher is None:
        annual_fetcher = fetch_ote_annual_prices
    incoming = _ote_workbook_new_data(annual_fetcher, cutoff, today, status["notes"])
    if incoming.empty:
        return status
    if currency_converter is None:
        currency_converter = convert_ote_czk_to_usd
    converted = _ote_workbook_converted(currency_converter, incoming)
    # Confirm every reusable row is actually empty before changing any cells.
    placements = []
    for record in converted:
        day = record["date"]
        matching = rows_by_date.get(day, [])
        if matching:
            target = matching[0]
            for column in price_cols:
                cell = ws.cell(target, column)
                if cell.data_type == "f" or (cell.value is not None and str(cell.value).strip().upper() not in {"", "-", "—", "N/A", "ND", "NAN"}):
                    raise ValueError(f"Řádek OTE pro {day} obsahuje nevyhodnocenou cenu nebo vzorec; nebyl přepsán.")
            placements.append((record, target))
        else:
            placements.append((record, None))
    all_headers = (_OTE_USD_MWH, _OTE_USD_MMBTU, _OTE_CZK_MWH, _OTE_CNB_RATE, _OTE_CNB_DATE, _OTE_SOURCE)
    for name in all_headers:
        if name.casefold() not in headers:
            column = ws.max_column + 1
            previous = ws.cell(1, column - 1)
            cell = ws.cell(1, column, name)
            if previous.has_style:
                cell._style = copy(previous._style)
            headers[name.casefold()] = column
    changed_cells = set()
    style_row = max(valid_rows, key=lambda item: (item[0], item[1]))[1] if valid_rows else None
    for record, target in placements:
        is_new = target is None
        target = ws.max_row + 1 if is_new else target
        if is_new:
            ws.cell(target, date_col, record["date"])
            changed_cells.add((ws.title, ws.cell(target, date_col).coordinate))
        values = {_OTE_USD_MWH: record["price_mwh"], _OTE_USD_MMBTU: record["price_mmbtu"],
                  _OTE_CZK_MWH: record["price_czk_mwh"], _OTE_CNB_RATE: record["czk_per_usd"],
                  _OTE_CNB_DATE: record["cnb_rate_date"], _OTE_SOURCE: record["source_url"]}
        for name, value in values.items():
            column = headers[name.casefold()]
            cell = ws.cell(target, column, value)
            if not cell.has_style and style_row is not None and ws.cell(style_row, column).has_style:
                cell._style = copy(ws.cell(style_row, column)._style)
            if name == _OTE_CNB_DATE:
                cell.number_format = "yyyy-mm-dd"
            elif name == _OTE_SOURCE:
                cell.data_type = "s"
            changed_cells.add((ws.title, cell.coordinate))
        if is_new:
            date_cell = ws.cell(target, date_col)
            if style_row is not None:
                date_cell._style = copy(ws.cell(style_row, date_col)._style)
            if not is_date_format(date_cell.number_format):
                date_cell.number_format = "yyyy-mm-dd"
        status["added_rows" if is_new else "updated_rows"] += 1
    staged_path = backup_path = None
    committed = False
    try:
        with tempfile.NamedTemporaryFile(prefix=".ote-update-", suffix=path.suffix, dir=path.parent, delete=False) as temporary:
            staged_path = Path(temporary.name)
        wb.save(staged_path)
        wb.close()
        _ote_restore_formula_caches(original_bytes, staged_path, changed_cells)
        shutil.copymode(path, staged_path)
        if path.read_bytes() != original_bytes:
            raise RuntimeError("Soubor OTE se během aktualizace změnil; nebyl přepsán.")
        with tempfile.NamedTemporaryFile(prefix=path.stem + ".backup-", suffix=path.suffix, dir=path.parent, delete=False) as temporary:
            backup_path = Path(temporary.name)
        shutil.copy2(path, backup_path)
        os.replace(staged_path, path)
        committed = True
    finally:
        wb.close()
        if staged_path is not None:
            staged_path.unlink(missing_ok=True)
        if not committed and backup_path is not None:
            backup_path.unlink(missing_ok=True)
    status["last_date"] = converted[-1]["date"].isoformat()
    status["backup_path"] = str(backup_path)
    return status


def read_ote(path: Path) -> pd.DataFrame:
    """Načte OTE z uvedených sloupců; chybějící cenu doplní převodem druhé jednotky."""
    if not path.exists():
        raise FileNotFoundError(f"Soubor OTE nebyl nalezen: {path.resolve()}")
    raw = pd.read_excel(path, sheet_name=0, engine="openpyxl")
    # Názvy sloupců vyhledáváme bez ohledu na velikost písmen a okrajové mezery.
    lookup = {str(col).strip().casefold(): col for col in raw.columns}
    required_date = "date"
    required_mwh = "spotová cena plynu ote (usd/mwh)"
    required_mmbtu = "spotová cena plynu ote (usd/mmbtu)"
    missing = [name for name in [required_date, required_mwh, required_mmbtu] if name not in lookup]
    if missing:
        # Cenový sloupec může chybět, protože ho lze dopočítat z druhé jednotky.
        if required_date not in lookup or (required_mwh not in lookup and required_mmbtu not in lookup):
            cols = ", ".join(str(c) for c in raw.columns)
            raise ValueError(
                "V souboru OTE nebyly nalezeny potřebné sloupce Date a alespoň jeden cenový sloupec. "
                f"Nalezené sloupce: {cols}"
            )
    date_col = lookup[required_date]
    mwh_col = lookup.get(required_mwh)
    mmbtu_col = lookup.get(required_mmbtu)
    result = pd.DataFrame({"date": pd.to_datetime(raw[date_col], errors="coerce")})
    result["price_mwh"] = _to_numeric(raw[mwh_col]) if mwh_col is not None else np.nan
    result["price_mmbtu"] = _to_numeric(raw[mmbtu_col]) if mmbtu_col is not None else np.nan
    # Když je v jedné jednotce cena prázdná, dopočítáme ji z dostupné druhé ceny.
    mask_mwh = result["price_mwh"].isna() & result["price_mmbtu"].notna()
    result.loc[mask_mwh, "price_mwh"] = result.loc[mask_mwh, "price_mmbtu"] * MMBTU_PER_MWH
    mask_mmbtu = result["price_mmbtu"].isna() & result["price_mwh"].notna()
    result.loc[mask_mmbtu, "price_mmbtu"] = result.loc[mask_mmbtu, "price_mwh"] / MMBTU_PER_MWH
    result = result.dropna(subset=["date", "price_mwh", "price_mmbtu"])
    result = result[(result["price_mwh"] > 0) & (result["price_mmbtu"] > 0)].copy()
    if result.empty:
        raise ValueError("V souboru OTE nezůstala žádná platná cenová pozorování.")
    # Porovnání se provádí na denní frekvenci. Hodinová či vícenásobná data OTE zprůměrujeme za den.
    result["date"] = result["date"].dt.normalize()
    result = result.groupby("date", as_index=False).agg(
        price_mwh=("price_mwh", "mean"),
        price_mmbtu=("price_mmbtu", "mean"),
    )
    if len(result) < 30:
        raise ValueError("Po vyčištění zůstalo v souboru OTE příliš málo platných denních pozorování.")
    return result.sort_values("date").reset_index(drop=True)


FED_EXCHANGE_RATE_SOURCE = "https://www.federalreserve.gov/RELEASES/H10/hist/dat00_eu.htm"
_FED_RATE_ANCHOR = pd.Timestamp("2000-01-03")
_FED_MONTHS = {name: number for number, name in enumerate(
    ("JAN", "FEB", "MAR", "APR", "MAY", "JUN", "JUL", "AUG", "SEP", "OCT", "NOV", "DEC"), 1)}


class _FedRateTables(HTMLParser):
    """Collect tables while honoring implicit cell/row endings in ordinary HTML."""
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.tables = []
        self._stack = []
        self._ignored_tag = None
        self.page_title = ""
        self._title_parts = []
        self._in_title = False

    @staticmethod
    def _finish_cell(table):
        if table["cell"] is not None:
            if table["row"] is None:
                table["row"] = []
            table["row"].append(" ".join("".join(table["cell"]).split()))
            table["cell"] = None

    @classmethod
    def _finish_row(cls, table):
        cls._finish_cell(table)
        if table["row"] is not None:
            table["rows"].append(table["row"])
            table["row"] = None

    def _finish_table(self):
        table = self._stack.pop()
        self._finish_row(table)
        self.tables.append(table["rows"])

    def handle_starttag(self, tag, attrs):
        tag = tag.lower()
        if self._ignored_tag is not None:
            return
        if tag == "title":
            self._in_title = True
            if self._title_parts:
                self._title_parts.append(" ")
            return
        if tag in ("script", "style"):
            self._ignored_tag = tag
            return
        if tag == "table":
            # A nested table gets its own rows; its contents never enter the outer cell.
            self._stack.append({"rows": [], "row": None, "cell": None})
        elif self._stack:
            table = self._stack[-1]
            if tag == "tr":
                self._finish_row(table)
                table["row"] = []
            elif tag in ("td", "th"):
                self._finish_cell(table)
                if table["row"] is None:
                    table["row"] = []
                table["cell"] = []
            elif tag in ("thead", "tbody", "tfoot"):
                self._finish_row(table)
            elif tag == "br" and table["cell"] is not None:
                table["cell"].append(" ")

    def handle_data(self, data):
        if self._ignored_tag is None and self._in_title:
            self._title_parts.append(data)
        elif self._ignored_tag is None and self._stack and self._stack[-1]["cell"] is not None:
            self._stack[-1]["cell"].append(data)

    def handle_endtag(self, tag):
        tag = tag.lower()
        if self._ignored_tag is not None:
            if tag == self._ignored_tag:
                self._ignored_tag = None
            return
        if tag == "title":
            self._in_title = False
            self.page_title = " ".join("".join(self._title_parts).split())
            return
        if not self._stack:
            return
        table = self._stack[-1]
        if tag in ("td", "th"):
            self._finish_cell(table)
        elif tag in ("tr", "thead", "tbody", "tfoot"):
            self._finish_row(table)
        elif tag == "table":
            self._finish_table()

    def close(self):
        super().close()
        self.page_title = " ".join("".join(self._title_parts).split())
        # HTML documents may end without explicitly closing their last cell/row/table.
        while self._stack:
            self._finish_table()


def _fed_header_label(text: str) -> str:
    """Normalize a header, permitting only a trailing footnote marker."""
    text = " ".join(text.split()).casefold()
    return re.sub(r"\s*(?:\[\d+\]|\(\d+\)|[0-9¹²³⁰⁴⁵⁶⁷⁸⁹*†‡]+)\s*$", "", text).strip()


def _fed_date(text: str) -> pd.Timestamp:
    match = re.fullmatch(r"(\d{1,2})-([A-Za-z]{3})-(\d{2}|\d{4})", text.strip())
    if not match or match[2].upper() not in _FED_MONTHS:
        raise ValueError("Neplatné datum v tabulce Fed: " + text[:80])
    day, month, raw_year = match.groups()
    year = int(raw_year) + (2000 if len(raw_year) == 2 else 0)
    try:
        return pd.Timestamp(year=year, month=_FED_MONTHS[month.upper()], day=int(day))
    except ValueError as exc:
        raise ValueError("Neplatné datum v tabulce Fed: " + text[:80]) from exc


def _validated_exchange_rates(rates: pd.DataFrame) -> pd.DataFrame:
    if not {"date", "usd_per_eur"}.issubset(rates.columns):
        raise ValueError("Kurzy Fed musí obsahovat sloupce date a usd_per_eur.")
    result = rates[["date", "usd_per_eur"]].copy()
    result["date"] = pd.to_datetime(result["date"], errors="coerce").dt.normalize()
    result["usd_per_eur"] = pd.to_numeric(result["usd_per_eur"], errors="coerce")
    invalid = result["date"].isna() | ~result["usd_per_eur"].map(
        lambda value: pd.notna(value) and math.isfinite(value) and value > 0)
    if invalid.any():
        raise ValueError("Kurzy Fed obsahují neplatné datum nebo nekladný či nekonečný kurz.")
    if (result.groupby("date")["usd_per_eur"].nunique() > 1).any():
        raise ValueError("Fed obsahuje rozdílné kurzy pro stejný den.")
    result = result.drop_duplicates("date").sort_values("date").reset_index(drop=True)
    if not result["date"].eq(_FED_RATE_ANCHOR).any():
        raise ValueError("V tabulce Fed chybí platný kurz pro 3. 1. 2000.")
    result["usd_per_eur"] = result["usd_per_eur"].astype(float)
    return result


def parse_exchange_rates_html(html: str) -> pd.DataFrame:
    parser = _FedRateTables()
    parser.feed(html)
    parser.close()
    rows = []
    found = False
    for table in parser.tables:
        positions = None
        for row in table:
            labels = [_fed_header_label(cell) for cell in row]
            if "date" in labels and "rate" in labels:
                positions = (labels.index("date"), labels.index("rate"))
                found = True
                continue
            if positions is None or not any(row):
                continue
            date_position, rate_position = positions
            if len(row) <= max(positions):
                raise ValueError("Tabulka Fed obsahuje neúplný řádek Date / Rate.")
            rate_date = _fed_date(row[date_position])
            rate_text = row[rate_position].strip()
            if rate_text.upper() in ("", "ND"):
                continue
            if not re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?", rate_text):
                raise ValueError("Neplatný kurz v tabulce Fed: " + rate_text[:80])
            value = float(rate_text)
            if not math.isfinite(value) or value <= 0:
                raise ValueError("Kurz v tabulce Fed musí být kladné konečné číslo.")
            rows.append((rate_date, value))
    if not found:
        details = f"Počet stažených tabulek: {len(parser.tables)}."
        if parser.page_title:
            details += " Název stažené stránky: " + parser.page_title[:120] + "."
        raise ValueError("Ve stažené odpovědi Fed nebyla rozpoznána tabulka s hlavičkami Date a Rate. " + details)
    return _validated_exchange_rates(pd.DataFrame(rows, columns=["date", "usd_per_eur"]))


def read_exchange_rates() -> pd.DataFrame:
    request = Request(FED_EXCHANGE_RATE_SOURCE, headers={
        "User-Agent": "Mozilla/5.0 (compatible; HenryHubOTEReport/1.0)", "Accept": "text/html"})
    try:
        with urlopen(request, timeout=20) as response:
            body = response.read(5_000_001)
            encoding = response.headers.get_content_charset() or "utf-8"
        if len(body) > 5_000_000:
            raise ValueError("Odpověď překračuje limit 5 MB.")
        html = body.decode(encoding)
    except (HTTPError, URLError, OSError, UnicodeError, LookupError, ValueError) as exc:
        raise RuntimeError("Kurzy USD/EUR se nepodařilo stáhnout z Fed: " + str(exc)[:250]) from exc
    try:
        return parse_exchange_rates_html(html)
    except ValueError as exc:
        raise RuntimeError("Staženou odpověď Fed se nepodařilo zpracovat: " + str(exc)) from exc


def resolve_exchange_rates(dates: pd.Series, rates: pd.DataFrame) -> pd.DataFrame:
    """Use equal-weight neighboring original quotes for every internal missing day."""
    source = _validated_exchange_rates(rates)
    target = dates if isinstance(dates, pd.Series) else pd.Series(dates)
    target_dates = pd.to_datetime(target, errors="coerce").dt.normalize()
    if target_dates.isna().any():
        raise ValueError("Cenová řada obsahuje neplatné datum pro převod měny.")
    observed_dates = source["date"].tolist()
    observed_values = source["usd_per_eur"].tolist()
    anchor_position = bisect_left(observed_dates, _FED_RATE_ANCHOR)
    resolved = []
    for target_date in target_dates:
        position = bisect_left(observed_dates, target_date)
        next_date = pd.NaT
        if target_date < _FED_RATE_ANCHOR:
            value = observed_values[anchor_position]
            previous_date, method = _FED_RATE_ANCHOR, "pre_2000"
        elif position < len(observed_dates) and observed_dates[position] == target_date:
            value = observed_values[position]
            previous_date, method = observed_dates[position], "exact"
        elif position == len(observed_dates):
            value = observed_values[-1]
            previous_date, method = observed_dates[-1], "last_available"
        else:
            value = (observed_values[position - 1] + observed_values[position]) / 2
            previous_date, next_date = observed_dates[position - 1], observed_dates[position]
            method = "neighbor_mean"
        resolved.append((target_date, value, previous_date, next_date, method))
    result = pd.DataFrame(resolved, index=target.index, columns=[
        "date", "usd_per_eur", "rate_date", "next_rate_date", "rate_method"])
    for column in ("date", "rate_date", "next_rate_date"):
        result[column] = pd.to_datetime(result[column])
    result["usd_per_eur"] = result["usd_per_eur"].astype(float)
    return result


def convert_to_eur(prices: pd.DataFrame, rates: pd.DataFrame) -> pd.DataFrame:
    """Convert both daily USD price units before the existing statistics pipeline."""
    result = prices.copy()
    resolved = resolve_exchange_rates(result["date"], rates)
    for column in resolved.columns:
        result[column] = resolved[column].to_numpy()
    for column in ("price_mmbtu", "price_mwh"):
        result[column] = result[column] / result["usd_per_eur"]
    return result.sort_values("date", kind="stable").reset_index(drop=True)


def rolling_mad_annualized(series: pd.Series, window: int) -> pd.Series:
    def mad(values: np.ndarray) -> float:
        median = np.median(values)
        return float(1.4826 * np.median(np.abs(values - median)) * math.sqrt(252))
    return series.rolling(window, min_periods=max(10, window // 2)).apply(mad, raw=True)


def piecewise_breaks_bic(values: np.ndarray, min_size: int = 18, max_breaks: int = 7) -> tuple[list[int], np.ndarray]:
    """Segmentuje řadu na úseky s konstantní úrovní a počet segmentů vybírá podle BIC."""
    y = np.asarray(values, dtype=float)
    n = len(y)
    if n == 0:
        return [], np.array([], dtype=float)
    if n < min_size * 2:
        return [], np.full(n, np.mean(y), dtype=float)
    cs = np.r_[0.0, np.cumsum(y)]
    cs2 = np.r_[0.0, np.cumsum(y * y)]

    def cost(i: int, j: int) -> float:
        if j <= i:
            return np.inf
        mean = (cs[j] - cs[i]) / (j - i)
        return max(0.0, cs2[j] - cs2[i] - (j - i) * mean * mean)

    max_segments = min(max_breaks + 1, n // min_size)
    dp = np.full((max_segments + 1, n + 1), np.inf)
    previous = np.full((max_segments + 1, n + 1), -1, dtype=int)
    dp[0, 0] = 0.0
    for k in range(1, max_segments + 1):
        for j in range(k * min_size, n + 1):
            candidates = range((k - 1) * min_size, j - min_size + 1)
            best = (np.inf, -1)
            for i in candidates:
                if np.isfinite(dp[k - 1, i]):
                    candidate = dp[k - 1, i] + cost(i, j)
                    if candidate < best[0]:
                        best = (candidate, i)
            if best[1] >= 0:
                dp[k, j], previous[k, j] = best

    choices: list[tuple[float, int]] = []
    for k in range(1, max_segments + 1):
        rss = dp[k, n]
        if np.isfinite(rss):
            # Malá konstanta brání log(0) u téměř konstantní řady.
            bic = n * np.log(max(rss / n, 1e-12)) + (2 * k - 1) * np.log(n)
            choices.append((float(bic), k))
    if not choices:
        return [], np.full(n, np.mean(y), dtype=float)

    _, best_k = min(choices)
    ends = [n]
    j = n
    for k in range(best_k, 0, -1):
        i = previous[k, j]
        if i < 0:
            return [], np.full(n, np.mean(y), dtype=float)
        if i > 0:
            ends.append(i)
        j = i
    breaks = sorted(index for index in ends if 0 < index < n)
    fitted = np.empty(n, dtype=float)
    boundaries = [0] + breaks + [n]
    for a, b in zip(boundaries[:-1], boundaries[1:]):
        fitted[a:b] = np.mean(y[a:b])
    return breaks, fitted


def exponential_trend(dates: pd.Series, prices: pd.Series) -> tuple[np.ndarray, float]:
    date_values = pd.to_datetime(dates)
    x = (date_values - date_values.min()).dt.days.to_numpy(dtype=float) / 365.2425
    price_values = np.asarray(prices, dtype=float)
    mask = np.isfinite(x) & np.isfinite(price_values) & (price_values > 0)
    if mask.sum() < 2:
        fitted = np.full(len(price_values), np.nan, dtype=float)
        return fitted, float("nan")
    beta, alpha = np.polyfit(x[mask], np.log(price_values[mask]), 1)
    fitted = np.exp(alpha + beta * x)
    return fitted, float(np.exp(beta) - 1)


def prepare_market(prices: pd.DataFrame) -> dict[str, Any]:
    """Připraví denní a měsíční řady, trendy, klouzavé statistiky a volatilitu pro jeden trh."""
    daily = prices.copy()
    for unit in UNITS:
        daily[f"price_{unit}_nominal"] = daily[f"price_{unit}"]
    keys = [(unit, basis) for unit in UNITS for basis in BASES]

    for unit, basis in keys:
        price_col = f"price_{unit}_{basis}"
        return_col = f"log_return_{unit}_{basis}"
        daily[return_col] = np.log(daily[price_col]).diff()
        indexed = pd.Series(daily[price_col].to_numpy(), index=pd.DatetimeIndex(daily["date"]))
        daily[f"ma12_{unit}_{basis}"] = indexed.rolling("365D", min_periods=1).mean().to_numpy()
        daily[f"median12_{unit}_{basis}"] = indexed.rolling("365D", min_periods=1).median().to_numpy()

    monthly = daily.set_index("date").resample("MS").agg(
        **{
            **{f"price_{unit}_{basis}": (f"price_{unit}_{basis}", "mean") for unit, basis in keys},
            **{f"median_{unit}_{basis}": (f"price_{unit}_{basis}", "median") for unit, basis in keys},
        }
    ).dropna(subset=["price_mmbtu_nominal"]).reset_index()

    breaks_by_key: dict[str, list[pd.Timestamp]] = {}
    trend_rates: dict[str, float] = {}
    for unit, basis in keys:
        key = f"{unit}_{basis}"
        price_col = f"price_{key}"
        median_col = f"median_{key}"
        p = monthly[price_col]
        if len(p) >= 3:
            frac = min(0.20, max(0.06, 36 / len(monthly)))
            monthly[f"ma12_{key}"] = p.rolling(12, min_periods=6).mean()
            monthly[f"median12_{key}"] = p.rolling(12, min_periods=6).median()
            monthly[f"loess_{key}"] = lowess(p, np.arange(len(p)), frac=frac, it=3, return_sorted=False)
            exp_fit, annual_rate = exponential_trend(monthly["date"], p)
            monthly[f"exp_{key}"] = exp_fit
            trend_rates[key] = annual_rate
            found_breaks, fitted = piecewise_breaks_bic(np.log(monthly[median_col].to_numpy()), min_size=18, max_breaks=7)
            monthly[f"breakfit_{key}"] = np.exp(fitted)
            breaks_by_key[key] = [pd.Timestamp(monthly.iloc[i]["date"]) for i in found_breaks]
        else:
            for curve in ["ma12", "median12", "loess", "exp", "breakfit"]:
                monthly[f"{curve}_{key}"] = np.nan
            trend_rates[key] = float("nan")
            breaks_by_key[key] = []

        # Interpolujeme měsíční trendy na denní data, aby byly zobrazené i v denním grafu.
        source = monthly.set_index("date")
        all_dates = source.index.union(pd.DatetimeIndex(daily["date"])).sort_values()
        for curve in ["loess", "exp", "breakfit"]:
            col = f"{curve}_{key}"
            interpolated = source[col].reindex(all_dates).interpolate(method="time").ffill().bfill()
            daily[col] = interpolated.reindex(pd.DatetimeIndex(daily["date"])).to_numpy()

    # Volatilita nominálních cen; jednotkový převod je konstantní.
    for unit, basis in keys:
        return_col = f"log_return_{unit}_{basis}"
        r = daily[return_col]
        daily[f"vol30_{unit}_{basis}"] = r.rolling(30, min_periods=15).std() * math.sqrt(252)
        daily[f"vol90_{unit}_{basis}"] = r.rolling(90, min_periods=45).std() * math.sqrt(252)
        daily[f"vol365_{unit}_{basis}"] = r.rolling(365, min_periods=183).std() * math.sqrt(252)
        daily[f"robust_vol90_{unit}_{basis}"] = rolling_mad_annualized(r, 90)
        daily[f"downside_vol90_{unit}_{basis}"] = r.where(r < 0).rolling(90, min_periods=23).std() * math.sqrt(252)
        daily[f"abs_change_pct_{unit}_{basis}"] = r.abs() * 100
        price_col = f"price_{unit}_{basis}"
        daily[f"price_vol90_{unit}_{basis}"] = daily[price_col].diff().rolling(90, min_periods=45).std()

    return {"daily": daily, "monthly": monthly, "breaks": breaks_by_key, "trend_rates": trend_rates}


def _make_trace(
    x, y, name: str, meta: dict[str, str], color: str,
    dash: str | None = None, yaxis: str = "y", showlegend: bool = True,
) -> go.Scatter:
    return go.Scatter(
        x=list(x), y=list(y), mode="lines", name=name,
        line={"color": color, "width": 2, "shape": "hv" if meta.get("overlay") == "diff_break" else "linear", **({"dash": dash} if dash else {})},
        yaxis=yaxis, meta=meta, visible=False, showlegend=showlegend,
    )


def build_market_figure(prepared: dict[str, Any], market_title: str) -> go.Figure:
    daily = prepared["daily"]
    monthly = prepared["monthly"]
    fig = go.Figure()

    for kind, frame in [("daily", daily), ("monthly", monthly)]:
        for unit in UNITS:
            for basis in BASES:
                key = f"{unit}_{basis}"
                for overlay, curve in OVERLAY_LABELS.items():
                    if overlay == "raw":
                        column = f"price_{key}"
                    elif overlay == "ma":
                        column = f"ma12_{key}"
                    elif overlay == "median":
                        column = f"median12_{key}"
                    elif overlay == "break":
                        column = f"breakfit_{key}"
                    else:
                        column = f"{overlay}_{key}"
                    fig.add_trace(_make_trace(
                        frame["date"], frame[column], curve,
                        {"kind": kind, "unit": unit, "basis": basis, "overlay": overlay},
                        OVERLAY_COLORS[overlay], dash="dash" if overlay == "exp" else ("dot" if overlay == "break" else None),
                        showlegend=True,
                    ))

    for unit in UNITS:
        for basis in BASES:
            key = f"{unit}_{basis}"
            for column, label, color in VOL_SPECS:
                column_name = f"{column}_{key}"
                fig.add_trace(_make_trace(
                    daily["date"], daily[column_name], label,
                    {"kind": "volatility", "unit": unit, "basis": basis, "overlay": column},
                    color, showlegend=True,
                ))

    # Události mají vlastní neviditelnou stopu; zapíná se tlačítkem pod příslušným grafem.
    fig.add_trace(go.Scatter(
        x=[pd.Timestamp(row[0]) for row in EVENTS], y=[0.96] * len(EVENTS),
        mode="markers", name="Události", showlegend=False,
        text=[row[1] for row in EVENTS], customdata=[row[2] for row in EVENTS],
        marker={"size": 8, "color": "rgba(214,39,40,0.45)"},
        hovertemplate="<b>%{text}</b><br>Datum: %{x|%d.%m.%Y}<br>Kategorie: %{customdata}<extra></extra>",
        yaxis="y3", meta={"kind": "events", "unit": "all", "basis": "all", "overlay": "events"}, visible=False,
    ))

    lower_row = {"2020-01-01", "2024-01-01"}
    for start, end, label, color in MARKET_REGIMES:
        fig.add_vrect(x0=start, x1=end, fillcolor=color, line_width=0, layer="below")
        start_date = pd.Timestamp(start)
        end_date = pd.Timestamp(end)
        midpoint = start_date + (end_date - start_date) / 2
        fig.add_annotation(
            x=midpoint, y=0.84 if start in lower_row else 0.94, xref="x", yref="paper",
            text=label, showarrow=False, xanchor="center", yanchor="middle", align="center",
            font={"size": 10, "color": "#334155"}, bgcolor=color,
            bordercolor="rgba(100,116,139,0.18)", borderwidth=0, borderpad=3,
        )

    for date in prepared["breaks"].get("mmbtu_nominal", []):
        fig.add_shape(
            type="line", x0=date, x1=date, y0=0, y1=0.78, xref="x", yref="paper",
            line={"width": 1, "dash": "dot", "color": "rgba(214,39,40,0.45)"}, name="structural-break",
        )

    fig.update_layout(
        template="plotly_white", height=700, hovermode="x unified",
        margin={"l": 70, "r": 70, "t": 90, "b": 55}, title={"text": ""},
        xaxis_title="Datum", yaxis_title="USD/MMBtu", yaxis={"domain": [0, 0.78]},
        yaxis3={"overlaying": "y", "anchor": "x", "domain": [0, 0.78], "range": [0, 1],
                "visible": False, "fixedrange": True, "showgrid": False, "showticklabels": False,
                "zeroline": False, "ticks": ""},
        # Stejná nativní legenda jako u porovnání cen.
        legend={"orientation": "h", "yanchor": "bottom", "y": 1.02, "xanchor": "left", "x": 0},
        showlegend=True,
    )
    fig.update_xaxes(rangeslider_visible=False)
    if market_title == "OTE":
        fig.update_xaxes(range=[daily["date"].min(), daily["date"].max()], autorange=False)
    return fig


def build_comparison_frame(hh: dict[str, Any], ote: dict[str, Any]) -> pd.DataFrame:
    """Spojí Henry Hub a OTE podle společného kalendářního dne."""
    price_cols = [f"price_{unit}_{basis}" for unit in UNITS for basis in BASES]
    left = hh["daily"][["date", *price_cols]].rename(columns={col: f"hh_{col}" for col in price_cols})
    right = ote["daily"][["date", *price_cols]].rename(columns={col: f"ote_{col}" for col in price_cols})
    comparison = left.merge(right, on="date", how="inner").sort_values("date").reset_index(drop=True)
    if comparison.empty:
        raise ValueError("Henry Hub a OTE nemají žádná společná kalendářní data pro porovnání.")
    for unit in UNITS:
        for basis in BASES:
            suffix = f"{unit}_{basis}"
            comparison[f"diff_{suffix}"] = comparison[f"ote_price_{suffix}"] - comparison[f"hh_price_{suffix}"]
    return comparison


def spread_analysis(comparison: pd.DataFrame, unit: str, basis: str) -> dict[str, Any]:
    """Trendy a BIC režimy rozdílu; záporné i nulové hodnoty zůstávají ve výpočtu."""
    suffix = f"{unit}_{basis}"
    common = comparison[["date", f"hh_price_{suffix}", f"ote_price_{suffix}", f"diff_{suffix}"]].dropna().set_index("date").sort_index()
    daily = common[f"diff_{suffix}"]
    monthly = daily.resample("MS").mean().dropna()
    medians = daily.resample("MS").median().reindex(monthly.index)
    values = monthly.to_numpy(dtype=float)
    if len(monthly) >= 3:
        frac = min(0.20, max(0.06, 36 / len(monthly)))
        smooth = lowess(values, np.arange(len(values)), frac=frac, it=3, return_sorted=False)
    else:
        smooth = values.copy()
    # Přímý nelineární fit A*exp(b*t), bez logaritmování a bez vyřazení záporných rozdílů.
    # Model má stálé znaménko; při změnách znaménka jde pouze o aproximaci.
    x = (monthly.index - monthly.index.min()).days.to_numpy(dtype=float)
    x = x / max(float(x.max()), 1.0)
    scale = max(float(np.max(np.abs(values))), 1.0)
    if len(values) >= 2:
        fits = [least_squares(lambda p: p[0] * np.exp(p[1] * x) - values / scale,
                              [float(np.mean(values) / scale), b], bounds=([-np.inf, -10], [np.inf, 10]))
                for b in (-1.0, 0.0, 1.0)]
        fit = min(fits, key=lambda result: np.sum(result.fun ** 2))
        exponential = scale * fit.x[0] * np.exp(fit.x[1] * x)
    else:
        exponential = values.copy()
    # BIC pracuje s cenovou úrovní měsíčních mediánů; log rozdílu není obecně definován.
    breaks, levels = piecewise_breaks_bic(medians.to_numpy(), min_size=18, max_breaks=7)
    curves = pd.DataFrame({"diff_loess": smooth, "diff_exp": exponential, "diff_break": levels}, index=monthly.index)
    daily_curves = {}
    union = curves.index.union(daily.index).sort_values()
    for key in curves:
        series = curves[key].reindex(union)
        # Režimy jsou schody, nikoli lineární přechody mezi režimy.
        interpolated = series.ffill().bfill() if key == "diff_break" else series.interpolate(method="time").ffill().bfill()
        daily_curves[key] = interpolated.reindex(daily.index)
    boundaries = [0] + breaks + [len(monthly)]
    rows = []
    label = UNITS[unit]["label"]
    for number, (a, b) in enumerate(zip(boundaries[:-1], boundaries[1:]), 1):
        start = monthly.index[a]
        end = monthly.index[b] if b < len(monthly) else daily.index.max() + pd.Timedelta(days=1)
        segment = daily[(daily.index >= start) & (daily.index < end)]
        rows.append({
            "Režim": number, "Od": max(start, daily.index.min()).date(),
            "Do": (end - pd.Timedelta(days=1)).date(),
            f"Úroveň režimu ({label})": _fmt_money(levels[a]),
            f"Průměr Henry Hub ({label})": _fmt_money(common.loc[segment.index, f"hh_price_{suffix}"].mean()),
            f"Průměr OTE ({label})": _fmt_money(common.loc[segment.index, f"ote_price_{suffix}"].mean()),
            f"Minimum rozdílu ({label})": _fmt_money(segment.min()),
            f"Maximum rozdílu ({label})": _fmt_money(segment.max()),
            f"Průměr rozdílu ({label})": _fmt_money(segment.mean()),
            f"Medián rozdílu ({label})": _fmt_money(segment.median()),
            f"MAE ({label})": _fmt_money(segment.abs().mean()),
            f"95. percentil absolutního rozdílu ({label})": _fmt_money(segment.abs().quantile(0.95)),
            "Podíl dní OTE > HH": _fmt_pct((segment > 0).mean()),
            "Počet měsíců": b - a, "Počet společných dní": len(segment),
        })
    return {"daily": daily_curves, "monthly": curves, "regimes": pd.DataFrame(rows)}


def spread_volatility_extras(diff: pd.Series) -> dict[str, pd.Series]:
    """Robustní MAD úrovně rozdílu za 90 kalendářních dní, min. 25 hodnot."""
    diff = diff.sort_index()
    def mad(values: np.ndarray) -> float:
        median = np.median(values)
        return float(1.4826 * np.median(np.abs(values - median)))
    return {
        "robust_vol90": diff.rolling("90D", min_periods=25).apply(mad, raw=True),
    }


def build_comparison_figure(comparison: pd.DataFrame) -> go.Figure:
    """Vytvoří interaktivní graf s cenami, rozdílem, klouzavými statistikami a volatilitou rozdílu."""
    fig = go.Figure()
    styles = {
        "hh": ("Henry Hub", "#1f77b4", None),
        "ote": ("OTE", "#ff7f0e", None),
        "diff": ("Rozdíl OTE − Henry Hub", "#2ca02c", None),
        "diff_ma": ("12M klouzavý průměr rozdílu", "#9467bd", "dash"),
        "diff_median": ("12M klouzavý medián rozdílu", "#d62728", "dot"),
        "diff_loess": ("LOESS rozdílu", "#17becf", None),
        "diff_exp": ("Exponenciální trend rozdílu", "#8c564b", "dash"),
        "diff_break": ("Strukturální režimy rozdílu", "#e377c2", "dot"),
    }

    # Denní řady: 365denní časové okno pro průměr a medián rozdílu.
    dated = comparison.set_index("date").sort_index()
    for unit in UNITS:
        for basis in BASES:
            suffix = f"{unit}_{basis}"
            hh_col, ote_col, diff_col = f"hh_price_{suffix}", f"ote_price_{suffix}", f"diff_{suffix}"
            diff_series = dated[diff_col]
            analysis = spread_analysis(comparison, unit, basis)
            daily_curves = {
                "hh": dated[hh_col],
                "ote": dated[ote_col],
                "diff": diff_series,
                "diff_ma": diff_series.rolling("365D", min_periods=30).mean(),
                "diff_median": diff_series.rolling("365D", min_periods=30).median(),
            }
            daily_curves.update(analysis["daily"])
            for overlay, values in daily_curves.items():
                label, color, dash = styles[overlay]
                fig.add_trace(_make_trace(
                    dated.index, values, label,
                    {"kind": "daily", "unit": unit, "basis": basis, "overlay": overlay},
                    color, dash=dash, yaxis="y2" if overlay.startswith("diff") else "y",
                ))

            # Měsíční ceny jsou průměry společných denních pozorování daného měsíce.
            monthly = dated[[hh_col, ote_col, diff_col]].resample("MS").agg({
                hh_col: "mean", ote_col: "mean", diff_col: "mean"
            }).dropna(how="all")
            monthly_median_diff = dated[diff_col].resample("MS").median().reindex(monthly.index)
            monthly_curves = {
                "hh": monthly[hh_col],
                "ote": monthly[ote_col],
                "diff": monthly[diff_col],
                "diff_ma": monthly[diff_col].rolling(12, min_periods=6).mean(),
                "diff_median": monthly_median_diff.rolling(12, min_periods=6).median(),
            }
            monthly_curves.update({key: analysis["monthly"][key].reindex(monthly.index) for key in analysis["monthly"]})
            for overlay, values in monthly_curves.items():
                label, color, dash = styles[overlay]
                fig.add_trace(_make_trace(
                    monthly.index, values, label,
                    {"kind": "monthly", "unit": unit, "basis": basis, "overlay": overlay},
                    color, dash=dash, yaxis="y2" if overlay.startswith("diff") else "y",
                ))

            # Volatilita spreadu: klouzavá směrodatná odchylka samotné řady rozdílu,
            # uváděná v cenové jednotce (není anualizována).
            volatility_specs = [
                ("vol30", "30denní volatilita rozdílu", "30D", 10, "#d62728"),
                ("vol90", "90denní volatilita rozdílu", "90D", 25, "#ff7f0e"),
                ("vol365", "365denní volatilita rozdílu", "365D", 90, "#1f77b4"),
            ]
            for overlay, label, window, min_periods, color in volatility_specs:
                values = diff_series.rolling(window, min_periods=min_periods).std()
                fig.add_trace(_make_trace(
                    dated.index, values, label,
                    {"kind": "volatility", "unit": unit, "basis": basis, "overlay": overlay},
                    color, yaxis="y",
                ))

            extra_styles = [
                ("robust_vol90", "90denní robustní MAD úrovně rozdílu", "#2ca02c"),
            ]
            extras = spread_volatility_extras(diff_series)
            for overlay, label, color in extra_styles:
                fig.add_trace(_make_trace(
                    dated.index, extras[overlay], label,
                    {"kind": "volatility", "unit": unit, "basis": basis, "overlay": overlay},
                    color, yaxis="y",
                ))

    # Události: neviditelné značky mají hover s názvem, datem a kategorií.
    fig.add_trace(go.Scatter(
        x=[pd.Timestamp(row[0]) for row in EVENTS], y=[0.96] * len(EVENTS),
        mode="markers", name="Události", showlegend=False,
        text=[row[1] for row in EVENTS], customdata=[row[2] for row in EVENTS],
        marker={"size": 8, "color": "rgba(214,39,40,0.45)"},
        hovertemplate="<b>%{text}</b><br>Datum: %{x|%d.%m.%Y}<br>Kategorie: %{customdata}<extra></extra>",
        yaxis="y3", meta={"kind": "events", "unit": "all", "basis": "all", "overlay": "events"}, visible=False,
    ))

    # Barevné epochy a jejich popisky používají stejné schéma jako grafy Henry Hub a OTE.
    lower_row = {"2020-01-01", "2024-01-01"}
    for start, end, label, color in MARKET_REGIMES:
        fig.add_vrect(x0=start, x1=end, fillcolor=color, line_width=0, layer="below")
        start_date, end_date = pd.Timestamp(start), pd.Timestamp(end)
        midpoint = start_date + (end_date - start_date) / 2
        fig.add_annotation(
            x=midpoint, y=0.84 if start in lower_row else 0.94, xref="x", yref="paper",
            text=label, showarrow=False, xanchor="center", yanchor="middle", align="center",
            font={"size": 10, "color": "#334155"}, bgcolor=color,
            bordercolor="rgba(100,116,139,0.18)", borderwidth=0, borderpad=3,
        )

    fig.update_layout(
        template="plotly_white", height=700, hovermode="x unified",
        margin={"l": 70, "r": 85, "t": 90, "b": 55}, title={"text": ""},
        xaxis_title="Datum", yaxis_title="Cena (USD/MMBtu)",
        yaxis={"domain": [0, 0.78]},
        yaxis2={"title": "Rozdíl OTE − Henry Hub", "overlaying": "y", "side": "right", "showgrid": False},
        yaxis3={"overlaying": "y", "anchor": "x", "domain": [0, 0.78], "range": [0, 1],
                "visible": False, "fixedrange": True, "showgrid": False, "showticklabels": False,
                "zeroline": False, "ticks": ""},
        legend={"orientation": "h", "yanchor": "bottom", "y": 1.02, "xanchor": "left", "x": 0},
    )
    fig.update_xaxes(rangeslider_visible=False,
                     range=[comparison["date"].min(), comparison["date"].max()], autorange=False)
    return fig


def _fmt_money(value: Any, digits: int = 2) -> str:
    return f"{value:,.{digits}f}".replace(",", " ") if pd.notna(value) and np.isfinite(value) else "—"


def _fmt_pct(value: Any, digits: int = 1) -> str:
    return f"{value:.{digits}%}" if pd.notna(value) and np.isfinite(value) else "—"


def _table_view(prepared: dict[str, Any], unit: str, basis: str) -> tuple[pd.DataFrame, pd.DataFrame, list[pd.Timestamp]]:
    daily_source = prepared["daily"]
    monthly_source = prepared["monthly"]
    key = f"{unit}_{basis}"
    daily = pd.DataFrame({
        "date": daily_source["date"],
        "price": daily_source[f"price_{key}"],
        "returns": daily_source[f"log_return_{key}"],
    })
    monthly = pd.DataFrame({
        "date": monthly_source["date"],
        "price": monthly_source[f"price_{key}"],
        "median": monthly_source[f"median_{key}"],
    })
    return daily, monthly, prepared["breaks"].get(key, [])


def statistics_tables(prepared: dict[str, Any], unit: str, basis: str) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """Vytvoří souhrnné, roční, sezónní a režimové statistiky pro konkrétní jednotku a cenovou bázi."""
    daily, monthly, break_dates = _table_view(prepared, unit, basis)
    price_unit = UNITS[unit]["label"]
    returns = daily["returns"].dropna()
    prices = daily["price"].dropna()
    summary = pd.DataFrame({
        "Metrika": [
            "Počet cen", "Od", "Do", f"Minimum ({price_unit})", f"Maximum ({price_unit})",
            f"Medián ({price_unit})", f"Průměr ({price_unit})", "Celková anualizovaná volatilita",
            "Robustní anualizovaná volatilita", "95. percentil absolutní denní změny",
            "99. percentil absolutní denní změny", "Největší denní růst", "Největší denní pokles",
            f"Podíl dnů nad 5 {price_unit}", f"Podíl dnů nad 10 {price_unit}", f"Podíl dnů nad 15 {price_unit}",
        ],
        "Hodnota": [
            len(prices), daily["date"].min().date(), daily["date"].max().date(),
            _fmt_money(prices.min()), _fmt_money(prices.max()), _fmt_money(prices.median()), _fmt_money(prices.mean()),
            _fmt_pct(returns.std() * math.sqrt(252)),
            _fmt_pct(1.4826 * np.median(np.abs(returns - np.median(returns))) * math.sqrt(252)),
            _fmt_pct(returns.abs().quantile(0.95), 2), _fmt_pct(returns.abs().quantile(0.99), 2),
            _fmt_pct(returns.max(), 2), _fmt_pct(returns.min(), 2),
            _fmt_pct((prices > 5).mean()), _fmt_pct((prices > 10).mean()), _fmt_pct((prices > 15).mean()),
        ],
    })

    annual_source = daily.copy().set_index("date")
    annual = annual_source.groupby(pd.Grouper(freq="YE")).agg(
        mean_price=("price", "mean"), median_price=("price", "median"), minimum=("price", "min"),
        maximum=("price", "max"), count=("price", "count"),
        volatility=("returns", lambda x: x.std() * math.sqrt(252)),
        p95=("returns", lambda x: x.abs().quantile(0.95)),
    ).dropna(subset=["mean_price"]).reset_index()
    annual["Rok"] = annual["date"].dt.year
    annual[f"Minimum ({price_unit})"] = annual["minimum"].map(_fmt_money)
    annual[f"Maximum ({price_unit})"] = annual["maximum"].map(_fmt_money)
    annual[f"Průměrná cena ({price_unit})"] = annual["mean_price"].map(_fmt_money)
    annual[f"Medián ceny ({price_unit})"] = annual["median_price"].map(_fmt_money)
    annual["Volatilita"] = annual["volatility"].map(_fmt_pct)
    annual["95. percentil absolutní změny"] = annual["p95"].map(lambda x: _fmt_pct(x, 2))
    annual["Počet záznamů"] = annual["count"].map(lambda x: f"{int(x):,}".replace(",", " "))
    annual = annual[[
        "Rok", f"Minimum ({price_unit})", f"Maximum ({price_unit})", f"Průměrná cena ({price_unit})",
        f"Medián ceny ({price_unit})", "Volatilita", "95. percentil absolutní změny", "Počet záznamů",
    ]]

    season_source = daily.copy()
    season_source["Měsíc"] = season_source["date"].dt.month
    season = season_source.groupby("Měsíc").agg(
        mean_price=("price", "mean"), median_price=("price", "median"), minimum=("price", "min"),
        maximum=("price", "max"), count=("price", "count"),
        volatility=("returns", lambda x: x.std() * math.sqrt(252)),
        p95=("returns", lambda x: x.abs().quantile(0.95)),
    ).reset_index()
    season[f"Minimum ({price_unit})"] = season["minimum"].map(_fmt_money)
    season[f"Maximum ({price_unit})"] = season["maximum"].map(_fmt_money)
    season[f"Průměrná cena ({price_unit})"] = season["mean_price"].map(_fmt_money)
    season[f"Medián ceny ({price_unit})"] = season["median_price"].map(_fmt_money)
    season["Volatilita"] = season["volatility"].map(_fmt_pct)
    season["95. percentil absolutní změny"] = season["p95"].map(lambda x: _fmt_pct(x, 2))
    season["Počet záznamů"] = season["count"].map(lambda x: f"{int(x):,}".replace(",", " "))
    season = season[[
        "Měsíc", f"Minimum ({price_unit})", f"Maximum ({price_unit})", f"Průměrná cena ({price_unit})",
        f"Medián ceny ({price_unit})", "Volatilita", "95. percentil absolutní změny", "Počet záznamů",
    ]]

    regime_rows = []
    if len(monthly):
        regime_boundaries = [pd.Timestamp(monthly["date"].min())] + break_dates + [pd.Timestamp(monthly["date"].max()) + pd.offsets.MonthBegin(1)]
        for index, (start, end) in enumerate(zip(regime_boundaries[:-1], regime_boundaries[1:]), start=1):
            segment_monthly = monthly[(monthly["date"] >= start) & (monthly["date"] < end)]
            segment_daily = daily[(daily["date"] >= start) & (daily["date"] < end)]
            segment_returns = np.log(segment_daily["price"]).diff().dropna()
            vol = segment_returns.std() * math.sqrt(252) if len(segment_returns) > 1 else np.nan
            p95 = segment_returns.abs().quantile(0.95) if len(segment_returns) else np.nan
            min_price = segment_daily["price"].min() if len(segment_daily) else np.nan
            max_price = segment_daily["price"].max() if len(segment_daily) else np.nan
            # BIC v prepare_market pracuje s logaritmy měsíčních mediánů.
            regime_level = np.exp(np.log(segment_monthly["median"]).mean())
            mean_price = segment_daily["price"].mean()
            median_price = segment_daily["price"].median()
            regime_rows.append({
                "Režim": index, "Od": start.date(), "Do": (end - pd.Timedelta(days=1)).date(),
                f"Úroveň režimu ({price_unit})": _fmt_money(regime_level),
                f"Minimum ({price_unit})": _fmt_money(min_price), f"Maximum ({price_unit})": _fmt_money(max_price),
                f"Průměrná cena ({price_unit})": _fmt_money(mean_price),
                f"Medián ceny ({price_unit})": _fmt_money(median_price),
                "Volatilita (anualizovaná)": _fmt_pct(vol), "95. percentil absolutní změny": _fmt_pct(p95, 2),
                "Počet měsíců": len(segment_monthly),
            })
    regimes = pd.DataFrame(regime_rows)
    return summary, annual, season, regimes


def comparison_tables(comparison: pd.DataFrame, unit: str, basis: str) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    unit_label = UNITS[unit]["label"]
    suffix = f"{unit}_{basis}"
    frame = comparison[["date", f"hh_price_{suffix}", f"ote_price_{suffix}", f"diff_{suffix}"]].copy()
    frame = frame.dropna().rename(columns={
        f"hh_price_{suffix}": "Henry Hub", f"ote_price_{suffix}": "OTE", f"diff_{suffix}": "Rozdíl",
    })
    diff = frame["Rozdíl"]
    abs_diff = diff.abs()
    correlation = frame["Henry Hub"].corr(frame["OTE"]) if len(frame) > 1 else np.nan
    rmse = math.sqrt(float(np.mean(np.square(diff)))) if len(diff) else np.nan
    summary = pd.DataFrame({
        "Metrika": [
            "Počet společných dní", "Od", "Do", f"Průměr Henry Hub ({unit_label})", f"Průměr OTE ({unit_label})",
            f"Průměrný rozdíl OTE − HH ({unit_label})", f"Medián rozdílu ({unit_label})",
            f"Nejmenší rozdíl ({unit_label})", f"Největší rozdíl ({unit_label})",
            f"Průměrná absolutní odchylka (MAE, {unit_label})", f"RMSE rozdílu ({unit_label})",
            f"95. percentil absolutního rozdílu ({unit_label})", "Korelace cen", "Podíl dní, kdy OTE > HH",
        ],
        "Hodnota": [
            len(frame), frame["date"].min().date() if len(frame) else "—", frame["date"].max().date() if len(frame) else "—",
            _fmt_money(frame["Henry Hub"].mean() if len(frame) else np.nan),
            _fmt_money(frame["OTE"].mean() if len(frame) else np.nan),
            _fmt_money(diff.mean() if len(frame) else np.nan), _fmt_money(diff.median() if len(frame) else np.nan),
            _fmt_money(diff.min() if len(frame) else np.nan), _fmt_money(diff.max() if len(frame) else np.nan),
            _fmt_money(abs_diff.mean() if len(frame) else np.nan), _fmt_money(rmse),
            _fmt_money(abs_diff.quantile(0.95) if len(frame) else np.nan), _fmt_pct(correlation),
            _fmt_pct((diff > 0).mean() if len(frame) else np.nan),
        ],
    })

    def aggregate_comparison(grouped, period: str) -> pd.DataFrame:
        table = grouped.agg(
            hh_mean=("Henry Hub", "mean"), ote_mean=("OTE", "mean"),
            diff_min=("Rozdíl", "min"), diff_max=("Rozdíl", "max"),
            diff_mean=("Rozdíl", "mean"), diff_median=("Rozdíl", "median"),
            mae=("Rozdíl", lambda x: x.abs().mean()),
            p95=("Rozdíl", lambda x: x.abs().quantile(0.95)),
            share=("Rozdíl", lambda x: (x > 0).mean()), count=("Rozdíl", "count"),
        ).dropna(subset=["diff_mean"]).reset_index()
        if period == "Rok":
            table[period] = table["date"].dt.year
        result = pd.DataFrame({period: table[period]})
        for title, key in [
            ("Průměr Henry Hub", "hh_mean"), ("Průměr OTE", "ote_mean"),
            ("Minimum rozdílu", "diff_min"), ("Maximum rozdílu", "diff_max"),
            ("Průměr rozdílu", "diff_mean"), ("Medián rozdílu", "diff_median"),
            ("MAE", "mae"), ("95. percentil absolutního rozdílu", "p95"),
        ]:
            result[f"{title} ({unit_label})"] = table[key].map(_fmt_money)
        result["Podíl dní OTE > HH"] = table["share"].map(_fmt_pct)
        result["Počet společných dní"] = table["count"].map(lambda x: f"{int(x):,}".replace(",", " "))
        return result

    annual = aggregate_comparison(frame.set_index("date").groupby(pd.Grouper(freq="YE")), "Rok")
    # Zachováváme sezónní měsíční statistiky: stejný měsíc napříč všemi roky.
    frame["Měsíc"] = frame["date"].dt.month
    season = aggregate_comparison(frame.groupby("Měsíc"), "Měsíc")
    return summary, annual, season


def html_table(df: pd.DataFrame, title: str) -> str:
    table = df.to_html(index=False, classes="data-table", border=0, escape=False)
    return (f"<h3>{title}</h3>"
            f'<div class="data-table-scroll" role="region" aria-label="{title}" tabindex="0">{table}</div>')


def market_table_variants(prepared: dict[str, Any], market_id: str, market_title: str) -> str:
    parts: list[str] = []
    for unit in UNITS:
        for basis in BASES:
            summary, annual, season, regimes = statistics_tables(prepared, unit, basis)
            unit_label = UNITS[unit]["label"]
            basis_label = BASES[basis]["label"]
            active = " active" if unit == "mmbtu" and basis == "nominal" else ""
            parts.append(
                f'<div class="table-variant{active}" data-unit="{unit}" data-basis="{basis}" '
                f'id="{market_id}-tables-{unit}-{basis}">'
                f'<p class="small">Cenová báze: {basis_label}; jednotka: {unit_label}.</p>'
                f'{html_table(summary, "Souhrnné statistiky")}'
                f'{html_table(regimes, "Algoritmicky nalezené cenové režimy")}'
                f'{html_table(season, "Měsíční statistiky")}'
                f'{html_table(annual, "Roční statistiky")}'
                f'</div>'
            )
    return "\n".join(parts)


def comparison_table_variants(comparison: pd.DataFrame) -> str:
    parts: list[str] = []
    for unit in UNITS:
        for basis in BASES:
            summary, annual, season = comparison_tables(comparison, unit, basis)
            regimes = spread_analysis(comparison, unit, basis)["regimes"]
            unit_label = UNITS[unit]["label"]
            basis_label = BASES[basis]["label"]
            active = " active" if unit == "mmbtu" and basis == "nominal" else ""
            parts.append(
                f'<div class="table-variant comparison-table-variant{active}" data-unit="{unit}" data-basis="{basis}" '
                f'id="comparison-tables-{unit}-{basis}">'
                f'<p class="small">Cenová báze: {basis_label}; jednotka: {unit_label}. Rozdíl je definován jako OTE minus Henry Hub.</p>'
                f'{html_table(summary, "Souhrnné statistiky")}'
                f'{html_table(regimes, "Algoritmicky nalezené cenové režimy")}'
                f'<p class="small">Režimy: BIC nad měsíčními mediány rozdílu, bez logaritmování. Průměry, medián, MAE a percentil se počítají ze společných denních pozorování. Exponenciální trend A·exp(b·t) se odhaduje přímo i ze záporných hodnot; nedokáže měnit znaménko.</p>'
                f'{html_table(season, "Měsíční statistiky")}'
                f'{html_table(annual, "Roční statistiky")}'
                f'</div>'
            )
    return "\n".join(parts)


def _market_kpi_html(prepared: dict[str, Any]) -> str:
    rate = _fmt_pct(prepared["trend_rates"].get("mmbtu_nominal", float("nan")), 2)
    return f'<div class="kpis"><div><b>Exponenciální trend – nominálně:</b> {rate} ročně</div></div>'



def build_excel_export(figure: go.Figure, currency: str, unit: str,
                       basis: str = "nominal") -> bytes:
    """Exportuje přesně čtyři řady grafu sekce 3; funguje i bez webového serveru.

    Měsíční trend se nepřepočítává z denní interpolace: přebírá se přímo
    z měsíčního grafu. Export vždy obsahuje celé společné období.
    """
    from openpyxl import Workbook
    from openpyxl.chart import LineChart, Reference
    from openpyxl.chart.axis import DateAxis
    from openpyxl.drawing.image import Image
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    label = currency + ("/MWh" if unit == "mwh" else "/MMBtu")
    overlays = ("hh", "ote", "diff", "diff_exp")
    names = ("Henry Hub", "OTE", "Rozdíl OTE − HH", "Exponenciální trend rozdílu")
    colors = ("1F77B4", "FF7F0E", "2CA02C", "8C564B")
    for kind, sheet_name in (("daily", "Denní ceny"), ("monthly", "Měsíční ceny")):
        columns = {}
        for trace in figure.data:
            meta = trace.meta or {}
            if (meta.get("kind") == kind and meta.get("unit") == unit
                    and meta.get("basis") == basis and meta.get("overlay") in overlays):
                columns[meta["overlay"]] = pd.Series(
                    np.asarray(trace.y, dtype=float), index=pd.to_datetime(trace.x))
        if set(columns) != set(overlays):
            raise ValueError(f"Export: chybějí řady pro {currency}/{unit}/{kind}.")
        data = pd.DataFrame({key: columns[key] for key in overlays}).sort_index()
        if data.empty:
            raise ValueError("Export: sekce 3 neobsahuje společná data.")
        ws = wb.create_sheet(sheet_name)
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.sheet_properties.tabColor = "82302B"
        ws.merge_cells("A1:D2")
        ws["A1"] = f"Henry Hub & OTE – {sheet_name.lower()}"
        ws["A1"].font = Font(size=18, bold=True, color="82302B")
        ws["A1"].alignment = Alignment(vertical="center")
        ws.merge_cells("A3:E3")
        ws["A3"] = f"{label} | Nominální ceny | Společné období: {data.index.min():%d.%m.%Y} – {data.index.max():%d.%m.%Y}"
        ws.merge_cells("A4:E4")
        ws["A4"] = "Měsíční ceny: průměry společných dnů; trend převzat z grafu sekce 3."
        ws.merge_cells("A5:E5")
        ws["A5"] = "Rozdíl = OTE − HH. Exponenciální trend popisuje rozdíl, nejde o předpověď."
        for r in (3, 4, 5):
            ws.cell(r, 1).font = Font(size=10, color="596579")
        logo = Image(BytesIO(base64.b64decode(PRINT_LOGO_DATA_URI.split(",", 1)[1])))
        logo.width = 167
        logo.height = 81
        ws.add_image(logo, "G1")
        header_row = 8
        ws.append([])
        headers = ["Datum" if kind == "daily" else "Měsíc"] + [f"{name} ({label})" for name in names]
        for col, name in enumerate(headers, 1):
            cell = ws.cell(header_row, col, name)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="82302B")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[header_row].height = 34
        for row, (date, values) in enumerate(data.iterrows(), header_row + 1):
            ws.cell(row, 1, date.to_pydatetime()).number_format = "dd.mm.yyyy" if kind == "daily" else "mm/yyyy"
            for col, value in enumerate(values, 2):
                ws.cell(row, col, float(value) if np.isfinite(value) else None).number_format = '#,##0.00;[Red]-#,##0.00'
        last = header_row + len(data)
        table = Table(displayName="DailyPrices" if kind == "daily" else "MonthlyPrices", ref=f"A8:E{last}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
        for col, width in enumerate((15, 23, 23, 25, 33), 1):
            ws.column_dimensions[get_column_letter(col)].width = width
        ws.freeze_panes = "B9"
        # Dva spojené čárové grafy se skutečnou datumovou osou, stejně jako HTML.
        chart = LineChart()
        chart.title = f"Henry Hub a OTE – {sheet_name.lower()}"
        chart.style = 13
        chart.width, chart.height = 29, 15
        chart.y_axis.title = f"Cena ({label})"
        chart.x_axis = DateAxis(axId=10, crossAx=100)
        chart.x_axis.axPos = "b"
        chart.x_axis.number_format = "mm/yyyy"
        chart.x_axis.title = "Datum" if kind == "daily" else "Měsíc"
        chart.x_axis.majorTimeUnit = "months"
        chart.y_axis.crosses = "min"
        chart.add_data(Reference(ws, min_col=2, max_col=3, min_row=8, max_row=last), titles_from_data=True)
        dates = Reference(ws, min_col=1, min_row=9, max_row=last)
        chart.set_categories(dates)
        right = LineChart()
        right.x_axis = chart.x_axis
        right.y_axis.axPos = "r"
        right.y_axis.axId = 200
        right.y_axis.title = f"Rozdíl OTE − HH ({label})"
        right.y_axis.crosses = "max"
        right.y_axis.majorGridlines = None
        right.add_data(Reference(ws, min_col=4, max_col=5, min_row=8, max_row=last), titles_from_data=True)
        right.set_categories(dates)
        for i, series in enumerate(list(chart.series) + list(right.series)):
            series.graphicalProperties.line.solidFill = colors[i]
            series.graphicalProperties.line.width = 19050
            series.marker.symbol = "none"
            series.smooth = False
            if i == 3:
                series.graphicalProperties.line.prstDash = "dash"
        chart += right
        chart.legend.position = "b"
        chart.display_blanks = "gap"
        ws.add_chart(chart, "G8")
        ws.sheet_view.zoomScale = 85
    result = BytesIO()
    wb.save(result)
    return result.getvalue()


def make_html(
    hh_figure: go.Figure,
    ote_figure: go.Figure,
    comparison_figure: go.Figure,
    hh_prepared: dict[str, Any],
    ote_prepared: dict[str, Any],
    comparison: pd.DataFrame,
    output: Path,
    eur: dict[str, Any],
    forward_data: dict[str, Any] | None = None,
    ote_forward_data: dict[str, Any] | None = None,
) -> None:
    if forward_data is None:
        forward_data = {"rows": [], "trade_date": None, "fetched_at": None, "error": "Data CME nebyla při vytváření reportu načtena."}
    if ote_forward_data is None:
        ote_forward_data = {"rows": [], "quote_date": None, "source_file": "Forwardy OTE.xlsx", "error": "Soubor Forwardy OTE.xlsx nebyl při vytváření reportu načten."}
    event_df = pd.DataFrame([row[:4] for row in EVENTS], columns=["Datum", "Událost", "Kategorie", "Možný kanál"])
    event_df["Datum"] = pd.to_datetime(event_df["Datum"]).dt.strftime("%d.%m.%Y")

    usd = {"hh": hh_prepared, "ote": ote_prepared, "comparison": comparison,
           "figures": {"hh": hh_figure, "ote": ote_figure, "comparison": comparison_figure}}
    bundles = {"USD": usd, "EUR": eur}
    figure_sets = {currency: {key: fig.to_plotly_json() for key, fig in bundle["figures"].items()}
                   for currency, bundle in bundles.items()}
    excel_exports = {
        f"{currency}:{unit}:{basis}": base64.b64encode(
            build_excel_export(bundle["figures"]["comparison"], currency, unit, basis)
        ).decode("ascii")
        for currency, bundle in bundles.items() for unit in UNITS for basis in BASES
    }
    table_parts = {"hh": [], "ote": [], "comparison": []}
    kpis = {"hh": [], "ote": []}
    def currency_wrap(content: str, currency: str) -> str:
        active = " active" if currency == "USD" else ""
        return f'<div class="currency-variant{active}" data-currency="{currency}">{content}</div>'
    old_labels = {key: value.copy() for key, value in UNITS.items()}
    try:
        for currency, bundle in bundles.items():
            for unit in UNITS:
                UNITS[unit]["label"] = currency + ("/MWh" if unit == "mwh" else "/MMBtu")
            for market, title in [("hh", "Henry Hub"), ("ote", "OTE")]:
                table_parts[market].append(currency_wrap(market_table_variants(bundle[market], market + '-' + currency, title), currency))
                kpis[market].append(currency_wrap(_market_kpi_html(bundle[market]), currency))
            tables = comparison_table_variants(bundle["comparison"])
            tables = tables.replace('id="comparison-tables-', f'id="comparison-{currency}-tables-')
            table_parts["comparison"].append(currency_wrap(tables, currency))
    finally:
        UNITS.update(old_labels)
    hh_tables, ote_tables, compare_tables = ("".join(table_parts[key]) for key in ("hh", "ote", "comparison"))
    event_table = html_table(event_df, "Události s potenciálním vlivem na cenu")

    html = r'''<!doctype html>
<html lang="cs">
<head>
<meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Henry Hub a OTE – interaktivní analýza spotových cen plynu</title>
<script>__PLOTLY_JS__</script>
<style>
:root { color-scheme: light; }
body { font-family: Segoe UI, Arial, sans-serif; margin: 0; background: #f4f6f8; color: #1f2937; }
.wrap { max-width: 1550px; margin: auto; padding: 20px; }
.panel { background: white; border-radius: 12px; padding: 16px 18px; margin-bottom: 16px; box-shadow: 0 2px 10px #00000010; }
.controls { display: grid; gap: 10px; }
.row { display: flex; flex-wrap: wrap; gap: 7px; align-items: center; }
.label { font-weight: 700; min-width: 145px; }
button { border: 1px solid #cbd5e1; background: #fff; padding: 8px 12px; border-radius: 8px; cursor: pointer; }
button.active { background: #0f6cbd; color: white; border-color: #0f6cbd; }
button:disabled { opacity: .45; cursor: not-allowed; }
.data-table { border-collapse: collapse; width: 100%; font-size: 13px; }
.data-table th { background: #e8f1fb; text-align: left; position: sticky; top: 0; white-space: nowrap; }
.data-table th, .data-table td { padding: 7px 9px; border-bottom: 1px solid #e5e7eb; }
.data-table tbody tr:nth-child(even) { background: #fafbfc; }
.data-table-scroll { display:block; width:100%; max-width:100%; overflow-x:auto; margin-bottom:14px; }
.data-table-scroll .data-table { width:max-content; min-width:100%; }
.data-table-scroll:focus-visible { outline:2px solid #832C29; outline-offset:2px; }
.table-scroll { overflow: auto; max-height: 430px; margin-bottom: 8px; }
.comparison-annual-scroll { display: block; width: 100%; max-width: 100%; overflow-x: auto; overflow-y: hidden; margin: 8px 0 14px; border-bottom: 1px solid #e5e7eb; }
.comparison-annual-scroll .data-table { width: max-content; min-width: 1250px; }
.comparison-annual-scroll h3 { margin-top: 14px; }
.kpis { display: flex; gap: 20px; flex-wrap: wrap; margin: 8px 0; }
.small { font-size: 12px; color: #64748b; }
.chart-title { font-size: 22px; line-height: 1.35; font-weight: 650; margin: 0 0 18px; color: #1f2937; }
.chart-panel { padding-top: 20px; }
.section-title { margin-top: 0; }
.table-variant { display: none; }
.table-variant.active { display: block; }
.glossary { display: grid; grid-template-columns: minmax(190px, .8fr) minmax(0, 2fr); column-gap: 18px; row-gap: 8px; margin: 8px 0 0; }
.glossary dt { font-weight: 650; color: #1f2937; }
.glossary dd { margin: 0; color: #475569; }
#forwardError { color:#9a3412; }
.section-intro { color: #475569; margin-top: 0; }
@media(max-width: 700px) {
  .wrap { padding: 10px; }
  .glossary { grid-template-columns: 1fr; row-gap: 3px; }
  .glossary dd { margin-bottom: 9px; }
  .label { min-width: 100%; }
}
html { scroll-behavior: smooth; scroll-padding-top: var(--header-height, 210px); }
#reportHeader { position: sticky; top: 0; z-index: 1000; background: #832C29; color: white; box-shadow: 0 3px 14px #0003; }
.header-inner { max-width: 1550px; margin: auto; padding: 14px 20px; }
#reportHeader h1 { font-size: 23px; margin: 0 0 12px; }
#reportHeader small { font-size: 14px; font-weight: 400; margin-left: 12px; }
#reportHeader nav { display: flex; flex-wrap: wrap; gap: 8px; margin-bottom: 12px; }
#reportHeader a { color: white; text-decoration: none; padding: 6px 10px; border: 1px solid #ffffff66; border-radius: 7px; }
#reportHeader a:hover { background: #ffffff22; }
#reportHeader button.active { background: #361311; border-color: white; }
#reportHeader input { width: 110px; border-radius: 6px; border: 1px solid #ddd; padding: 8px; }
#fxStatus { font-size: 12px; margin: 10px 0 0; }
.data-table th { position: static; }
.panel { overflow-x: auto; }
@media(max-width:700px) { #reportHeader h1 { font-size: 18px; } .header-inner { padding: 10px; } #reportHeader .row { gap: 5px; } }
.brand-row { display:flex; align-items:center; gap:20px; margin-bottom:12px; }
.brand-logo { width:160px; height:75px; object-fit:contain; background:white; padding:5px 12px; border-radius:8px; }
#reportHeader .brand-row h1 { margin:0; }
#reportHeader button { background:transparent; color:white; border:1px solid #ffffff88; }
#reportHeader button.active, #reportHeader nav a.active { background:#fff; color:#832C29; border-color:#fff; font-weight:800; box-shadow:0 0 0 2px #ffffff55; }
#reportHeader button.active::before { content:"✓ "; }
#reportHeader a:focus-visible, #reportHeader button:focus-visible { outline:3px solid #ffda83; outline-offset:3px; }
.currency-variant { display:none; }
.currency-variant.active { display:block; }
@media(max-width:700px) { .brand-logo { width:110px; height:52px; } .brand-row { gap:12px; } }
#reportHeader .header-inner { display:flex; align-items:center; flex-wrap:nowrap; gap:6px; max-width:none; padding:8px 18px; overflow-x:auto; }
#reportHeader .brand-row { display:flex; align-items:center; flex:0 0 auto; gap:10px; margin:0; }
#reportHeader .brand-logo { width:100px; height:46px; padding:0; border:0; border-radius:0; background:transparent; object-fit:contain; }
#reportHeader h1 { font-size:17px; white-space:nowrap; margin:0; }
#reportHeader small { font-size:11px; margin-left:6px; }
#reportHeader nav { flex:0 0 auto; flex-wrap:nowrap; gap:4px; margin:0 0 0 20px; }
#reportHeader .row { flex:0 0 auto; flex-wrap:nowrap; gap:4px; white-space:nowrap; margin-left:16px; }
#reportHeader button, #reportHeader nav a { white-space:nowrap; font-size:12px; padding:6px 8px; }
#reportHeader .row span { font-size:12px; }
.report-footer { text-align:center; padding:22px 12px 26px; color:#7b8490; font-size:12px; }
/* Posuvník patří konkrétní široké tabulce, nikoli celé sekci. */
#comparison-section { overflow:visible; min-width:0; }
#comparison-section .currency-variant,
#comparison-section .comparison-table-variant { min-width:0; max-width:100%; }
.comparison-regimes-scroll { display:block; width:100%; max-width:100%; overflow-x:scroll; overflow-y:hidden; margin:8px 0 14px; border-bottom:1px solid #e5e7eb; }
.comparison-regimes-scroll .data-table { width:max-content; min-width:100%; }
.comparison-regimes-scroll:focus-visible { outline:2px solid #832C29; outline-offset:2px; }
.comparison-season-table { width:100%; max-width:100%; }
.comparison-season-table .data-table { width:100%; table-layout:fixed; }
.comparison-season-table .data-table th,
.comparison-season-table .data-table td { white-space:normal; overflow-wrap:anywhere; }
#reportHeader .currency-label { margin-left:10px; }
#reportHeader #pdfExportBtn { margin-left:0; flex:0 0 auto; background:white; color:#832C29; font-weight:700; border:1px solid white; padding:8px 12px; }
#reportHeader #excelExportBtn { margin-left:auto; flex:0 0 auto; background:white; color:#832C29; font-weight:700; border:1px solid white; padding:8px 12px; }
#pdfDialog { border:0; border-radius:14px; width:min(650px, calc(100vw - 40px)); max-height:85vh; padding:24px; color:#1f2937; box-shadow:0 20px 80px #0005; }
#pdfDialog::backdrop { background:#0007; }
#pdfDialog h2 { margin:0 0 10px; }
#pdfDialog fieldset { border:1px solid #dde2e8; border-radius:8px; margin:12px 0; padding:10px 12px; }
#pdfDialog legend { font-weight:700; }
#pdfDialog label { display:flex; align-items:flex-start; gap:9px; padding:6px 0; cursor:pointer; }
#pdfDialog input { accent-color:#832C29; width:17px; height:17px; flex-shrink:0; }
.pdf-actions { display:flex; gap:8px; justify-content:flex-end; margin-top:16px; }
#pdfConfirm { background:#832C29; color:white; border-color:#832C29; }
#pdfStatus { min-height:1.3em; font-size:13px; }
</style>
</head>
<body>
<header id="reportHeader">
  <div class="header-inner">
    <div class="brand-row"><img class="brand-logo" src="__LOGO__" alt="LAMA energy"><h1>Henry Hub &amp; OTE <small>Spotové ceny plynu</small></h1></div>
    <nav aria-label="Části reportu">
      <a href="#hh-section">1. Henry Hub</a><a href="#ote-section">2. OTE</a>
      <a href="#comparison-section">3. Porovnání</a><a href="#forward-section">4. Forward</a><a href="#events-section">Události</a><a href="#glossary-section">Slovníček</a>
    </nav>
    <div class="row">
      <span>Jednotka</span><button class="unit-btn active" data-unit="mmbtu">MMBtu</button><button class="unit-btn" data-unit="mwh">MWh</button>
      <span class="currency-label">Měna</span><button class="currency-btn active" data-currency="USD">USD</button><button class="currency-btn" data-currency="EUR">EUR</button>
    </div>
    <button id="excelExportBtn" type="button" title="Dva listy: denní a měsíční ceny sekce 3, aktuální měna a jednotka, celé období.">Exportovat do Excelu</button>
    <button id="pdfExportBtn" type="button" aria-haspopup="dialog">Exportovat do PDF</button>
  </div>
</header>
<dialog id="pdfDialog" aria-labelledby="pdfDialogTitle">
  <h2 id="pdfDialogTitle">Exportovat do PDF</h2>
  <p class="small">Vyberte grafy a tabulky. Export zachová aktuální měnu, jednotku, zobrazené křivky a přiblížení grafů. V následném tiskovém dialogu vyberte „Uložit jako PDF“.</p>
  <div class="row"><button id="pdfSelectAll" type="button">Vybrat vše</button><button id="pdfSelectNone" type="button">Zrušit výběr</button></div>
  <div id="pdfChoices"></div>
  <p id="pdfStatus" role="status" aria-live="polite"></p>
  <div class="pdf-actions"><button id="pdfCancel" type="button">Zrušit</button><button id="pdfConfirm" type="button">Exportovat</button></div>
</dialog>
<div class="wrap">
  <div class="panel"><p>Srovnání amerického Henry Hub a českého trhu OTE. Henry Hub se načítá z listu Data 1 souboru <a href="https://www.eia.gov/dnav/ng/hist/rngwhhdd.htm" target="_blank" rel="noopener noreferrer">EIA</a>; OTE z prvního listu původního Excelu, do kterého se při vytváření reportu doplňují novější dny z ročních zpráv OTE. Porovnání používá pouze společné kalendářní dny. Více cen ve stejném dni se zprůměruje.</p><p class="small">1 MWh = 3,4121416 MMBtu. OTE − HH: kladný rozdíl znamená dražší OTE. Přepnutí měny převádí grafy i cenové statistiky. EUR = USD / Rate z <a href="https://www.federalreserve.gov/RELEASES/H10/hist/dat00_eu.htm" target="_blank" rel="noopener noreferrer">Federal Reserve</a> (počet USD za 1 EUR). Chybějící den včetně ND se doplní průměrem nejbližšího předchozího a následujícího dostupného kurzu. Před 3. 1. 2000 se použije kurz z 3. 1. 2000; po posledním dostupném dni (__FX_END__) poslední zveřejněný kurz. Trendy, výnosy a statistiky se počítají z převedených denních cen. Přepínače platí pro všechny čtyři sekce.</p></div>

  <div class="panel" id="hh-section">
    <h2 class="section-title">1. Henry Hub - spotové ceny plynu</h2>
    __HH_KPI__
    <div class="controls">
      <div class="row"><span class="label">Zobrazení</span>
        <button data-market="hh" data-kind="daily" class="kind-btn active">Denní ceny</button>
        <button data-market="hh" data-kind="monthly" class="kind-btn">Měsíční ceny</button>
        <button data-market="hh" data-kind="volatility" class="kind-btn">Volatilita</button>
      </div>
      <div class="row"><span class="label">Překryvné řady</span>
        <button data-market="hh" data-overlay="raw" class="overlay-btn active">Cena</button>
        <button data-market="hh" data-overlay="ma" class="overlay-btn active">12M průměr</button>
        <button data-market="hh" data-overlay="median" class="overlay-btn active">12M medián</button>
        <button data-market="hh" data-overlay="loess" class="overlay-btn active" title="LOESS je lokální vyhlazovací metoda odhadující trend podle blízkých pozorování.">LOESS</button>
        <button data-market="hh" data-overlay="exp" class="overlay-btn" title="Exponenciální trend popisuje dlouhodobé složené tempo růstu či poklesu, nikoli předpověď budoucí ceny.">Exponenciální trend</button>
        <button data-market="hh" data-overlay="break" class="overlay-btn active" title="Strukturální režimy dělí historii na cenové úseky; počet zlomů vybírá kritérium BIC.">Strukturální režimy</button>
      </div>
      <div class="row"><span class="label">Volatilita</span>
        <button data-market="hh" data-vol="vol30" class="vol-btn active">30 dní</button>
        <button data-market="hh" data-vol="vol90" class="vol-btn active">90 dní</button>
        <button data-market="hh" data-vol="vol365" class="vol-btn active">365 dní</button>
        <button data-market="hh" data-vol="robust_vol90" class="vol-btn" title="MAD je robustní odhad kolísání méně citlivý na extrémní hodnoty.">Robustní MAD</button>
        <button data-market="hh" data-vol="downside_vol90" class="vol-btn" title="Kolísavost záporných denních logaritmických výnosů.">Volatilita poklesů</button>
        <button data-market="hh" data-vol="abs_change_pct" class="vol-btn" title="Absolutní velikost denního logaritmického výnosu v procentech.">Abs. denní změna</button>
        <button data-market="hh" data-vol="price_vol90" class="vol-btn" title="Směrodatná odchylka denních změn ceny v aktuálně zvolené jednotce.">Změna ceny</button>
      </div>
      <div class="row"><span class="label">Události</span><button data-market="hh" class="events-btn">Značky událostí</button></div>
    </div>
    <div class="chart-panel"><h3 id="hhChartTitle" class="chart-title">Henry Hub: denní ceny</h3><div id="hhChart" style="height:700px"></div></div>
    __HH_TABLES__
  </div>

  <div class="panel" id="ote-section">
    <h2 class="section-title">2. OTE – spotové ceny plynu</h2>
    <p class="small">Poslední den OTE v reportu: <strong>__OTE_END__</strong>. Nově doplněné ceny NC BAL v Kč/MWh se převádějí na USD podle stejného nebo posledního předchozího kurzu ČNB. Roční zpráva V0 je zveřejňována průběžně do D−3.</p>
    __OTE_KPI__
    <div class="controls">
      <div class="row"><span class="label">Zobrazení</span>
        <button data-market="ote" data-kind="daily" class="kind-btn active">Denní ceny</button>
        <button data-market="ote" data-kind="monthly" class="kind-btn">Měsíční ceny</button>
        <button data-market="ote" data-kind="volatility" class="kind-btn">Volatilita</button>
      </div>
      <div class="row"><span class="label">Překryvné řady</span>
        <button data-market="ote" data-overlay="raw" class="overlay-btn active">Cena</button>
        <button data-market="ote" data-overlay="ma" class="overlay-btn active">12M průměr</button>
        <button data-market="ote" data-overlay="median" class="overlay-btn active">12M medián</button>
        <button data-market="ote" data-overlay="loess" class="overlay-btn active" title="LOESS je lokální vyhlazovací metoda odhadující trend podle blízkých pozorování.">LOESS</button>
        <button data-market="ote" data-overlay="exp" class="overlay-btn" title="Exponenciální trend popisuje dlouhodobé složené tempo růstu či poklesu, nikoli předpověď budoucí ceny.">Exponenciální trend</button>
        <button data-market="ote" data-overlay="break" class="overlay-btn active" title="Strukturální režimy dělí historii na cenové úseky; počet zlomů vybírá kritérium BIC.">Strukturální režimy</button>
      </div>
      <div class="row"><span class="label">Volatilita</span>
        <button data-market="ote" data-vol="vol30" class="vol-btn active">30 dní</button>
        <button data-market="ote" data-vol="vol90" class="vol-btn active">90 dní</button>
        <button data-market="ote" data-vol="vol365" class="vol-btn active">365 dní</button>
        <button data-market="ote" data-vol="robust_vol90" class="vol-btn" title="MAD je robustní odhad kolísání méně citlivý na extrémní hodnoty.">Robustní MAD</button>
        <button data-market="ote" data-vol="downside_vol90" class="vol-btn" title="Kolísavost záporných denních logaritmických výnosů.">Volatilita poklesů</button>
        <button data-market="ote" data-vol="abs_change_pct" class="vol-btn" title="Absolutní velikost denního logaritmického výnosu v procentech.">Abs. denní změna</button>
        <button data-market="ote" data-vol="price_vol90" class="vol-btn" title="Směrodatná odchylka denních změn ceny v aktuálně zvolené jednotce.">Změna ceny</button>
      </div>
      <div class="row"><span class="label">Události</span><button data-market="ote" class="events-btn">Značky událostí</button></div>
    </div>
    <div class="chart-panel"><h3 id="oteChartTitle" class="chart-title">OTE: denní ceny</h3><div id="oteChart" style="height:700px"></div></div>
    __OTE_TABLES__
  </div>

  <div class="panel" id="comparison-section">
    <h2 class="section-title">3. Porovnání Henry Hub a OTE</h2>
    <p class="section-intro">Kladný rozdíl znamená, že je OTE dražší než Henry Hub; záporný rozdíl znamená, že je levnější. Cenové řady a rozdíl mají při zobrazení cen samostatné svislé osy. Průměr a medián rozdílu jsou klouzavé statistiky přibližně za posledních 12 měsíců.</p>
    <div class="controls comparison-controls">
      <div class="row"><span class="label">Zobrazení</span>
        <button data-comparison-kind="daily" class="comparison-kind-btn active">Denní ceny</button>
        <button data-comparison-kind="monthly" class="comparison-kind-btn">Měsíční ceny</button>
        <button data-comparison-kind="volatility" class="comparison-kind-btn">Volatilita rozdílu</button>
      </div>
      <div class="row"><span class="label">Křivky</span>
        <button data-comparison-overlay="hh" class="comparison-overlay-btn active">Henry Hub</button>
        <button data-comparison-overlay="ote" class="comparison-overlay-btn active">OTE</button>
        <button data-comparison-overlay="diff" class="comparison-overlay-btn active">Rozdíl OTE − HH</button>
        <button data-comparison-overlay="diff_ma" class="comparison-overlay-btn">Průměr rozdílu (12M)</button>
        <button data-comparison-overlay="diff_median" class="comparison-overlay-btn">Medián rozdílu (12M)</button>
        <button data-comparison-overlay="diff_loess" class="comparison-overlay-btn" title="Lokální vyhlazení měsíčního rozdílu OTE − HH.">LOESS</button>
        <button data-comparison-overlay="diff_exp" class="comparison-overlay-btn" title="Přímý fit A·exp(b·t), včetně záporných hodnot. Model nemění znaménko a není předpovědí.">Exponenciální trend</button>
        <button data-comparison-overlay="diff_break" class="comparison-overlay-btn" title="BIC segmentace měsíčních mediánů rozdílu v původní cenové úrovni.">Strukturální režimy</button>
      </div>
      <div class="row"><span class="label">Předvolby křivek</span>
        <button data-comparison-preset="prices" class="comparison-preset-btn">Obě ceny</button>
        <button data-comparison-preset="difference" class="comparison-preset-btn">Rozdíl + průměr/medián</button>
        <button data-comparison-preset="all" class="comparison-preset-btn">Vše současně</button>
      </div>
      <div class="row"><span class="label">Volatilita rozdílu</span>
        <button data-comparison-vol="vol30" class="comparison-vol-btn active" title="Směrodatná odchylka rozdílu cen v klouzavém 30denním časovém okně; není anualizovaná.">30 dní</button>
        <button data-comparison-vol="vol90" class="comparison-vol-btn active" title="Směrodatná odchylka rozdílu cen v klouzavém 90denním časovém okně; není anualizovaná.">90 dní</button>
        <button data-comparison-vol="vol365" class="comparison-vol-btn active" title="Směrodatná odchylka rozdílu cen v klouzavém 365denním časovém okně; není anualizovaná.">365 dní</button>
        <button data-comparison-vol="robust_vol90" class="comparison-vol-btn" title="1,4826 × medián absolutních odchylek úrovně rozdílu od mediánu v 90 kalendářních dnech; nejméně 25 hodnot. V cenové jednotce, bez anualizace.">Robustní MAD</button>
      </div>
      <div class="row"><span class="label">Události</span>
        <button id="comparisonEventsBtn" class="comparison-events-btn">Značky událostí</button>

      </div>
    </div>
    <div class="chart-panel"><h3 id="comparisonChartTitle" class="chart-title">Porovnání Henry Hub a OTE – denní ceny</h3><div id="comparisonChart" style="height:700px"></div></div>
    __COMPARISON_TABLES__
  </div>

  <div class="panel" id="forward-section">
    <h2 class="section-title">4. Forward</h2>
    <p class="section-intro">Henry Hub: vypořádací ceny CME podle měsíce dodání kontraktu (výchozí USD/MMBtu). OTE: ceny z jediného řádku Excelu s datem nejbližším dni vytvoření reportu (výchozí EUR/MWh). Měsíce jsou samostatné body, kvartály vodorovné úsečky přes tři měsíce a Cal přes celý rok.</p>
    <div class="row"><a href="https://www.cmegroup.com/markets/energy/natural-gas/natural-gas.settlements.html" target="_blank" rel="noopener noreferrer">Zdroj: CME Group</a></div>
    <p class="small">Datum vypořádání: <strong id="forwardTradeDate">—</strong> · Data načtena: <span id="forwardFetchedAt">—</span></p>
    <p class="small" id="forwardFxInfo" hidden></p>
    <p id="forwardError" role="alert" hidden></p>
    <p class="small">OTE · vybraný den: <strong id="oteForwardQuoteDate">—</strong> · Soubor: <span id="oteForwardSource">—</span></p>
    <p class="small" id="oteForwardFxInfo" hidden></p>
    <p class="small" id="oteForwardMissing" hidden></p>
    <p id="oteForwardError" role="alert" hidden></p>
    <div class="chart-panel"><h3 id="forwardChartTitle" class="chart-title">Henry Hub a OTE – Forward</h3><div id="forwardChart" style="height:600px"></div></div>
    <p class="small">Osa X zobrazuje období dodávky, osa Y cenu ve zvolené měně a jednotce. Body OTE leží na začátku příslušného měsíce. Každá úsečka představuje jednu cenu pro celé období; jednotlivé produkty OTE se navzájem nespojují. Chybějící ceny se vynechají.</p>
  </div>

  <div class="panel" id="events-section">
    __EVENT_TABLE__
    <p class="small">Tabulka událostí navrhuje možné přenosové kanály. Časová shoda sama o sobě neprokazuje kauzalitu.</p>
  </div>

  <div class="panel" id="glossary-section">
    <h2>Slovníček statistických pojmů a metod</h2>
    <dl class="glossary">
      <dt>USD/MMBtu a USD/MWh</dt><dd>Dvě jednotky ceny energie. 1 MWh odpovídá 3,4121416 MMBtu. Pro převod ceny z USD/MMBtu na USD/MWh se násobí 3,4121416; opačným směrem se dělí.</dd>
      <dt>LOESS</dt><dd>Lokální vyhlazovací metoda, která odhaduje trend z blízkých pozorování a vytváří hladkou křivku.</dd>
      <dt>Exponenciální trend</dt><dd>Trend odhadnutý exponenciální funkcí. Roční míra vyjadřuje odhadované složené tempo růstu či poklesu, nikoli předpověď.</dd>
      <dt>Strukturální režim</dt><dd>Úsek časové řady s přibližně stálou typickou cenovou úrovní; hranice mezi úseky se označují jako strukturální zlomy.</dd>
      <dt>BIC</dt><dd>Bayesovské informační kritérium, které porovnává vhodnost modelů a penalizuje jejich složitost.</dd>
      <dt>Logaritmický výnos</dt><dd>Rozdíl přirozených logaritmů dvou po sobě jdoucích cen. Používá se k výpočtu volatility a relativních denních změn.</dd>
      <dt>Volatilita</dt><dd>Míra kolísání výnosů, zde založená na směrodatné odchylce denních logaritmických výnosů.</dd>
      <dt>Anualizovaná volatilita</dt><dd>Volatilita přepočtená na roční měřítko násobením směrodatné odchylky denních výnosů odmocninou z 252.</dd>
      <dt>Robustní volatilita MAD</dt><dd>Odhad volatility založený na mediánu absolutních odchylek od mediánu; je méně citlivý na extrémní pozorování.</dd>
      <dt>Volatilita poklesů</dt><dd>Kolísavost počítaná pouze ze záporných denních logaritmických výnosů.</dd>
      <dt>95. percentil absolutní změny</dt><dd>Hranice, pod kterou leží 95 % absolutních denních logaritmických výnosů; používá se jako ukazatel velikosti výraznějších denních pohybů.</dd>
      <dt>MAE</dt><dd>Průměrná absolutní chyba. Zde jde o průměrnou absolutní vzdálenost mezi denní cenou Henry Hub a cenou OTE.</dd>
      <dt>Volatilita rozdílu OTE − HH</dt><dd>30/90/365 dní: směrodatná odchylka úrovně rozdílu v příslušném klouzavém kalendářním okně. Robustní MAD: 1,4826 × medián absolutních odchylek úrovně rozdílu od mediánu za 90 kalendářních dní; vyžaduje alespoň 25 hodnot. Vše je v cenové jednotce bez anualizace. Ukazatele měří rozptýlení úrovně rozdílu OTE − HH, nikoli jeho denních změn.</dd>
      <dt>RMSE</dt><dd>Odmocnina z průměru čtverců rozdílů. Větší chyby penalizuje výrazněji než MAE.</dd>
    </dl>
  </div>
</div>
<script>
const figureSets = __FIGURE_SETS__;
// Plotly mutates supplied objects, so each render receives a separate copy.
let figures = structuredClone(figureSets.USD);
const excelExports = __EXCEL_EXPORTS__;
let globalState = { unit: 'mmbtu', basis: 'nominal', currency: 'USD' };
const marketState = {
  hh: { kind: 'daily', overlays: new Set(['raw', 'ma', 'median', 'loess', 'break']), vols: new Set(['vol30', 'vol90', 'vol365']), events: false },
  ote: { kind: 'daily', overlays: new Set(['raw', 'ma', 'median', 'loess', 'break']), vols: new Set(['vol30', 'vol90', 'vol365']), events: false }
};
document.getElementById('excelExportBtn').addEventListener('click', () => {
  const button = document.getElementById('excelExportBtn');
  button.disabled = true;
  try {
    const {currency, unit, basis} = globalState;
    const encoded = excelExports[`${currency}:${unit}:${basis}`];
    if (!encoded) throw new Error('Pro vybrané parametry nejsou dostupná data.');
    const bytes = Uint8Array.from(atob(encoded), char => char.charCodeAt(0));
    const url = URL.createObjectURL(new Blob([bytes], {
      type: 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    }));
    const link = document.createElement('a');
    link.href = url;
    link.download = `Henry_Hub_OTE_${currency}_${unit === 'mwh' ? 'MWh' : 'MMBtu'}.xlsx`;
    document.body.appendChild(link);
    link.click();
    link.remove();
    setTimeout(() => URL.revokeObjectURL(url), 60000);
  } catch (error) {
    alert('Export do Excelu se nezdařil: ' + error.message);
  } finally {
    button.disabled = false;
  }
});
const comparisonState = { kind: 'daily', overlays: new Set(['hh', 'ote', 'diff']), vols: new Set(['vol30', 'vol90', 'vol365']), events: false };
const chartIds = { hh: 'hhChart', ote: 'oteChart', comparison: 'comparisonChart' };
const eventData = __EVENT_DATA__;
const unitLabels = { mmbtu: 'USD/MMBtu', mwh: 'USD/MWh' };
const basisLabels = { nominal: "nominálně" };

function bindLegend(id) {
  const chart = document.getElementById(chartIds[id]);
  chart.on('plotly_legendclick', event => {
    const meta = chart.data[event.curveNumber].meta;
    if (!meta || meta.kind === 'events') return false;
    const state = id === 'comparison' ? comparisonState : marketState[id];
    const selected = state.kind === 'volatility' ? state.vols : state.overlays;
    selected.has(meta.overlay) ? selected.delete(meta.overlay) : selected.add(meta.overlay);
    if (id === 'comparison') {
      document.querySelectorAll('.comparison-overlay-btn').forEach(b => b.classList.toggle('active', state.overlays.has(b.dataset.comparisonOverlay)));
      document.querySelectorAll('.comparison-vol-btn').forEach(b => b.classList.toggle('active', state.vols.has(b.dataset.comparisonVol)));
      document.querySelectorAll('.comparison-preset-btn').forEach(b => b.classList.remove('active'));
      updateComparisonChart();
    } else {
      document.querySelectorAll(`.overlay-btn[data-market="${id}"]`).forEach(b => b.classList.toggle('active', state.overlays.has(b.dataset.overlay)));
      document.querySelectorAll(`.vol-btn[data-market="${id}"]`).forEach(b => b.classList.toggle('active', state.vols.has(b.dataset.vol)));
      updateMarketChart(id);
    }
    return false;
  });
  chart.on('plotly_legenddoubleclick', () => false);
}
let changingCurrency = false;
function syncHeader() {
  document.querySelectorAll('#reportHeader .unit-btn, #reportHeader .currency-btn').forEach(button => {
    const selected = button.dataset.unit ? button.dataset.unit === globalState.unit : button.dataset.currency === globalState.currency;
    button.classList.toggle('active', selected);
    button.setAttribute('aria-pressed', String(selected));
  });
}
async function applyCurrency(currency) {
  if (changingCurrency || currency === globalState.currency) return;
  changingCurrency = true;
  document.querySelectorAll('.currency-btn, .unit-btn, #excelExportBtn, #pdfExportBtn').forEach(b => b.disabled = true);
  try {
    await chartsReady;
    globalState.currency = currency;
    figures = structuredClone(figureSets[currency]);
    unitLabels.mmbtu = currency + '/MMBtu'; unitLabels.mwh = currency + '/MWh';
    await Promise.all(Object.entries(figures).map(([id, fig]) => Plotly.react(chartIds[id], fig.data, fig.layout, {responsive:true, displaylogo:false})));
    refreshAll();
    await forwardReady;
    syncHeader();
  } finally {
    changingCurrency = false;
    document.querySelectorAll('.currency-btn, .unit-btn, #excelExportBtn, #pdfExportBtn').forEach(b => b.disabled = false);
  }
}
document.querySelectorAll('.currency-btn').forEach(b => b.addEventListener('click', () => applyCurrency(b.dataset.currency)));
const navLinks = [...document.querySelectorAll('#reportHeader nav a')];
function setActiveSection(id) {
  navLinks.forEach(link => {
    const active = link.hash === '#' + id;
    link.classList.toggle('active', active);
    if (active) link.setAttribute('aria-current','location'); else link.removeAttribute('aria-current');
  });
}
let navigatingUntil = 0;
navLinks.forEach(link => link.addEventListener('click', () => {
  navigatingUntil = Date.now() + 1200;
  setActiveSection(link.hash.slice(1));
}));
window.addEventListener('scroll', () => {
  if (Date.now() < navigatingUntil) return;
  const offset = document.getElementById('reportHeader').offsetHeight + 40;
  let current = navLinks[0].hash.slice(1);
  navLinks.forEach(link => { if (document.querySelector(link.hash).getBoundingClientRect().top <= offset) current = link.hash.slice(1); });
  setActiveSection(current);
}, {passive:true});
setActiveSection(location.hash.slice(1) || 'hh-section');
syncHeader();
new ResizeObserver(() => document.documentElement.style.setProperty('--header-height', (document.getElementById('reportHeader').offsetHeight + 16) + 'px')).observe(document.getElementById('reportHeader'));

const chartsReady = Promise.all([
Plotly.newPlot(chartIds.hh, figures.hh.data, figures.hh.layout, { responsive: true, displaylogo: false }).then(() => { bindLegend('hh'); updateMarketChart('hh'); }),
Plotly.newPlot(chartIds.ote, figures.ote.data, figures.ote.layout, { responsive: true, displaylogo: false }).then(() => { bindLegend('ote'); updateMarketChart('ote'); }),
Plotly.newPlot(chartIds.comparison, figures.comparison.data, figures.comparison.layout, { responsive: true, displaylogo: false }).then(() => { bindLegend('comparison'); updateComparisonChart(); })
]);

// Forward is embedded when the report is generated; opening it never downloads data.
const forwardData = __FORWARD_DATA__;
const oteForwardData = __OTE_FORWARD_DATA__;
function forwardTime(value, dateOnly = false) {
  if (!value) return '—';
  const date = new Date(dateOnly ? value + 'T00:00:00' : value);
  return Number.isNaN(date.getTime()) ? value : dateOnly ? date.toLocaleDateString('cs-CZ') : date.toLocaleString('cs-CZ');
}
function forwardRateDescription(data) {
  const rateDate = forwardTime(data.fx_rate_date, true);
  if (data.fx_rate_method === 'neighbor_mean') return 'průměr kurzů z ' + rateDate + ' a ' + forwardTime(data.fx_next_rate_date, true);
  if (data.fx_rate_method === 'last_available') return 'poslední dostupný kurz z ' + rateDate;
  return 'kurz z ' + rateDate;
}
function renderForward() {
  document.getElementById('forwardTradeDate').textContent = forwardTime(forwardData.trade_date, true);
  document.getElementById('forwardFetchedAt').textContent = forwardTime(forwardData.fetched_at);
  document.getElementById('oteForwardQuoteDate').textContent = forwardTime(oteForwardData.quote_date, true);
  document.getElementById('oteForwardSource').textContent = oteForwardData.source_file || 'Forwardy OTE.xlsx';
  const label = unitLabels[globalState.unit];
  const isEur = globalState.currency === 'EUR';
  const fxRate = forwardData.fx_rate;
  const oteFxRate = oteForwardData.fx_rate;
  const missingFx = isEur && forwardData.rows.length > 0 && (!Number.isFinite(fxRate) || fxRate <= 0);
  const missingOteFx = !isEur && oteForwardData.rows.length > 0 && (!Number.isFinite(oteFxRate) || oteFxRate <= 0);
  const errorBox = document.getElementById('forwardError');
  errorBox.hidden = !forwardData.error && !missingFx;
  errorBox.textContent = forwardData.error ? 'Data CME se při vytváření reportu nepodařilo načíst. ' + forwardData.error : missingFx ? 'Pro převod Forwardu CME do EUR chybí platný kurz k datu vypořádání.' : '';
  const oteErrorBox = document.getElementById('oteForwardError');
  oteErrorBox.hidden = !oteForwardData.error && !missingOteFx;
  oteErrorBox.textContent = oteForwardData.error || (missingOteFx ? 'Pro převod forwardů OTE do USD chybí platný kurz k vybranému datu.' : '');
  const fxInfo = document.getElementById('forwardFxInfo');
  fxInfo.hidden = !isEur || !forwardData.rows.length || missingFx;
  fxInfo.textContent = !fxInfo.hidden ? 'CME · převod podle data vypořádání: 1 EUR = ' + fxRate.toLocaleString('cs-CZ', {maximumFractionDigits: 6}) + ' USD (' + forwardRateDescription(forwardData) + ', Federal Reserve).' : '';
  const oteFxInfo = document.getElementById('oteForwardFxInfo');
  oteFxInfo.hidden = isEur || !oteForwardData.rows.length || missingOteFx;
  oteFxInfo.textContent = !oteFxInfo.hidden ? 'OTE · převod podle vybraného dne: 1 EUR = ' + oteFxRate.toLocaleString('cs-CZ', {maximumFractionDigits: 6}) + ' USD (' + forwardRateDescription(oteForwardData) + ', Federal Reserve).' : '';
  const skipped = oteForwardData.skipped_products || [];
  document.getElementById('oteForwardMissing').hidden = !skipped.length;
  document.getElementById('oteForwardMissing').textContent = skipped.length ? 'Ve vybraném řádku chybí ceny: ' + skipped.join(', ') + '. Ceny z jiných dnů se nedoplňují.' : '';
  const rows = missingFx ? [] : (forwardData.rows || []);
  const oteRows = missingOteFx ? [] : (oteForwardData.rows || []);
  const factor = (globalState.unit === 'mwh' ? __MMBTU_PER_MWH__ : 1) / (isEur ? fxRate : 1);
  const oteFactor = (isEur ? 1 : oteFxRate) / (globalState.unit === 'mwh' ? 1 : __MMBTU_PER_MWH__);
  const dates = [];
  const traces = [];
  if (rows.length) {
    traces.push({
      type: 'scatter', mode: 'lines+markers', name: 'Henry Hub · CME', meta: {source: 'cme'},
      x: rows.map(row => row.contract_date), y: rows.map(row => row.settle * factor),
      customdata: rows.map(row => row.month),
      line: {color: '#832C29', width: 2}, marker: {size: 5},
      hovertemplate: 'CME · %{customdata}<br>Settle: %{y:.3f} ' + label + '<extra></extra>'
    });
    dates.push(...rows.map(row => row.contract_date));
  }
  const months = oteRows.filter(row => row.kind === 'month');
  if (months.length) traces.push({
    type: 'scatter', mode: 'markers', name: 'OTE · měsíce', meta: {source: 'ote', kind: 'month'},
    x: months.map(row => row.delivery_start), y: months.map(row => row.price * oteFactor),
    customdata: months.map(row => row.product), marker: {color: '#1565c0', size: 10},
    hovertemplate: 'OTE · %{customdata}<br>Cena: %{y:.3f} ' + label + '<extra></extra>'
  });
  for (const [kind, name, color] of [['quarter', 'OTE · kvartály', '#00897b'], ['calendar', 'OTE · roky (Cal)', '#6d4c9c']]) {
    oteRows.filter(row => row.kind === kind).forEach((row, index) => {
      const endDay = new Date(row.delivery_end + 'T12:00:00');
      endDay.setDate(endDay.getDate() - 1);
      const period = forwardTime(row.delivery_start, true) + ' – ' + endDay.toLocaleDateString('cs-CZ');
      traces.push({
        type: 'scatter', mode: 'lines', name, meta: {source: 'ote', kind, product: row.product},
        legendgroup: 'ote-' + kind, showlegend: index === 0,
        x: [row.delivery_start, row.delivery_end], y: [row.price * oteFactor, row.price * oteFactor],
        customdata: [[row.product, period], [row.product, period]], line: {color, width: 4},
        hovertemplate: 'OTE · %{customdata[0]}<br>%{customdata[1]}<br>Cena: %{y:.3f} ' + label + '<extra></extra>'
      });
    });
  }
  for (const row of oteRows) dates.push(row.delivery_start, row.kind === 'month' ? row.delivery_start : row.delivery_end);
  dates.sort();
  const first = dates.length ? new Date(dates[0]) : null;
  const last = dates.length ? new Date(dates[dates.length - 1]) : null;
  const spanMonths = first && last ? (last.getUTCFullYear() - first.getUTCFullYear()) * 12 + last.getUTCMonth() - first.getUTCMonth() + 1 : 1;
  const notes = [];
  if (rows.length) notes.push('CME: ' + forwardTime(forwardData.trade_date, true));
  if (oteRows.length) notes.push('OTE: ' + forwardTime(oteForwardData.quote_date, true));
  document.getElementById('forwardChartTitle').textContent = 'Henry Hub a OTE – Forward (' + label + ') · ' + (notes.join(' · ') || 'Data nejsou k dispozici');
  return Plotly.react('forwardChart', traces, {
    template: 'plotly_white', margin: {l: 70, r: 30, t: 55, b: 85},
    xaxis: {title: {text: 'Období dodávky'}, type: 'date', tickformat: '%b %Y', tickangle: -45, tick0: dates[0], dtick: 'M' + Math.max(1, Math.ceil(spanMonths / (window.innerWidth < 700 ? 6 : 18)))},
    yaxis: {title: {text: 'Cena (' + label + ')'}, tickformat: '.3f'},
    uirevision: 'forward:' + globalState.currency + ':' + globalState.unit,
    showlegend: true, legend: {orientation: 'h', x: 0, y: 1.12, groupclick: 'togglegroup'},
    annotations: traces.length ? [] : [{text: 'Data nejsou k dispozici', x: .5, y: .5, xref: 'paper', yref: 'paper', showarrow: false}]
  }, {responsive: true, displaylogo: false});
}
let forwardReady = renderForward();

function addEventShapes(fig, enabled) {
  const shapes = (fig.layout.shapes || []).filter(shape => shape.name !== 'event');
  if (enabled) {
    eventData.forEach(event => shapes.push({
      type: 'line', x0: event.date, x1: event.date, y0: 0, y1: 0.78,
      xref: 'x', yref: 'paper', name: 'event', layer: 'below',
      line: { color: 'rgba(0,0,0,.22)', width: 1 }
    }));
  }
  return shapes;
}

// Epochs and historic events must not extend OTE/comparison into years without data.
const timeRangeViews = {};
function dataTimeRange(id, state, fig) {
  if (id === 'hh') return {};
  const view = globalState.currency + ':' + state.kind;
  if (timeRangeViews[id] === view) return {}; // Preserve manual zoom when toggling curves.
  timeRangeViews[id] = view;
  let first = null, last = null;
  for (const trace of fig.data) {
    const meta = trace.meta || {};
    if (meta.kind !== state.kind || meta.unit !== globalState.unit || meta.basis !== globalState.basis) continue;
    (trace.x || []).forEach((date, i) => {
      const value = trace.y[i];
      if (value == null || !Number.isFinite(Number(value))) return;
      const time = new Date(date).getTime();
      if (!Number.isFinite(time)) return;
      first = first === null ? time : Math.min(first, time);
      last = last === null ? time : Math.max(last, time);
    });
  }
  return first === null ? {} : {'xaxis.range': [new Date(first).toISOString(), new Date(last).toISOString()], 'xaxis.autorange': false};
}

function updateMarketChart(id) {
  const state = marketState[id];
  const fig = figures[id];
  const visible = fig.data.map(trace => {
    const meta = trace.meta || {};
    if (meta.kind === 'events') return state.events;
    if (state.kind === 'volatility') {
      return meta.kind === 'volatility' && meta.unit === globalState.unit && meta.basis === globalState.basis ? (state.vols.has(meta.overlay) ? true : 'legendonly') : false;
    }
    return meta.kind === state.kind && meta.unit === globalState.unit && meta.basis === globalState.basis ? (state.overlays.has(meta.overlay) ? true : 'legendonly') : false;
  });
  const titlePrefix = id === 'hh' ? 'Henry Hub' : 'OTE';
  const viewTitle = state.kind === 'volatility' ? 'volatilita' : (state.kind === 'daily' ? 'denní ceny' : 'měsíční ceny');
  document.getElementById(id + 'ChartTitle').textContent = `${titlePrefix}: ${viewTitle} (${unitLabels[globalState.unit]}, ${basisLabels[globalState.basis]})`;
  Plotly.restyle(chartIds[id], { visible });
  Plotly.relayout(chartIds[id], {
    ...dataTimeRange(id, state, fig),
    'yaxis.title.text': state.kind === 'volatility' ? 'Volatilita / změna' : unitLabels[globalState.unit],
    shapes: addEventShapes(fig, state.events)
  });
}

function updateComparisonChart() {
  const state = comparisonState;
  const fig = figures.comparison;
  const visible = fig.data.map(trace => {
    const meta = trace.meta || {};
    if (meta.kind === 'events') return state.events;
    if (state.kind === 'volatility') {
      return meta.kind === 'volatility' && meta.unit === globalState.unit && meta.basis === globalState.basis ? (state.vols.has(meta.overlay) ? true : 'legendonly') : false;
    }
    return meta.kind === state.kind && meta.unit === globalState.unit && meta.basis === globalState.basis ? (state.overlays.has(meta.overlay) ? true : 'legendonly') : false;
  });
  const titleKind = state.kind === 'volatility' ? 'volatilita rozdílu' : (state.kind === 'daily' ? 'denní ceny' : 'měsíční ceny');
  document.getElementById('comparisonChartTitle').textContent = `Porovnání Henry Hub a OTE – ${titleKind} (${unitLabels[globalState.unit]}, ${basisLabels[globalState.basis]})`;

  const shapes = (fig.layout.shapes || []).filter(shape => shape.name !== 'comparison-event');
  if (state.events) {
    eventData.forEach(event => shapes.push({
      type: 'line', x0: event.date, x1: event.date, y0: 0, y1: 0.78,
      xref: 'x', yref: 'paper', name: 'comparison-event', layer: 'below',
      line: { color: 'rgba(0,0,0,.22)', width: 1 }
    }));
  }
  Plotly.restyle(chartIds.comparison, { visible });
  Plotly.relayout(chartIds.comparison, {
    ...dataTimeRange("comparison", state, fig),
    'yaxis.title.text': state.kind === 'volatility' ? `Volatilita rozdílu (${unitLabels[globalState.unit]})` : `Cena (${unitLabels[globalState.unit]})`,
    'yaxis2.title.text': state.kind === 'volatility' ? '' : `Rozdíl OTE − HH (${unitLabels[globalState.unit]})`,
    'yaxis.visible': true,
    'yaxis2.visible': state.kind !== 'volatility',
    shapes: shapes
  });
}

function updateAllTables() {
  document.querySelectorAll('.currency-variant').forEach(el => el.classList.toggle('active', el.dataset.currency === globalState.currency));
  document.querySelectorAll('.table-variant').forEach(table => {
    table.classList.toggle('active', table.dataset.unit === globalState.unit && table.dataset.basis === globalState.basis);
  });
}

function refreshAll() {
  updateMarketChart('hh');
  updateMarketChart('ote');
  updateComparisonChart();
  updateAllTables();
  forwardReady = forwardReady.then(renderForward);
}

document.querySelectorAll('.unit-btn').forEach(button => button.addEventListener('click', () => {
  globalState.unit = button.dataset.unit;
  syncHeader();
  document.querySelectorAll('.unit-btn').forEach(item => item.classList.toggle('active', item === button));
  refreshAll();
}));
document.querySelectorAll('.kind-btn').forEach(button => button.addEventListener('click', () => {
  const state = marketState[button.dataset.market];
  state.kind = button.dataset.kind;
  document.querySelectorAll(`.kind-btn[data-market="${button.dataset.market}"]`).forEach(item => item.classList.toggle('active', item === button));
  updateMarketChart(button.dataset.market);
}));
document.querySelectorAll('.overlay-btn').forEach(button => button.addEventListener('click', () => {
  const state = marketState[button.dataset.market];
  const overlay = button.dataset.overlay;
  button.classList.toggle('active');
  button.classList.contains('active') ? state.overlays.add(overlay) : state.overlays.delete(overlay);
  updateMarketChart(button.dataset.market);
}));
document.querySelectorAll('.vol-btn').forEach(button => button.addEventListener('click', () => {
  const state = marketState[button.dataset.market];
  const vol = button.dataset.vol;
  button.classList.toggle('active');
  button.classList.contains('active') ? state.vols.add(vol) : state.vols.delete(vol);
  updateMarketChart(button.dataset.market);
}));
document.querySelectorAll('.events-btn').forEach(button => button.addEventListener('click', () => {
  const state = marketState[button.dataset.market];
  state.events = !state.events;
  button.classList.toggle('active', state.events);
  updateMarketChart(button.dataset.market);
}));

document.querySelectorAll('.comparison-kind-btn').forEach(button => button.addEventListener('click', () => {
  comparisonState.kind = button.dataset.comparisonKind;
  document.querySelectorAll('.comparison-kind-btn').forEach(item => item.classList.toggle('active', item === button));
  updateComparisonChart();
}));
document.querySelectorAll('.comparison-overlay-btn').forEach(button => button.addEventListener('click', () => {
  const overlay = button.dataset.comparisonOverlay;
  button.classList.toggle('active');
  button.classList.contains('active') ? comparisonState.overlays.add(overlay) : comparisonState.overlays.delete(overlay);
  document.querySelectorAll('.comparison-preset-btn').forEach(item => item.classList.remove('active'));
  updateComparisonChart();
}));
document.querySelectorAll('.comparison-preset-btn').forEach(button => button.addEventListener('click', () => {
  const preset = button.dataset.comparisonPreset;
  const sets = {
    prices: ['hh', 'ote'],
    difference: ['diff', 'diff_ma', 'diff_median'],
    all: ['hh', 'ote', 'diff', 'diff_ma', 'diff_median', 'diff_loess', 'diff_exp', 'diff_break']
  };
  comparisonState.overlays = new Set(sets[preset]);
  document.querySelectorAll('.comparison-overlay-btn').forEach(item => item.classList.toggle('active', comparisonState.overlays.has(item.dataset.comparisonOverlay)));
  document.querySelectorAll('.comparison-preset-btn').forEach(item => item.classList.toggle('active', item === button));
  if (preset === 'prices' || preset === 'difference' || preset === 'all') {
    comparisonState.kind = comparisonState.kind === 'volatility' ? 'daily' : comparisonState.kind;
    document.querySelectorAll('.comparison-kind-btn').forEach(item => item.classList.toggle('active', item.dataset.comparisonKind === comparisonState.kind));
  }
  updateComparisonChart();
}));
document.querySelectorAll('.comparison-vol-btn').forEach(button => button.addEventListener('click', () => {
  const overlay = button.dataset.comparisonVol;
  button.classList.toggle('active');
  button.classList.contains('active') ? comparisonState.vols.add(overlay) : comparisonState.vols.delete(overlay);
  updateComparisonChart();
}));
document.getElementById('comparisonEventsBtn').addEventListener('click', event => {
  comparisonState.events = !comparisonState.events;
  event.currentTarget.classList.toggle('active', comparisonState.events);
  updateComparisonChart();
});

// Export is fully offline: Plotly snapshots + browser print-to-PDF.
const pdfDialog = document.getElementById('pdfDialog');
let pdfItems = [];
let pdfBusy = false;
function preparePdfChoices() {
  const target = document.getElementById('pdfChoices');
  target.replaceChildren(); pdfItems = [];
  function group(title, items) {
    const fieldset = document.createElement('fieldset');
    const legend = document.createElement('legend'); legend.textContent = title;
    fieldset.append(legend);
    items.forEach(item => {
      const index = pdfItems.push(item) - 1;
      const label = document.createElement('label');
      const input = document.createElement('input'); input.type = 'checkbox'; input.checked = true; input.value = String(index);
      const caption = document.createElement('span'); caption.textContent = item.title;
      label.append(input, caption); fieldset.append(label);
    });
    target.append(fieldset);
  }
  for (const [id, title] of [['hh','Henry Hub'], ['ote','OTE'], ['comparison','Porovnání Henry Hub a OTE']]) {
    const section = document.getElementById(id + '-section');
    const items = [{type:'chart', id:chartIds[id], title:document.getElementById(id + 'ChartTitle').textContent}];
    section.querySelectorAll('.currency-variant.active .table-variant.active table').forEach(table => {
      items.push({type:'table', node:table, title:title + ' – ' + (table.closest('.data-table-scroll')?.previousElementSibling?.textContent || table.previousElementSibling?.textContent || 'Tabulka')});
    });
    group(title, items);
  }
  if (forwardData.rows.length || oteForwardData.rows.length) group('Forward · ' + unitLabels[globalState.unit], [{type:'chart', id:'forwardChart', title:document.getElementById('forwardChartTitle').textContent}]);
  group('Další části', [
    {type:'section', node:document.getElementById('events-section'), title:'Události s potenciálním vlivem na cenu'},
    {type:'section', node:document.getElementById('glossary-section'), title:'Slovníček statistických pojmů a metod'}
  ]);
}
document.getElementById('pdfExportBtn').addEventListener('click', async () => {
  await Promise.all([chartsReady, forwardReady]);
  preparePdfChoices(); document.getElementById('pdfStatus').textContent = ''; pdfDialog.showModal();
});
document.getElementById('pdfCancel').addEventListener('click', () => pdfDialog.close());
for (const [id, checked] of [['pdfSelectAll',true], ['pdfSelectNone',false]]) {
  document.getElementById(id).addEventListener('click', () => document.querySelectorAll('#pdfChoices input').forEach(input => input.checked = checked));
}
pdfDialog.addEventListener('cancel', event => { if (pdfBusy) event.preventDefault(); });
// Tiskové stránky se skládají přímo ve formátu A4 na šířku.
// Díky tomu se obrázek loga NEPOZICUJE pomocí position:fixed nebo záporného bottom.
// Všechny stránky mají vlastní logo a tabulky se dělí pouze mezi řádky.
function buildPdfSheet(doc, pages) {
  const sheet = doc.createElement('article');
  sheet.className = 'pdf-sheet';
  const body = doc.createElement('div'); body.className = 'pdf-sheet-body';
  const footer = doc.createElement('div'); footer.className = 'pdf-sheet-footer';
  footer.textContent = 'Henry Hub a OTE · Autor: Marek Božoň, BSc s využitím AI.';
  const number = doc.createElement('span'); number.className = 'pdf-sheet-number';
  number.textContent = String(pages.length + 1);
  const logo = doc.createElement('img');
  logo.className = 'pdf-sheet-logo'; logo.alt = 'LAMA energy'; logo.src = '__PRINT_LOGO__';
  sheet.append(body, footer, number, logo);
  doc.body.append(sheet);
  const page = {sheet, body}; pages.push(page);
  return page;
}

function buildPdfPages(doc, selected) {
  const pages = [];
  let current = buildPdfSheet(doc, pages);
  const fits = () => current.body.scrollHeight <= current.body.clientHeight + 1;
  const nextPage = () => { current = buildPdfSheet(doc, pages); };

  // Nerozdělitelný prvek se umístí na dosavadní stranu, jen pokud se tam vejde.
  function addBlock(block) {
    const hadContent = current.body.children.length > 0;
    current.body.append(block);
    if (!fits()) {
      block.remove();
      if (hadContent) nextPage();
      current.body.append(block);
      if (!fits()) {
        throw new Error('Obsah se nevejde na stránku A4: ' +
          (block.querySelector('h2')?.textContent || block.className));
      }
    }
  }

  function addTable(title, sourceTable) {
    const rows = [...sourceTable.querySelectorAll('tbody > tr')];
    const makePart = continuing => {
      const wrapper = doc.createElement('section'); wrapper.className = 'pdf-table-block';
      const h = doc.createElement('h2');
      h.textContent = title + (continuing ? ' (pokračování)' : '');
      const table = doc.createElement('table'); table.className = 'pdf-table';
      const cols = sourceTable.querySelector('colgroup');
      if (cols) table.append(doc.importNode(cols, true));
      if (sourceTable.tHead) table.append(doc.importNode(sourceTable.tHead, true));
      const tbody = doc.createElement('tbody'); table.append(tbody);
      wrapper.append(h, table);
      return {wrapper, tbody};
    };
    let part = makePart(false);
    for (let i = 0; i < rows.length; ++i) {
      const row = doc.importNode(rows[i], true);
      part.tbody.append(row);
      if (!part.wrapper.isConnected) current.body.append(part.wrapper);
      if (fits()) continue;
      row.remove();
      // Zbývá málo prostoru i na záhlaví tabulky a první řádek.
      if (part.tbody.rows.length === 0) {
        part.wrapper.remove();
        nextPage();
        current.body.append(part.wrapper);
        part.tbody.append(row);
        if (!fits()) throw new Error('Příliš vysoký řádek v tabulce: ' + title);
      } else {
        nextPage();
        part = makePart(true);
        part.tbody.append(row);
        current.body.append(part.wrapper);
        if (!fits()) throw new Error('Příliš vysoký řádek v tabulce: ' + title);
      }
    }
    if (!rows.length) addBlock(part.wrapper);
  }

  // Titulní hlavička a PRVNÍ GRAF sdílejí první stránku.
  const brand = doc.createElement('header'); brand.className = 'pdf-brand';
  const topLogo = doc.createElement('img');
  topLogo.className = 'pdf-brand-logo'; topLogo.src = '__PRINT_LOGO__'; topLogo.alt = 'LAMA energy';
  const heading = doc.createElement('h1'); heading.textContent = 'Henry Hub a OTE – spotové ceny plynu';
  brand.append(topLogo, heading);
  const settings = doc.createElement('p'); settings.className = 'pdf-settings';
  settings.textContent = `Všechny sekce: ${unitLabels[globalState.unit]} · Nominální ceny · Rozdíl: OTE − HH`;
  current.body.append(brand, settings);

  // Grafy se tisknou jako první. Následující tabulky pak plynule vyplní
  // zbylé stránky; graf mezi tabulkami nevytváří polovinu prázdné A4.
  const printOrder = [...selected.filter(item => item.type === 'chart'),
                      ...selected.filter(item => item.type !== 'chart')];
  for (const item of printOrder) {
    if (item.type === 'chart') {
      const block = doc.createElement('section'); block.className = 'pdf-chart';
      const h = doc.createElement('h2'); h.textContent = item.title;
      const img = doc.createElement('img'); img.alt = item.title; img.src = item.png;
      block.append(h, img); addBlock(block);
    } else if (item.type === 'table') {
      addTable(item.title, item.node);
    } else if (item.type === 'section') {
      const clone = item.node;
      const table = clone.querySelector('table');
      if (table) {
        addTable(item.title, table);
        clone.querySelectorAll('p').forEach(p => {
          const block = doc.createElement('p'); block.className = 'pdf-note';
          block.textContent = p.textContent; addBlock(block);
        });
      } else {
        const entries = [...clone.querySelectorAll('dl > dt')];
        let first = true;
        for (const dt of entries) {
          const pair = doc.createElement('section'); pair.className = 'pdf-glossary-item';
          if (first) {
            const h = doc.createElement('h2'); h.textContent = item.title; pair.append(h);
          }
          const title = doc.createElement('strong'); title.textContent = dt.textContent;
          const def = doc.createElement('p'); def.textContent = dt.nextElementSibling?.textContent || '';
          pair.append(title, def); addBlock(pair);
          first = false;
        }
      }
    }
  }
  return pages;
}

document.getElementById('pdfConfirm').addEventListener('click', async () => {
  if (pdfBusy) return;
  const selected = [...document.querySelectorAll('#pdfChoices input:checked')].map(input => pdfItems[Number(input.value)]);
  const status = document.getElementById('pdfStatus');
  if (!selected.length) { status.textContent = 'Zaškrtněte alespoň jednu položku.'; return; }
  pdfBusy = true;
  pdfDialog.querySelectorAll('button, input').forEach(el => el.disabled = true);
  let frame;
  try {
    await Promise.all([chartsReady, forwardReady]);
    status.textContent = 'Připravuji vybrané části pro PDF…';
    // PNG snímek grafu respektuje právě nastavené řady, zoom i měnu.
    for (const item of selected) {
      if (item.type === 'chart') {
        item.png = await Plotly.toImage(document.getElementById(item.id),
          {format: 'png', width: 1400, height: 700, scale: 2});
      }
    }
    document.getElementById('pdfPrintFrame')?.remove();
    frame = document.createElement('iframe');
    frame.id = 'pdfPrintFrame'; frame.title = 'Podklady pro tisk PDF';
    frame.style.cssText = 'position:fixed;left:-12000px;top:0;width:1123px;height:794px;border:0;';
    document.body.append(frame);
    const doc = frame.contentDocument;
    doc.open();
    doc.write(`<!doctype html><html lang="cs"><head><meta charset="utf-8">
      <title>Henry Hub a OTE – report</title>
      <style>
        @page { size: A4 landscape; margin: 0; }
        * { box-sizing: border-box; }
        html,body { padding:0; margin:0; background:white; }
        body { font-family:Arial,sans-serif; color:#18202a; font-size:10pt;
          -webkit-print-color-adjust:exact; print-color-adjust:exact; }
        .pdf-sheet { position:relative; width:297mm; height:209.8mm;
          padding:12mm 12mm 19mm; overflow:hidden; background:white;
          break-after:page; page-break-after:always; }
        .pdf-sheet:last-child { break-after:auto; page-break-after:auto; }
        .pdf-sheet-body { height:100%; overflow:hidden; }
        .pdf-sheet-logo { position:absolute; right:12mm; bottom:3mm;
          display:block; width:26mm; height:12.7mm; object-fit:contain; }
        .pdf-sheet-footer { position:absolute; left:12mm; bottom:6mm;
          max-width:170mm; color:#697586; font-size:8pt; }
        .pdf-sheet-number { position:absolute; bottom:6mm; right:42mm;
          color:#697586; font-size:8pt; }
        .pdf-brand { height:22mm; display:flex; align-items:center; gap:7mm;
          border-bottom:1px solid #bec5ce; padding-bottom:2mm; margin:0 0 2mm; }
        .pdf-brand-logo { display:block; width:38mm; height:18mm; object-fit:contain; }
        h1 { font-size:18pt; margin:0; line-height:1.2; }
        h2 { font-size:11.5pt; line-height:1.25; margin:0 0 2mm; }
        .pdf-settings { margin:0 0 3mm; line-height:1.3; font-size:9pt; }
        .pdf-chart { margin:0 0 2mm; }
        .pdf-chart img { display:block; width:100%; height:136.5mm;
          object-fit:contain; }
        .pdf-table-block { margin:0 0 3mm; }
        .pdf-table { border-collapse:collapse; width:100%; table-layout:fixed;
          font-size:7.6pt; line-height:1.25; }
        .pdf-table th, .pdf-table td { padding:1.7mm 1.5mm;
          border-bottom:.25mm solid #d8dde3;
          white-space:normal!important; overflow-wrap:anywhere; vertical-align:top; }
        .pdf-table th { background:#e9edf2; text-align:left; }
        .pdf-note { font-size:8.5pt; margin:2mm 0 3mm; color:#4c5563; }
        .pdf-glossary-item { margin:0 0 2.5mm; font-size:9pt; line-height:1.35; }
        .pdf-glossary-item strong { font-size:9pt; }
        .pdf-glossary-item p { margin:1mm 0 0; }
        @media print { .pdf-sheet { outline:0; } }
      </style></head><body></body></html>`);
    doc.close();
    if (doc.fonts?.ready) await doc.fonts.ready;
    const pages = buildPdfPages(doc, selected);
    await Promise.all([...doc.images].map(img => img.decode()));
    await new Promise(resolve => frame.contentWindow.requestAnimationFrame(() => frame.contentWindow.requestAnimationFrame(resolve)));
    // Ověříme znovu po dekódování obrázků (změna velikosti nemá posunout obsah).
    for (const page of pages) {
      if (page.body.scrollHeight > page.body.clientHeight + 1)
        throw new Error('Stránka ' + (pages.indexOf(page) + 1) + ' překračuje vyhrazenou oblast.');
    }
    pdfDialog.close();
    frame.contentWindow.focus(); frame.contentWindow.print();
    // Iframe musí zůstat v DOM, dokud prohlížeč nedokončí tisk.
  } catch (error) {
    frame?.remove();
    status.textContent = 'Export se nepodařilo připravit: ' + error.message;
    if (!pdfDialog.open) pdfDialog.showModal();
  } finally {
    pdfBusy = false;
    pdfDialog.querySelectorAll('button, input').forEach(el => el.disabled = false);
  }
});

</script>
<footer class="report-footer">Autor: Marek Božoň, BSc s využitím AI.</footer>
</body></html>'''

    replacements = {
        "__MMBTU_PER_MWH__": json.dumps(MMBTU_PER_MWH),
        "__FORWARD_DATA__": json.dumps(forward_data, ensure_ascii=False, allow_nan=False).replace("<", "\\u003c"),
        "__OTE_FORWARD_DATA__": json.dumps(ote_forward_data, ensure_ascii=False, allow_nan=False).replace("<", "\\u003c"),
        "__EXCEL_EXPORTS__": json.dumps(excel_exports),
        "__LOGO__": LOGO_DATA_URI,
        "__PRINT_LOGO__": PRINT_LOGO_DATA_URI,
        "__FX_END__": eur["fx_end"],
        "__OTE_END__": ote_prepared["daily"]["date"].max().strftime("%d.%m.%Y"),
        "__FIGURE_SETS__": json.dumps(figure_sets, cls=PlotlyJSONEncoder),
        "__PLOTLY_JS__": get_plotlyjs(),
        "__HH_KPI__": "".join(kpis["hh"]),
        "__OTE_KPI__": "".join(kpis["ote"]),
        "__HH_TABLES__": hh_tables,
        "__OTE_TABLES__": ote_tables,
        "__COMPARISON_TABLES__": compare_tables,
        "__EVENT_TABLE__": event_table,
        "__EVENT_DATA__": json.dumps(
            [{"date": row[0], "name": row[1], "category": row[2]} for row in EVENTS],
            ensure_ascii=False,
        ),
    }
    for marker, value in replacements.items():
        html = html.replace(marker, value)
    output.write_text(html, encoding="utf-8")


def resolve_input(name: str) -> Path:
    path = Path(name).expanduser()
    if path.is_file():
        return path
    # Notebook nemá __file__; vstupy hledáme i v pracovním adresáři.
    roots = [Path.cwd(), Path.cwd() / "upload"]
    script_file = globals().get("__file__")
    if script_file:
        roots.extend([Path(script_file).resolve().parent,
                      Path(script_file).resolve().parent / "upload"])
    names = [name]
    if name == INPUT_OTE_DEFAULT:
        names.append("spotové ceny OTE - plyn.xlsx")
    for root in roots:
        for candidate in names:
            if (root / candidate).is_file():
                return root / candidate
    searched = ", ".join(str(root) for root in dict.fromkeys(roots))
    raise FileNotFoundError(
        f"Vstupní soubor nebyl nalezen: {name}. Hledané složky: {searched}. "
        "V Jupyteru/Colabu nahrajte Excel do pracovního adresáře "
        "nebo nastavte příslušnou konstantu INPUT_OTE_DEFAULT / INPUT_OTE_FORWARD_DEFAULT na jeho cestu."
    )


def load_ote_forward_for_report(name: str = INPUT_OTE_FORWARD_DEFAULT) -> dict[str, Any]:
    """Read the local quotation file once; retain the other report sections on failure."""
    try:
        data = read_ote_forward(resolve_input(name))
        return {**data, "error": None}
    except Exception as exc:
        return {"rows": [], "quote_date": None, "fetched_at": None,
                "source_file": Path(name).name, "source_currency": "EUR", "source_unit": "mwh",
                "skipped_products": [], "error": "Forwardy OTE se nepodařilo načíst. " + str(exc)}


def attach_forward_exchange_rate(data: dict[str, Any], rates: pd.DataFrame, date_field: str) -> None:
    """Each source uses its own quotation day, never a delivery period's day."""
    if not data["rows"]:
        return
    resolved = resolve_exchange_rates(pd.Series([pd.Timestamp(data[date_field])]), rates).iloc[0]
    data.update(
        fx_rate=float(resolved["usd_per_eur"]),
        fx_rate_date=resolved["rate_date"].strftime("%Y-%m-%d"),
        fx_next_rate_date=resolved["next_rate_date"].strftime("%Y-%m-%d") if pd.notna(resolved["next_rate_date"]) else None,
        fx_rate_method=resolved["rate_method"],
    )


def main() -> None:
    parser = argparse.ArgumentParser(description="Interaktivní analýza spotových cen zemního plynu Henry Hub a OTE")
    parser.add_argument("--ote-input", default=INPUT_OTE_DEFAULT, help="Vstupní Excel soubor s cenami OTE")
    parser.add_argument("--ote-forward-input", default=INPUT_OTE_FORWARD_DEFAULT, help="Excel s forwardy OTE v EUR/MWh")
    parser.add_argument("--output", default=OUTPUT_DEFAULT, help="Výstupní HTML soubor")
    args = parser.parse_args([]) if "ipykernel" in sys.modules else parser.parse_args()

    ote_path = resolve_input(args.ote_input)
    print("Doplňuji původní Excel OTE z ročních zpráv (Europe/Prague)...")
    ote_update = update_ote_workbook(ote_path)
    for note in ote_update["notes"]:
        print("OTE: " + note)
    last_ote = pd.Timestamp(ote_update["last_date"]).strftime("%d.%m.%Y") if ote_update["last_date"] else "bez dat"
    if ote_update["added_rows"] or ote_update["updated_rows"]:
        print(f"OTE: přidáno {ote_update['added_rows']} řádků, doplněno {ote_update['updated_rows']} připravených řádků; data do {last_ote}.")
        print(f"Záloha původního OTE souboru: {ote_update['backup_path']}")
    else:
        print(f"OTE: nejsou dostupná novější data; soubor nebyl změněn. Poslední den: {last_ote}.")
    print("Stahuji denní spotové ceny Henry Hub z EIA (list Data 1)...")
    hh_prices = read_henry_hub()
    ote_prices = read_ote(ote_path)

    print("Připravuji data Henry Hub...")
    hh_prepared = prepare_market(hh_prices)
    print("Připravuji data OTE...")
    ote_prepared = prepare_market(ote_prices)
    comparison = build_comparison_frame(hh_prepared, ote_prepared)

    hh_figure = build_market_figure(hh_prepared, "Henry Hub")
    ote_figure = build_market_figure(ote_prepared, "OTE")
    comparison_figure = build_comparison_figure(comparison)
    print("Načítám historické kurzy USD za 1 EUR z Federal Reserve...")
    rates = read_exchange_rates()
    print("Připravuji EUR podle historických kurzů...")
    eur_hh = prepare_market(convert_to_eur(hh_prices, rates))
    eur_ote = prepare_market(convert_to_eur(ote_prices, rates))
    eur_comparison = build_comparison_frame(eur_hh, eur_ote)
    eur = {"fx_end": rates["date"].max().strftime("%d.%m.%Y"), "hh": eur_hh, "ote": eur_ote, "comparison": eur_comparison,
           "figures": {"hh": build_market_figure(eur_hh, "Henry Hub"),
                       "ote": build_market_figure(eur_ote, "OTE"),
                       "comparison": build_comparison_figure(eur_comparison)}}
    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    print("Načítám aktuální forwardovou křivku CME…")
    forward_data = load_forward_for_report()
    attach_forward_exchange_rate(forward_data, rates, "trade_date")
    if forward_data["error"]:
        print("Forward: " + forward_data["error"])
    print("Načítám forwardy OTE z Excelu (EUR/MWh)…")
    ote_forward_data = load_ote_forward_for_report(args.ote_forward_input)
    attach_forward_exchange_rate(ote_forward_data, rates, "quote_date")
    if ote_forward_data["error"]:
        print(ote_forward_data["error"])
    else:
        print(f"Forward OTE: vybrán {pd.Timestamp(ote_forward_data['quote_date']):%d.%m.%Y}, {len(ote_forward_data['rows'])} cen.")
    make_html(hh_figure, ote_figure, comparison_figure, hh_prepared, ote_prepared, comparison, output, eur, forward_data, ote_forward_data)
    print(f"Hotovo: {output.resolve()}")
    print(f"Henry Hub: {len(hh_prepared['daily']):,} denních záznamů".replace(",", " "))
    print(f"OTE: {len(ote_prepared['daily']):,} denních záznamů".replace(",", " "))
    print(f"Společné dny pro porovnání: {len(comparison):,}".replace(",", " "))



if __name__ == "__main__":
    main()
