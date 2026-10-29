"""Exercise discovery, real ZIP/XLSX parsing, CNB conversion and persistence together."""
from datetime import date, timedelta
from email.message import Message
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
import importlib.util
import unittest
from zipfile import ZipFile

import pandas as pd
from openpyxl import Workbook, load_workbook

spec=importlib.util.spec_from_file_location('report_ote_pipeline',Path(__file__).resolve().parents[1]/'henry_hub_ote_forward.py')
m=importlib.util.module_from_spec(spec)
spec.loader.exec_module(m)


def report_zip(year,rows):
    book=Workbook()
    ws=book.active
    ws.title='Odchylky NC BAL'
    ws.append(['Plynárenský den','Použitelná cena pro záporné vyrovnávací množství (Kč/MWh)'])
    for day,amount in rows:ws.append([pd.Timestamp(day).to_pydatetime(),amount])
    workbook=BytesIO()
    book.save(workbook)
    book.close()
    result=BytesIO()
    with ZipFile(result,'w') as archive:
        archive.writestr(f'Rocni_zprava_o_trhu_plyn_{year}_V0.xlsx',workbook.getvalue())
    return result.getvalue()


class OtePipelineTests(unittest.TestCase):
    def test_real_two_year_pipeline_and_repeated_update_use_no_future_fx_or_duplicate_write(self):
        pages={year:m.OTE_ANNUAL_SOURCE_TEMPLATE.format(year=year) for year in (2025,2026)}
        zips={year:f'https://www.ote-cr.cz/pubweb/attachments/62_162/{year}/Rocni_zprava_o_trhu_{year}_V0_plyn.zip' for year in pages}
        responses={}
        for year in pages:
            responses[pages[year]]=(f'<h1>Roční zpráva</h1><a href="{zips[year]}">Plyn {year} V0</a>'.encode(),'text/html')
        responses[zips[2025]]=(report_zip(2025,[('2025-12-30',9999),('2025-12-31',500)]),'application/zip')
        responses[zips[2026]]=(report_zip(2026,[('2026-01-01',600),('2026-01-02',700),('2026-01-07',9999)]),'application/zip')
        responses[m.CNB_YEARLY_RATE_TEMPLATE.format(year=2025)]=(b'Datum|1 USD\n31.12.2025|25,0\n','text/plain')
        responses[m.CNB_YEARLY_RATE_TEMPLATE.format(year=2026)]=(b'Datum|1 USD\n02.01.2026|20,0\n','text/plain')
        class Response:
            def __init__(self,body,kind):
                self.body=body
                self.headers=Message()
                self.headers['Content-Type']=kind+'; charset=utf-8'
            def __enter__(self):return self
            def __exit__(self,*args):pass
            def read(self,limit):
                if len(self.body)>limit:raise AssertionError('Unbounded request')
                return self.body
        calls=[]
        def download(request,timeout):
            calls.append(request.full_url)
            return Response(*responses[request.full_url])
        with TemporaryDirectory() as directory:
            path=Path(directory)/'original_ote.xlsx'
            book=Workbook()
            ws=book.active
            ws.append(['Date','Spotová cena plynu OTE (USD/MWh)'])
            for offset in range(30):ws.append([date(2025,12,1)+timedelta(days=offset),20.])
            book.save(path)
            book.close()
            original=path.read_bytes()
            with patch.object(m,'urlopen',side_effect=download):
                first=m.update_ote_workbook(path,today=date(2026,1,4))
                self.assertEqual(calls,[pages[2025],zips[2025],pages[2026],zips[2026],m.CNB_YEARLY_RATE_TEMPLATE.format(year=2025),m.CNB_YEARLY_RATE_TEMPLATE.format(year=2026)])
                self.assertEqual(first['added_rows'],3)
                self.assertEqual(first['last_date'],'2026-01-02')
                self.assertEqual(Path(first['backup_path']).read_bytes(),original)
                reported=m.read_ote(path).set_index('date')
                self.assertEqual(len(reported),33)
                self.assertEqual(reported.loc['2025-12-30','price_mwh'],20.)
                self.assertEqual(reported.loc['2025-12-31','price_mwh'],20.)
                self.assertEqual(reported.loc['2026-01-01','price_mwh'],24.)
                self.assertEqual(reported.loc['2026-01-02','price_mwh'],35.)
                updated=path.read_bytes()
                calls.clear()
                second=m.update_ote_workbook(path,today=date(2026,1,4))
                self.assertEqual(calls,[pages[2026],zips[2026]])
                self.assertEqual(second['added_rows'],0)
                self.assertIsNone(second['backup_path'])
                self.assertEqual(path.read_bytes(),updated)
                self.assertEqual(len(list(Path(directory).glob('original_ote.backup-*.xlsx'))),1)


if __name__=='__main__':unittest.main()
