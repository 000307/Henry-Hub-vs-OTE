"""Fixture-only persistence regressions; no OTE/CNB requests are performed."""
from __future__ import annotations

from datetime import date, datetime, timedelta
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import TestCase, main
from unittest.mock import Mock, patch
from io import BytesIO
import importlib.util
import math
import unittest
import zipfile
import xml.etree.ElementTree as ET

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.datetime import to_excel

_spec = importlib.util.spec_from_file_location("ote_workbook_update", Path(__file__).resolve().parents[1] / "henry_hub_ote_forward.py")
m = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(m)

DATE = "Date"
USD = "Spotová cena plynu OTE (USD/MWh)"
MMBTU = "Spotová cena plynu OTE (USD/MMBtu)"
CZK = "Spotová cena plynu OTE (Kč/MWh)"
RATE = "Kurz ČNB (Kč/USD)"
RATE_DATE = "Datum kurzu ČNB"
SOURCE = "Zdroj nových dat OTE"
CONVERSION = 3.4121416
XML_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


def annual_frame(*pairs):
    return pd.DataFrame([
        {"date": pd.Timestamp(day), "price_czk_mwh": amount,
         "source_url": "https://www.ote-cr.cz/fixture/" + str(pd.Timestamp(day).year) + ".zip"}
        for day, amount in pairs
    ], columns=["date", "price_czk_mwh", "source_url"])


def converted_frame(frame):
    result = frame.copy()
    result["czk_per_usd"] = 20.0
    result["price_mwh"] = result["price_czk_mwh"] / 20.0
    result["price_mmbtu"] = result["price_mwh"] / CONVERSION
    result["cnb_rate_date"] = pd.to_datetime(result["date"]) - pd.Timedelta(days=1)
    return result


def inject_cached_values(path, values):
    """Insert known Excel caches without recalculating the fixture's formulas."""
    target = BytesIO()
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(target, "w") as dest:
        for info in source.infolist():
            body = source.read(info.filename)
            if info.filename in values:
                root = ET.fromstring(body)
                for cell in root.iter(XML_NS + "c"):
                    if cell.attrib.get("r") in values[info.filename]:
                        node = cell.find(XML_NS + "v")
                        if node is None:
                            node = ET.SubElement(cell, XML_NS + "v")
                        node.text = str(values[info.filename][cell.attrib["r"]])
                body = ET.tostring(root, encoding="utf-8")
            dest.writestr(info, body)
    path.write_bytes(target.getvalue())


class UpdateOTEPersistenceTests(TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.path = self.root / "ote.xlsx"

    def book(self, rows, headers=None):
        wb = Workbook()
        ws = wb.active
        ws.title = "OTE historie"
        ws.append(headers or [DATE, USD, MMBTU])
        for row in rows:
            ws.append(row)
        wb.save(self.path)
        wb.close()

    def read(self, *, data_only=False):
        wb = load_workbook(self.path, data_only=data_only)
        self.addCleanup(wb.close)
        return wb

    def assert_only_original(self, original):
        self.assertEqual(self.path.read_bytes(), original)
        self.assertEqual([p.name for p in self.root.iterdir()], [self.path.name])

    def test_year_boundary_fetches_each_year_once_and_converts_only_new(self):
        self.book([[datetime(2025, 12, 30), 20, 20 / CONVERSION]])
        frames = {
            2025: annual_frame(("2025-12-30", 99999), ("2025-12-31", 400)),
            2026: annual_frame(("2026-01-01", 500), ("2026-01-02", 600), ("2026-01-03", 700)),
        }
        fetch = Mock(side_effect=lambda year: frames[year])
        convert = Mock(side_effect=converted_frame)
        original = self.path.read_bytes()
        result = m.update_ote_workbook(self.path, today=date(2026, 1, 2), annual_fetcher=fetch, currency_converter=convert)
        self.assertEqual([call.args[0] for call in fetch.call_args_list], [2025, 2026])
        self.assertEqual(convert.call_count, 1)
        self.assertEqual(pd.to_datetime(convert.call_args.args[0]["date"]).dt.strftime("%Y-%m-%d").tolist(),
                         ["2025-12-31", "2026-01-01", "2026-01-02"])
        self.assertEqual(result["added_rows"], 3)
        self.assertEqual(result["updated_rows"], 0)
        self.assertEqual(result["last_date"], "2026-01-02")
        self.assertEqual(Path(result["backup_path"]).read_bytes(), original)
        ws = self.read().worksheets[0]
        self.assertEqual(ws.cell(2, 2).value, 20)
        headers = {cell.value: cell.column for cell in ws[1]}
        self.assertEqual(ws.cell(3, headers[CZK]).value, 400)
        self.assertEqual(ws.cell(3, headers[RATE]).value, 20)
        self.assertEqual(ws.cell(3, headers[USD]).value, 20)
        self.assertAlmostEqual(ws.cell(3, headers[MMBTU]).value, 20 / CONVERSION)
        self.assertEqual(ws.cell(3, headers[SOURCE]).value, "https://www.ote-cr.cz/fixture/2025.zip")
        self.assertEqual(ws.cell(3, headers[RATE_DATE]).value.date(), date(2025, 12, 30))

    def test_current_workbook_is_exact_noop_without_callbacks(self):
        self.book([[datetime(2026, 1, 2), 0, -1]])
        original = self.path.read_bytes()
        fetch, convert = Mock(), Mock()
        result = m.update_ote_workbook(self.path, today=date(2026, 1, 2), annual_fetcher=fetch, currency_converter=convert)
        fetch.assert_not_called()
        convert.assert_not_called()
        self.assertEqual((result["added_rows"], result["updated_rows"]), (0, 0))
        self.assertIsNone(result["backup_path"])
        self.assert_only_original(original)

    def test_annual_data_with_no_new_days_avoids_fx_backup_and_write(self):
        self.book([[datetime(2026, 1, 1), 12, None]])
        original = self.path.read_bytes()
        convert = Mock()
        result = m.update_ote_workbook(self.path, today=date(2026, 1, 3),
            annual_fetcher=Mock(return_value=annual_frame(("2026-01-01", 999))), currency_converter=convert)
        convert.assert_not_called()
        self.assertIsNone(result["backup_path"])
        self.assert_only_original(original)

    def test_repeated_update_is_idempotent(self):
        self.book([[datetime(2026, 1, 1), 10, None]])
        fetch = Mock(return_value=annual_frame(("2026-01-02", 400)))
        convert = Mock(side_effect=converted_frame)
        first = m.update_ote_workbook(self.path, today=date(2026, 1, 2), annual_fetcher=fetch, currency_converter=convert)
        current = self.path.read_bytes()
        names = sorted(p.name for p in self.root.iterdir())
        second = m.update_ote_workbook(self.path, today=date(2026, 1, 2), annual_fetcher=fetch, currency_converter=convert)
        self.assertEqual((second["added_rows"], second["updated_rows"]), (0, 0))
        self.assertIsNone(second["backup_path"])
        self.assertEqual(self.path.read_bytes(), current)
        self.assertEqual(sorted(p.name for p in self.root.iterdir()), names)
        self.assertEqual(convert.call_count, 1)
        self.assertEqual(fetch.call_count, 1)
        self.assertTrue(Path(first["backup_path"]).exists())

    def test_future_blank_dates_do_not_advance_cutoff_and_existing_blank_day_is_filled(self):
        self.book([[datetime(2026, 1, 1), 10, None], [datetime(2026, 1, 2), None, None],
                   [datetime(2026, 12, 31), None, None]])
        result = m.update_ote_workbook(self.path, today=date(2026, 1, 2),
            annual_fetcher=Mock(return_value=annual_frame(("2026-01-02", 400), ("2026-01-03", 500))),
            currency_converter=converted_frame)
        self.assertEqual((result["added_rows"], result["updated_rows"]), (0, 1))
        self.assertEqual(result["last_date"], "2026-01-02")
        ws = self.read().worksheets[0]
        self.assertEqual(ws.max_row, 4)
        self.assertEqual(ws["B3"].value, 20)
        self.assertIsNone(ws["B4"].value)

    def test_preserves_zero_negative_and_prior_prices(self):
        self.book([[datetime(2026, 1, 1), -10, None], [datetime(2026, 1, 2), 0, None]])
        convert = Mock(side_effect=converted_frame)
        m.update_ote_workbook(self.path, today=date(2026, 1, 3),
            annual_fetcher=Mock(return_value=annual_frame(("2026-01-01", 9999), ("2026-01-02", 9999), ("2026-01-03", 600))),
            currency_converter=convert)
        ws = self.read().worksheets[0]
        self.assertEqual([ws["B2"].value, ws["B3"].value, ws["B4"].value], [-10, 0, 30])
        self.assertEqual(len(convert.call_args.args[0]), 1)

    def test_adds_missing_usd_unit_without_backfilling_historical_rows(self):
        self.book([[datetime(2026, 1, 1), 10, "keep"]], headers=[DATE, USD, "Poznámka"])
        m.update_ote_workbook(self.path, today=date(2026, 1, 2),
            annual_fetcher=Mock(return_value=annual_frame(("2026-01-02", 400))), currency_converter=converted_frame)
        ws = self.read().worksheets[0]
        headers = {cell.value: cell.column for cell in ws[1]}
        self.assertEqual(ws["C2"].value, "keep")
        self.assertIsNone(ws.cell(2, headers[MMBTU]).value)
        self.assertAlmostEqual(ws.cell(3, headers[MMBTU]).value, 20 / CONVERSION)
        self.assertEqual(ws.cell(3, headers[USD]).value, 20)

    def test_cached_formula_history_and_other_sheet_and_styles_are_preserved(self):
        self.book([[datetime(2026, 1, 1), "=1+2", None]], headers=[DATE, USD, MMBTU])
        wb = load_workbook(self.path)
        ws = wb.worksheets[0]
        ws["B2"].number_format = "0.00000"
        ws["B2"].font = Font(bold=True, color="00FF0000")
        ws["B2"].fill = PatternFill(fill_type="solid", fgColor="0000FF00")
        ws["D1"] = "Výpočet"
        ws["D2"] = "=B2*2"
        other = wb.create_sheet("Jiný list")
        other["A1"] = "untouched"
        other["B1"] = "=3+4"
        other["A1"].number_format = "@"
        wb.save(self.path)
        wb.close()
        inject_cached_values(self.path, {"xl/worksheets/sheet1.xml": {"B2": 3, "D2": 6},
                                         "xl/worksheets/sheet2.xml": {"B1": 7}})
        before = self.path.read_bytes()
        result = m.update_ote_workbook(self.path, today=date(2026, 1, 2),
            annual_fetcher=Mock(return_value=annual_frame(("2026-01-02", 400))), currency_converter=converted_frame)
        wb = self.read()
        ws = wb.worksheets[0]
        self.assertEqual(ws["B2"].value, "=1+2")
        self.assertEqual(ws["D2"].value, "=B2*2")
        self.assertEqual(ws["B2"].number_format, "0.00000")
        self.assertTrue(ws["B2"].font.bold)
        self.assertEqual(ws["B2"].fill.fgColor.rgb, "0000FF00")
        self.assertEqual(wb["Jiný list"]["B1"].value, "=3+4")
        self.assertEqual(wb["Jiný list"]["A1"].value, "untouched")
        cached = self.read(data_only=True)
        self.assertEqual(cached.worksheets[0]["B2"].value, 3)
        self.assertEqual(cached.worksheets[0]["D2"].value, 6)
        self.assertEqual(cached["Jiný list"]["B1"].value, 7)
        self.assertEqual(Path(result["backup_path"]).read_bytes(), before)

    def test_excel_serial_czech_and_iso_dates(self):
        self.book([[to_excel(datetime(2026, 1, 1)), 10, None], ["2.1.2026", 11, None], ["2026-01-03", 12, None]])
        result = m.update_ote_workbook(self.path, today=date(2026, 1, 4),
            annual_fetcher=Mock(return_value=annual_frame(("2026-01-04", 400))), currency_converter=converted_frame)
        self.assertEqual(result["last_date"], "2026-01-04")
        self.assertEqual(result["added_rows"], 1)

    def test_current_year_unpublished_is_noop_with_note(self):
        self.book([[datetime(2026, 1, 1), 10, None]])
        original = self.path.read_bytes()
        convert = Mock()
        result = m.update_ote_workbook(self.path, today=date(2026, 1, 2),
            annual_fetcher=Mock(return_value=None), currency_converter=convert)
        self.assertTrue(any("zatím nezveřejněno" in str(note).lower() for note in result["notes"]))
        convert.assert_not_called()
        self.assert_only_original(original)

    def test_missing_past_year_errors_before_any_mutation(self):
        self.book([[datetime(2025, 12, 30), 10, None]])
        original = self.path.read_bytes()
        convert = Mock()
        with self.assertRaises(Exception):
            m.update_ote_workbook(self.path, today=date(2026, 1, 2),
                annual_fetcher=Mock(return_value=None), currency_converter=convert)
        convert.assert_not_called()
        self.assert_only_original(original)

    def test_invalid_date_with_finite_existing_price_errors(self):
        self.book([["neplatné datum", 10, None]])
        original = self.path.read_bytes()
        fetch = Mock()
        with self.assertRaises(Exception):
            m.update_ote_workbook(self.path, today=date(2026, 1, 2), annual_fetcher=fetch, currency_converter=converted_frame)
        fetch.assert_not_called()
        self.assert_only_original(original)

    def test_fetch_failure_preserves_original_and_directory(self):
        self.book([[datetime(2026, 1, 1), 10, None]])
        original = self.path.read_bytes()
        with self.assertRaisesRegex(RuntimeError, "fetch failed"):
            m.update_ote_workbook(self.path, today=date(2026, 1, 2),
                annual_fetcher=Mock(side_effect=RuntimeError("fetch failed")), currency_converter=converted_frame)
        self.assert_only_original(original)

    def test_fx_failure_preserves_original_and_directory(self):
        self.book([[datetime(2026, 1, 1), 10, None]])
        original = self.path.read_bytes()
        with self.assertRaisesRegex(RuntimeError, "fx failed"):
            m.update_ote_workbook(self.path, today=date(2026, 1, 2),
                annual_fetcher=Mock(return_value=annual_frame(("2026-01-02", 400))),
                currency_converter=Mock(side_effect=RuntimeError("fx failed")))
        self.assert_only_original(original)

    def test_staged_save_failure_preserves_original_and_directory(self):
        self.book([[datetime(2026, 1, 1), 10, None]])
        original = self.path.read_bytes()
        with patch("openpyxl.workbook.workbook.Workbook.save", side_effect=OSError("save failed")):
            with self.assertRaisesRegex(OSError, "save failed"):
                m.update_ote_workbook(self.path, today=date(2026, 1, 2),
                    annual_fetcher=Mock(return_value=annual_frame(("2026-01-02", 400))), currency_converter=converted_frame)
        self.assert_only_original(original)

    def test_backups_are_unique_and_represent_previous_originals(self):
        self.book([[datetime(2026, 1, 1), 10, None]])
        original = self.path.read_bytes()
        first = m.update_ote_workbook(self.path, today=date(2026, 1, 2),
            annual_fetcher=Mock(return_value=annual_frame(("2026-01-02", 400))), currency_converter=converted_frame)
        next_original = self.path.read_bytes()
        second = m.update_ote_workbook(self.path, today=date(2026, 1, 3),
            annual_fetcher=Mock(return_value=annual_frame(("2026-01-02", 400), ("2026-01-03", 500))), currency_converter=converted_frame)
        self.assertNotEqual(first["backup_path"], second["backup_path"])
        self.assertEqual(Path(first["backup_path"]).read_bytes(), original)
        self.assertEqual(Path(second["backup_path"]).read_bytes(), next_original)

    def test_bad_new_price_is_rejected_before_fx_or_save(self):
        self.book([[datetime(2026, 1, 1), 10, None]])
        original = self.path.read_bytes()
        convert = Mock()
        with self.assertRaises(Exception):
            m.update_ote_workbook(self.path, today=date(2026, 1, 2),
                annual_fetcher=Mock(return_value=annual_frame(("2026-01-02", float("inf")))), currency_converter=convert)
        convert.assert_not_called()
        self.assert_only_original(original)

    def test_nonfinite_converted_price_is_rejected_before_write(self):
        self.book([[datetime(2026, 1, 1), 10, None]])
        original = self.path.read_bytes()
        def bad_converter(frame):
            out = converted_frame(frame)
            out.loc[out.index[0], "price_mwh"] = float("nan")
            return out
        with self.assertRaises(Exception):
            m.update_ote_workbook(self.path, today=date(2026, 1, 2),
                annual_fetcher=Mock(return_value=annual_frame(("2026-01-02", 400))), currency_converter=bad_converter)
        self.assert_only_original(original)


    def test_formula_cache_restore_failure_preserves_original_and_directory(self):
        self.book([[datetime(2026, 1, 1), 10, None]])
        original = self.path.read_bytes()
        with patch.object(m, "_ote_restore_formula_caches", side_effect=OSError("cache restore failed")):
            with self.assertRaisesRegex(OSError, "cache restore failed"):
                m.update_ote_workbook(self.path, today=date(2026, 1, 2),
                    annual_fetcher=Mock(return_value=annual_frame(("2026-01-02", 400))), currency_converter=converted_frame)
        self.assert_only_original(original)

    def test_atomic_commit_failure_removes_staged_and_backup_without_changing_original(self):
        self.book([[datetime(2026, 1, 1), 10, None]])
        original = self.path.read_bytes()
        with patch.object(m.os, "replace", side_effect=OSError("replace failed")):
            with self.assertRaisesRegex(OSError, "replace failed"):
                m.update_ote_workbook(self.path, today=date(2026, 1, 2),
                    annual_fetcher=Mock(return_value=annual_frame(("2026-01-02", 400))), currency_converter=converted_frame)
        self.assert_only_original(original)

    def test_missing_mwh_header_is_added_without_backfilling_mmbtu_history(self):
        self.book([[datetime(2026, 1, 1), 2]], headers=[DATE, MMBTU])
        m.update_ote_workbook(self.path, today=date(2026, 1, 2),
            annual_fetcher=Mock(return_value=annual_frame(("2026-01-02", 400))), currency_converter=converted_frame)
        ws = self.read().worksheets[0]
        headers = {cell.value: cell.column for cell in ws[1]}
        self.assertEqual(ws["B2"].value, 2)
        self.assertIsNone(ws.cell(2, headers[USD]).value)
        self.assertEqual(ws.cell(3, headers[USD]).value, 20)
        self.assertAlmostEqual(ws.cell(3, headers[MMBTU]).value, 20 / CONVERSION)

    def test_conflicting_source_prices_abort_without_fx_or_mutation(self):
        self.book([[datetime(2026, 1, 1), 10, None]])
        original = self.path.read_bytes()
        convert = Mock()
        with self.assertRaises(Exception):
            m.update_ote_workbook(self.path, today=date(2026, 1, 2),
                annual_fetcher=Mock(return_value=annual_frame(("2026-01-02", 400), ("2026-01-02", 500))),
                currency_converter=convert)
        convert.assert_not_called()
        self.assert_only_original(original)

    def test_new_date_with_existing_uncached_formula_is_not_overwritten(self):
        self.book([[datetime(2026, 1, 1), 10, None], [datetime(2026, 1, 2), "=1+2", None]])
        original = self.path.read_bytes()
        with self.assertRaises(Exception):
            m.update_ote_workbook(self.path, today=date(2026, 1, 2),
                annual_fetcher=Mock(return_value=annual_frame(("2026-01-02", 400))), currency_converter=converted_frame)
        self.assert_only_original(original)

if __name__ == "__main__":
    main(verbosity=2)
