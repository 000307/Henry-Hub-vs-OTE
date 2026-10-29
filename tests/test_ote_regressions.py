"""Regression cases for annual-data validation and workbook/report boundaries."""
from datetime import date, datetime, timedelta
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import Mock, patch
import importlib.util
import unittest
from zipfile import ZipFile

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles.numbers import is_date_format

spec=importlib.util.spec_from_file_location('report_ote_regressions',Path(__file__).resolve().parents[1]/'henry_hub_ote_forward.py')
m=importlib.util.module_from_spec(spec)
spec.loader.exec_module(m)


def archive(year,values):
    wb=Workbook()
    ws=wb.active
    ws.title='Odchylky NC BAL'
    ws.append(['Plynárenský den','Použitelná cena pro záporné vyrovnávací množství (Kč/MWh)'])
    for number,value in enumerate(values,1):
        ws.append([datetime(year,1,number),value])
    stream=BytesIO()
    wb.save(stream)
    result=BytesIO()
    with ZipFile(result,'w') as z:
        z.writestr(f'Rocni_zprava_o_trhu_plyn_{year}_V0.xlsx',stream.getvalue())
    return result.getvalue()


class OteBoundaryRegressions(unittest.TestCase):
    def test_only_genuine_missing_current_prices_can_represent_empty_template(self):
        year=pd.Timestamp.now(tz='Europe/Prague').year
        for missing in ('ND',None,'','—'):
            with self.subTest(missing=missing):
                self.assertTrue(m.parse_ote_annual_zip(archive(year,[missing]),year).empty)
        with self.assertRaises(ValueError):
            m.parse_ote_annual_zip(archive(year-1,['ND']),year-1)

    def test_invalid_current_prices_are_never_silently_treated_as_no_data(self):
        year=pd.Timestamp.now(tz='Europe/Prague').year
        for bad in ('garbage','inf','-inf','1.2junk'):
            for values in ([bad],[500,bad]):
                with self.subTest(values=values),self.assertRaisesRegex(ValueError,'neplatnou nebo nekonečnou cenu'):
                    m.parse_ote_annual_zip(archive(year,values),year)

    def test_new_date_remains_readable_when_history_uses_text_number_format(self):
        with TemporaryDirectory() as directory:
            path=Path(directory)/'ote.xlsx'
            wb=Workbook()
            ws=wb.active
            ws.append(['Date','Spotová cena plynu OTE (USD/MWh)'])
            for number in range(30):
                ws.append([(date(2026,1,1)+timedelta(days=number)).isoformat(),20.])
                ws.cell(number+2,1).number_format='@'
            wb.save(path)
            wb.close()
            original=path.read_bytes()
            source=pd.DataFrame({'date':[pd.Timestamp('2026-01-31')],'price_czk_mwh':[500.],
                                 'source_url':['https://www.ote-cr.cz/fixture.zip']})
            def convert(frame):
                result=frame.copy()
                result['czk_per_usd']=25.
                result['price_mwh']=result['price_czk_mwh']/25.
                result['price_mmbtu']=result['price_mwh']/m.MMBTU_PER_MWH
                result['cnb_rate_date']=pd.Timestamp('2026-01-30')
                return result
            result=m.update_ote_workbook(path,today=date(2026,1,31),annual_fetcher=lambda year:source,currency_converter=convert)
            self.assertEqual(result['added_rows'],1)
            self.assertEqual(Path(result['backup_path']).read_bytes(),original)
            wb=load_workbook(path)
            try:
                self.assertTrue(is_date_format(wb.active['A32'].number_format))
                self.assertEqual(wb.active['A2'].value,'2026-01-01')
                self.assertEqual(wb.active['A2'].number_format,'@')
            finally:wb.close()
            reported=m.read_ote(path)
            self.assertEqual(len(reported),31)
            self.assertEqual(reported['date'].max(),pd.Timestamp('2026-01-31'))

    def test_historical_v2_fallback_uses_prague_calendar_year(self):
        fake=Mock()
        fake.Timestamp.now.return_value=pd.Timestamp('2027-01-01',tz='Europe/Prague')
        with patch.object(m,'pd',fake):
            self.assertEqual(m._ote_pick({'historical2026V2':2},2026,'ZIP'),'historical2026V2')
        fake.Timestamp.now.assert_called_once_with(tz='Europe/Prague')


if __name__=='__main__':unittest.main()
